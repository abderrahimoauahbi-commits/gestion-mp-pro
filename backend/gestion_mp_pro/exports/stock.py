"""Classeurs du stock (routes 106 et 236) : le journal des mouvements et la liste des lots, avec
les mêmes filtres que l'écran (ARCHITECTURE § 11).

Conventions de Gestion MP, en attendant `exports/classeur.py` (lot D9) : en-tête en gras, volets
figés, filtre automatique, vraies dates Excel, kilos `#,##0.000`, montants `#,##0.00`, texte
commençant par `=`, `+`, `-` ou `@` gardé en texte. **Sans prix ni valeur** pour qui ne les voit
pas (barrière 4 : un classeur n'est pas filtré par la route).
"""

from __future__ import annotations

from collections.abc import Callable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from gestion_mp_pro.services.lots import STATUTS_LOT, LotSolde
from gestion_mp_pro.services.mouvements import LigneMouvement

FORMATS = {"kg": "#,##0.000", "montant": "#,##0.00", "prix": "#,##0.0000", "date": "DD/MM/YYYY"}
DANGEREUX = ("=", "+", "-", "@")


@dataclass(frozen=True)
class Colonne:
    """Une colonne : titre, lecture de la valeur, format, monétaire (absente sans les prix)."""

    titre: str
    valeur: Callable[[Any], object]
    format: str = ""
    monetaire: bool = False
    largeur: int = 14


def _cellule(valeur: object) -> object:
    """Une valeur pour Excel : nombres en `float`, instants sans fuseau."""
    if isinstance(valeur, Decimal):
        return float(valeur)
    if isinstance(valeur, datetime):
        return valeur.replace(tzinfo=None)
    return valeur


def classeur(
    titre: str, colonnes: Sequence[Colonne], lignes: Sequence[Any], voit_prix: bool
) -> bytes:
    """Un classeur d'une feuille ; les colonnes monétaires seulement si `voit_prix`."""
    gardees = [c for c in colonnes if voit_prix or not c.monetaire]
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101 — un classeur neuf a sa feuille
    ws.title = titre[:31]
    ws.append([c.titre for c in gardees])
    for cellule in ws[1]:
        cellule.font = Font(bold=True)
    for ligne in lignes:
        ws.append([_cellule(c.valeur(ligne)) for c in gardees])
        for cellule in ws[ws.max_row]:
            if isinstance(cellule.value, str) and cellule.value.startswith(DANGEREUX):
                cellule.data_type = "s"  # jamais une formule
    for i, c in enumerate(gardees, start=1):
        lettre = get_column_letter(i)
        ws.column_dimensions[lettre].width = c.largeur
        if c.format:
            for (cellule,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cellule.number_format = FORMATS[c.format]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    sortie = BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


COLONNES_MOUVEMENTS = (
    Colonne("Date", lambda x: x.date, "date", largeur=12),
    Colonne("Saisi le", lambda x: x.saisi_le, "date", largeur=12),
    Colonne("Mouvement", lambda x: x.libelle, largeur=24),
    Colonne("Document", lambda x: x.document_numero, largeur=16),
    Colonne("Référence", lambda x: x.reference, largeur=30),
    Colonne("Lot", lambda x: x.lot, largeur=18),
    Colonne("Magasin", lambda x: x.magasin, largeur=18),
    Colonne("Quantité (kg)", lambda x: x.quantite, "kg"),
    Colonne("Prix / kg", lambda x: x.prix, "prix", monetaire=True),
    Colonne("Devise", lambda x: x.devise, monetaire=True, largeur=8),
    Colonne("Cours", lambda x: x.cours, "prix", monetaire=True),
    Colonne("Coût unitaire (MAD)", lambda x: x.cout_mad, "prix", monetaire=True),
    Colonne("Valeur (MAD)", lambda x: x.valeur_mad, "montant", monetaire=True),
    Colonne("Motif", lambda x: x.motif, largeur=24),
    Colonne("Commentaire", lambda x: x.commentaire, largeur=30),
    Colonne("Par", lambda x: x.auteur, largeur=18),
    Colonne("Annule le mouvement n°", lambda x: x.contrepasse_id),
)
"""Route 106."""


def mouvements(lignes: Sequence[LigneMouvement], voit_prix: bool) -> bytes:
    """Le journal des mouvements, le plus récent d'abord."""
    return classeur("Mouvements", COLONNES_MOUVEMENTS, lignes, voit_prix)


def _magasins(x: LotSolde) -> str:
    return ", ".join(f"{m.magasin} : {m.quantite:.3f}".replace(".", ",") for m in x.magasins)


COLONNES_LOTS = (
    Colonne("Lot", lambda x: x.lot, largeur=18),
    Colonne("Statut", lambda x: STATUTS_LOT[x.statut], largeur=22),
    Colonne("Référence", lambda x: x.reference, largeur=30),
    Colonne("Désignation", lambda x: x.designation, largeur=24),
    Colonne("Rôle", lambda x: x.role, largeur=10),
    Colonne("Fournisseur", lambda x: x.fournisseur, largeur=22),
    Colonne("Entré le", lambda x: x.entre_le, "date", largeur=12),
    Colonne("Dernier mouvement", lambda x: x.historique.dernier_mouvement, "date"),
    Colonne("Entré (kg)", lambda x: x.historique.entre, "kg"),
    Colonne("Sorti (kg)", lambda x: x.historique.sorti, "kg"),
    Colonne("Reste (kg)", lambda x: x.quantite, "kg"),
    Colonne("Disponible (kg)", lambda x: x.disponible, "kg"),
    Colonne("En quarantaine (kg)", lambda x: x.quarantaine, "kg"),
    Colonne("Bobines", lambda x: x.bobines),
    Colonne("Magasins", _magasins, largeur=30),
    Colonne(
        "Ordres servis", lambda x: ", ".join(o.numero for o in x.historique.ordres), largeur=30
    ),
)
"""Route 236 (aucune colonne monétaire : la même pour tous les rôles)."""


def lots(lignes: Sequence[LotSolde], voit_prix: bool) -> bytes:
    """La liste des lots, telle que l'onglet Lots la filtre."""
    return classeur("Lots", COLONNES_LOTS, lignes, voit_prix)


def nom_de_fichier(prefixe: str, jour: date) -> str:
    """« mouvements-2026-10-08.xlsx »."""
    return f"{prefixe}-{jour.isoformat()}.xlsx"
