"""Le bon fournisseur dans un classeur (route 165) : exactement le contenu du document HTML
(`services/bon_fournisseur`), en anglais, sans aucun prix, montant ni devise, sur le papier de
l'usine (raison sociale en rouge, « HEAD OFFICE », pied « BANKS » et identifiants légaux).
Mise en page A4 portrait, une page de large, en-tête du tableau répété.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font
from openpyxl.worksheet.worksheet import Worksheet

from gestion_mp_pro.api.erreurs import Refus
from gestion_mp_pro.documents.rendu import LIBELLES, Societe
from gestion_mp_pro.exports.classeur import DANGEREUX, FOND_ENTETE, GRAS, TRAIT, cellule
from gestion_mp_pro.services.bon_fournisseur import (
    CONFIRMATION,
    SIGNATURE_ACHETEUR,
    SIGNATURE_FOURNISSEUR,
    TITRE,
    BonFournisseur,
)

ROUGE = "D40000"
LARGEURS = (8, 24, 34, 24, 10, 16, 18)
FORMAT_KG = "#,##0.###"
MOTS = LIBELLES["en"]


def _ligne(ws: Worksheet, valeurs: list[object], police: Font | None = None) -> int:
    """Ajoute une ligne (texte jamais formule) ; rend son rang."""
    ws.append([cellule(v) for v in valeurs])
    rang = ws.max_row
    for c in ws[rang]:
        if isinstance(c.value, str) and c.value.startswith(DANGEREUX):
            c.data_type = "s"
        if police is not None and c.value is not None:
            c.font = police
    return rang


def _papier(ws: Worksheet, s: Societe) -> None:
    """L'en-tête du papier de l'usine (lignes vides omises)."""
    _ligne(ws, [s.nom], Font(name="Times New Roman", size=18, color=ROUGE))
    lignes = [s.secteur, s.forme_juridique, f"{MOTS['siege']}   {s.adresse}" if s.adresse else "",
              s.ville, s.telephone]  # fmt: skip
    for texte in lignes:
        if texte:
            _ligne(ws, [texte], Font(name="Times New Roman", size=11))
    ws.append([])


def _blocs(ws: Worksheet, bon: BonFournisseur) -> None:
    """Le titre, les dates, le fournisseur, où livrer."""
    _ligne(ws, [TITRE, bon.titre], Font(bold=True, size=13))
    for libelle, valeur in (("Date", bon.date), ("Expected delivery", bon.livraison)):
        ws.cell(_ligne(ws, [libelle, valeur]), 1).font = GRAS
    ws.append([])
    _ligne(ws, ["Supplier"], Font(bold=True, size=12))
    for libelle, valeur in [("Name", bon.fournisseur), *bon.fournisseur_lignes]:
        ws.cell(_ligne(ws, [libelle, valeur]), 1).font = GRAS
    ws.append([])
    _ligne(ws, ["Deliver to"], Font(bold=True, size=12))
    for libelle, valeur in bon.livrer_a:
        ws.cell(_ligne(ws, [libelle, valeur]), 1).font = GRAS
    ws.append([])


def _tableau(ws: Worksheet, bon: BonFournisseur) -> int:
    """Les lignes et le total ; rend le rang de l'en-tête du tableau."""
    entete = _ligne(ws, list(bon.colonnes))
    for c in ws[entete]:
        c.font, c.fill, c.border = GRAS, FOND_ENTETE, Border(bottom=TRAIT)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for x in bon.lignes:
        valeurs: list[object] = [
            x.rang,
            x.reference,
            x.designation,
            x.conditionnement,
            x.unites,
            x.kg,
        ]
        if bon.avec_metres:
            valeurs.append(round(x.metres) if x.metres is not None else None)
        rang = _ligne(ws, valeurs)
        ws.cell(rang, 5).number_format = FORMAT_KG
        ws.cell(rang, 6).number_format = FORMAT_KG
        if bon.avec_metres:
            ws.cell(rang, 7).number_format = "#,##0"
    total = _ligne(ws, ["", "Total ordered", "", "", None, bon.total_kg])
    for c in ws[total]:
        c.font, c.border = GRAS, Border(top=TRAIT)
    ws.cell(total, 6).number_format = FORMAT_KG
    return entete


def _fin(ws: Worksheet, bon: BonFournisseur, societe: Societe) -> None:
    """Remarques, confirmation attendue, signatures, pied du papier."""
    if bon.remarques:
        ws.append([])
        ws.cell(_ligne(ws, ["Remarks", bon.remarques]), 1).font = GRAS
    ws.append([])
    _ligne(ws, [CONFIRMATION])
    ws.append([])
    _ligne(ws, [f"For {societe.nom or 'the buyer'}", "", "", SIGNATURE_FOURNISSEUR[0]])
    _ligne(ws, [SIGNATURE_ACHETEUR, "", "", SIGNATURE_FOURNISSEUR[1]])
    pied = []
    if societe.banques:
        pied.append(f"{MOTS['banques']}       " + "     ".join(societe.banques))
    if societe.identifiants:
        pied.append("     ".join(f"{lib} : {v}" for lib, v in societe.identifiants))
    if pied:
        ws.append([])
    for texte in pied:
        _ligne(ws, [texte], Font(name="Times New Roman", size=10))


def classeur(bon: BonFournisseur, societe: Societe) -> bytes:
    """Le classeur du bon ; 409 sans raison sociale : le bon ne part pas sans elle (A-18)."""
    if not societe.reglee:
        raise Refus.etat(
            "La raison sociale de l'usine n'est pas réglée (Réglages › Paramètres › Société) : "
            "le bon ne part pas sans elle."
        )
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101 — un classeur neuf a sa feuille
    ws.title = TITRE
    _papier(ws, societe)
    _blocs(ws, bon)
    entete = _tableau(ws, bon)
    _fin(ws, bon, societe)
    for i, largeur in enumerate(LARGEURS[: len(bon.colonnes)]):
        ws.column_dimensions["ABCDEFG"[i]].width = largeur
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True  # type: ignore[union-attr]
    ws.print_title_rows = f"{entete}:{entete}"
    sortie = BytesIO()
    wb.save(sortie)
    return sortie.getvalue()
