"""Les conventions communes des classeurs Excel (ARCHITECTURE § 11 ; Gestion MP `exports.py`).

Un classeur exporté sert à travailler ailleurs — trier, filtrer, envoyer. Chaque feuille a donc :
quelques lignes de titre (« … au JJ/MM/AAAA », les filtres de l'écran), l'en-tête en gras sur
fond gris, figé et filtrable ; des nombres restés nombres, au format de leur grandeur (§ 9.3) ;
de **vraies dates** Excel (triables), les instants à l'heure de l'usine ; une largeur lisible ;
et rien qui puisse s'exécuter à l'ouverture : un texte qui commence par `=`, `+`, `-` ou `@`
reste un texte, les caractères de contrôle refusés par Excel sont retirés.

**Aucun prix pour qui ne les voit pas** (barrière 4, ARCHITECTURE § 5.3) : une colonne
`monetaire=True` disparaît du classeur quand `voit_prix` est faux, et, dans un classeur pour qui
ne voit pas les prix, une colonne dont le titre dit un prix, un montant, une valeur, un coût ou
un cours sans être déclarée monétaire est une faute de programmation (`ValueError`, vue par les
essais du lot).

    from gestion_mp_pro.exports.classeur import Colonne, Feuille, classeur, reponse

    COLONNES = (Colonne("Date", lambda x: x.date, "date"),
                Colonne("Quantité (kg)", lambda x: x.kg, "kg"),
                Colonne("Valeur (MAD)", lambda x: x.valeur_mad, "montant", monetaire=True))
    contenu = classeur([Feuille("Stock", COLONNES, lignes, entete=[titre("Stock", jour)])],
                       voit_prix=qui.voit_prix)
    return reponse(contenu, "stock", jour)
"""

from __future__ import annotations

import re
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from enum import Enum
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from starlette.responses import Response

from gestion_mp_pro.documents.format import a_l_usine, date_fr

TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

FORMATS = {
    "kg": "#,##0.0",
    "kg_precis": "#,##0.000",
    "entier": "#,##0",
    "m2": "#,##0",
    "montant": "#,##0.00",
    "prix": "#,##0.00##",
    "cours": "0.0000",
    "pourcent": "0.0%",
    "date": "DD/MM/YYYY",
    "horodatage": "DD/MM/YYYY HH:MM",
}
"""Le format Excel de chaque grandeur (ARCHITECTURE § 9.3) ; `pourcent` attend une fraction."""

GRAVITES = {
    "bloquant": "Bloquant",
    "avertissement": "À corriger",
    "information": "Pour information",
}
"""Les gravités des contrôles dites en français (comme l'écran)."""

DANGEREUX = ("=", "+", "-", "@")
TITRE_MONETAIRE = re.compile(
    r"\b(prix|montants?|valeurs?|co[uû]ts?|cours|cmup|budget|revient|tva|frais|[ée]conomie"
    r"|droits de douane|MAD|USD|EUR)\b",
    re.IGNORECASE,
)
GRAS = Font(bold=True)
FOND_ENTETE = PatternFill("solid", fgColor="E8ECF3")
TRAIT = Side(style="thin", color="B8C0CC")
LARGEUR_MAX = 50
ECHANTILLON = 300
COLONNES_AVANT_VOLET = 4
"""Au-delà de 4 colonnes, la première (ce dont parle la ligne) reste figée à gauche."""


@dataclass(frozen=True)
class Colonne:
    """Une colonne : titre, lecture de la valeur dans une ligne, grandeur (clé de `FORMATS` ou
    format Excel), monétaire (absente pour qui ne voit pas les prix), largeur (sinon calculée)."""

    titre: str
    valeur: Callable[[Any], object]
    format: str = ""
    monetaire: bool = False
    largeur: int | None = None


@dataclass(frozen=True)
class Feuille:
    """Une feuille tabulaire : son titre d'onglet, ses colonnes, ses lignes, et les lignes de titre
    écrites au-dessus de l'en-tête (`titre(...)`, `filtres(...)`)."""

    titre: str
    colonnes: Sequence[Colonne]
    lignes: Sequence[Any]
    entete: Sequence[str] = field(default_factory=tuple)


def titre(texte: str, jour: date) -> str:
    """« Journal des actions au 08/10/2026 »."""
    return f"{texte} au {date_fr(jour)}"


def exporte_le(texte: str, jour: date) -> str:
    """« Fournisseurs — exporté le 08/10/2026 » (une liste sans date de référence)."""
    return f"{texte} — exporté le {date_fr(jour)}"


def filtres(valeurs: Iterable[tuple[str, object]]) -> str:
    """« Filtres — par : Karim · du : 01/10/2026 » (les filtres vides ne s'écrivent pas) ;
    « » si aucun filtre n'est posé."""
    dits = [
        f"{nom} : {date_fr(v) if isinstance(v, date) else v}"
        for nom, v in valeurs
        if v not in (None, "")
    ]
    return f"Filtres — {' · '.join(dits)}" if dits else ""


def cellule(valeur: object) -> object:
    """Une valeur sûre pour Excel : nombres en `float`, instants à l'heure de l'usine sans
    fuseau (Excel n'en connaît pas), texte sans caractère de contrôle, booléen « oui » / « »,
    énumération sa valeur."""
    if isinstance(valeur, bool):
        return "oui" if valeur else ""
    if isinstance(valeur, Enum):
        return valeur.value
    if isinstance(valeur, Decimal):
        return float(valeur)
    if isinstance(valeur, datetime):
        return a_l_usine(valeur).replace(tzinfo=None, microsecond=0)
    if isinstance(valeur, str):
        return ILLEGAL_CHARACTERS_RE.sub("", valeur)
    return valeur


def colonnes_permises(colonnes: Sequence[Colonne], voit_prix: bool) -> list[Colonne]:
    """Les colonnes que ce lecteur reçoit ; pour qui ne voit pas les prix, refuse une colonne
    dont le titre dit un montant sans qu'elle soit déclarée monétaire."""
    for c in colonnes:
        if not voit_prix and not c.monetaire and TITRE_MONETAIRE.search(c.titre):
            raise ValueError(f"colonne « {c.titre} » : monétaire, à déclarer `monetaire=True`")
    return [c for c in colonnes if voit_prix or not c.monetaire]


def _nom_d_onglet(texte: str, pris: set[str]) -> str:
    """Un nom d'onglet permis par Excel (31 caractères, sans `[]:*?/\\`), unique."""
    base = re.sub(r"[\[\]:*?/\\]", " ", texte).strip()[:31] or "Feuille"
    nom, rang = base, 2
    while nom.lower() in pris:
        suffixe = f" ({rang})"
        nom, rang = base[: 31 - len(suffixe)] + suffixe, rang + 1
    pris.add(nom.lower())
    return nom


def _ajouter(ws: Worksheet, valeurs: Sequence[object]) -> None:
    """Une ligne ; un texte qui ressemble à une formule reste un texte."""
    ws.append([cellule(v) for v in valeurs])
    for c in ws[ws.max_row]:
        if isinstance(c.value, str) and c.value.startswith(DANGEREUX):
            c.data_type = "s"


def _format(colonne: Colonne) -> str:
    return FORMATS.get(colonne.format, colonne.format)


def _largeur(colonne: Colonne, valeurs: Sequence[object]) -> float:
    """Largeur lisible : le titre, ou les valeurs les plus longues (300 premières lignes)."""
    if colonne.largeur is not None:
        return colonne.largeur
    if _format(colonne) in (FORMATS["date"], FORMATS["horodatage"]):
        return 12 if _format(colonne) == FORMATS["date"] else 17
    longueurs = [len(str(v)) for v in valeurs[:ECHANTILLON] if v is not None]
    return min(LARGEUR_MAX, max([len(colonne.titre) * 0.9, *longueurs, 8]) + 2)


def ecrire_feuille(wb: Workbook, f: Feuille, voit_prix: bool, pris: set[str]) -> Worksheet:
    """Ajoute une feuille au classeur, aux conventions communes."""
    gardees = colonnes_permises(f.colonnes, voit_prix)
    ws = wb.create_sheet(_nom_d_onglet(f.titre, pris))
    lignes_titre = [t for t in f.entete if t]
    for texte in lignes_titre:
        _ajouter(ws, [texte])
        ws.cell(ws.max_row, 1).font = GRAS
    if lignes_titre:
        ws.append([])
    _ajouter(ws, [c.titre for c in gardees])
    rang_entete = ws.max_row  # après la ligne vide (qu'`append([])` ne compte pas)
    for c in ws[rang_entete]:
        c.font, c.fill, c.border = GRAS, FOND_ENTETE, Border(bottom=TRAIT)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    valeurs = [[c.valeur(ligne) for c in gardees] for ligne in f.lignes]
    for v in valeurs:
        _ajouter(ws, v)
    for i, c in enumerate(gardees, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _largeur(c, [v[i - 1] for v in valeurs])
        if c.format:
            for (x,) in ws.iter_rows(min_row=rang_entete + 1, min_col=i, max_col=i):
                if isinstance(x.value, int | float | date):
                    x.number_format = _format(c)
    ws.freeze_panes = ws.cell(rang_entete + 1, 2 if len(gardees) > COLONNES_AVANT_VOLET else 1)
    if valeurs and gardees:
        dernier = get_column_letter(len(gardees))
        ws.auto_filter.ref = f"A{rang_entete}:{dernier}{ws.max_row}"
    return ws


def classeur(feuilles: Sequence[Feuille], voit_prix: bool) -> bytes:
    """Le classeur (xlsx) de ces feuilles, dans l'ordre."""
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]
    pris: set[str] = set()
    for f in feuilles:
        ecrire_feuille(wb, f, voit_prix, pris)
    sortie = BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


def nom_de_fichier(prefixe: str, jour: date) -> str:
    """« journal-des-actions-2026-10-08.xlsx » (lettres, chiffres, tirets seulement)."""
    propre = re.sub(r"[^A-Za-z0-9._-]+", "-", prefixe).strip("-") or "export"
    return f"{propre}-{jour.isoformat()}.xlsx"


def reponse(contenu: bytes, prefixe: str, jour: date) -> Response:
    """Le classeur à télécharger (`attachment`)."""
    return Response(
        contenu,
        media_type=TYPE_XLSX,
        headers={"Content-Disposition": f'attachment; filename="{nom_de_fichier(prefixe, jour)}"'},
    )


REPONSES_XLSX: dict[int | str, dict[str, Any]] = {200: {"content": {TYPE_XLSX: {}}}}
"""À fusionner dans les `responses` d'une route d'export (OpenAPI : un classeur, pas du JSON)."""
