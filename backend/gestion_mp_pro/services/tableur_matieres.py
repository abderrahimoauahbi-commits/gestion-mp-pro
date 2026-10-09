"""Le catalogue par tableur : aperçu puis application (routes 79–80 ; R-REF-49…54, A-03, P7-9…12).

Chaque ligne passe par **les mêmes règles que la fiche** (`matieres.verifier`, bornes et
longueurs comprises, rôle d'une matière en recette, A-02) : ce que l'aperçu accepte,
l'application l'écrit, sans erreur serveur. Une référence en double dans le fichier est une
erreur (casse ignorée) ; « nan » ou l'infini ne sont pas des nombres ; « En service » vaut oui
ou non (une faute de frappe ne désactive rien) ; les codes se retrouvent sans tenir compte de
la casse. Une case vide ne change rien ; la colonne « Vider » nomme les colonnes à vider ; la
colonne « Version » (celle du modèle) refuse une ligne dont la fiche a changé depuis.

L'aperçu rend une **empreinte** du plan d'écriture ; l'application refait le plan et refuse
(409) s'il diffère : on écrit exactement ce qu'on a vu. Une matière absente du tableur n'est
jamais supprimée. Tout ou rien.
"""

from __future__ import annotations

import hashlib
import json
import math
from collections.abc import Iterator
from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import select
from sqlalchemy.orm import Session

from gestion_mp_pro.api.erreurs import Refus
from gestion_mp_pro.modeles import Matiere, Utilisateur
from gestion_mp_pro.modeles import longueurs as L
from gestion_mp_pro.modeles.enums import ObjetJournal
from gestion_mp_pro.services import audit
from gestion_mp_pro.services.matieres import BORNES, reference_libre, verifier
from gestion_mp_pro.services.tableur_referentiels import (
    ACTIF,
    COLONNES,
    INTITULES,
    LIBELLES,
    NE_SE_VIDE_PAS,
    VERSION,
    VIDE,
    VIDER,
    Referentiels,
    genre,
    referentiels,
)

FEUILLE = "Matières"
LIGNES_MAX = 10000
ANCIENS = {"Poids d'une bobine": "kg / bobine"}
LONGUEURS = {
    "reference": L.REFERENCE,
    "designation": L.DESIGNATION,
    "couleur_code": L.COULEUR_CODE,
    "unite_achat": L.UNITE_ACHAT,
    "groupe_equiv": L.GROUPE_EQUIV,
}


@dataclass
class LigneLue:
    """Une ligne du tableur : ce qu'elle change, ou ce qui ne va pas."""

    ligne: int
    reference: str
    etat: str = "inchangee"
    matiere_id: int | None = None
    version: int | None = None
    ecrire: dict[str, Any] = field(default_factory=dict)
    changements: list[dict[str, str]] = field(default_factory=list)
    erreurs: list[str] = field(default_factory=list)


@dataclass
class Apercu:
    """Le plan de l'import : une entrée par ligne lue, et son empreinte."""

    lignes: list[LigneLue]
    empreinte: str

    def compte(self, etat: str) -> int:
        """Le nombre de lignes dans un état (nouvelle, modifiee, inchangee, erreur)."""
        return sum(1 for x in self.lignes if x.etat == etat)

    @property
    def appliquable(self) -> bool:
        """Aucune erreur, et au moins une matière à écrire."""
        return not self.compte("erreur") and any(x.ecrire for x in self.lignes)


def _feuille(contenu: bytes) -> Iterator[tuple[Any, ...]]:
    """Les lignes de la feuille « Matières » (à défaut, la feuille active)."""
    try:
        classeur = load_workbook(BytesIO(contenu), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — un fichier illisible se dit en clair
        raise Refus.saisie(
            "Tableur illisible : le fichier n'est pas un classeur Excel (.xlsx) lisible.",
            code="tableur_invalide",
        ) from exc
    feuille = classeur[FEUILLE] if FEUILLE in classeur.sheetnames else classeur.active
    if feuille is None:
        raise Refus.saisie("Tableur illisible : aucune feuille.", code="tableur_invalide")
    return feuille.iter_rows(values_only=True)  # type: ignore[no-any-return]


def _entetes(premiere: tuple[Any, ...]) -> dict[str, int]:
    """Colonne → position, par intitulé (les anciens intitulés restent lus)."""
    noms = [ANCIENS.get(str(x or "").strip(), str(x or "").strip()) for x in premiere]
    index = {nom: i for i, nom in enumerate(noms) if nom}
    if "Référence" not in index:
        raise Refus.saisie(
            "La colonne « Référence » est obligatoire (télécharger le modèle).",
            code="tableur_invalide",
        )
    return index


def _fr(x: float) -> str:
    """« 10 000 000 », « 2,5 » : un nombre à la française."""
    if x == int(x):
        return f"{int(x):,}".replace(",", " ")
    return f"{x:g}".replace(".", ",")


def _vide(v: object) -> bool:
    """Une case vide (rien, ou des espaces)."""
    return v is None or (isinstance(v, str) and not v.strip())


def _nombre(libelle: str, champ: str, v: object, erreurs: list[str]) -> Decimal | None:
    """Un nombre fini dans ses bornes (virgule acceptée)."""
    try:
        x = float(str(v).strip().replace(" ", "").replace(",", "."))
    except ValueError:
        x = math.nan
    if not math.isfinite(x):
        erreurs.append(f"{libelle} : pas un nombre (« {v} »)")
        return None
    bas, haut = BORNES[champ]
    if not bas <= x <= haut:
        erreurs.append(f"{libelle} : entre {_fr(bas)} et {_fr(haut)} — reçu {_fr(x)}")
        return None
    return Decimal(str(x))


def _cellule(r: Referentiels, colonne: Any, v: object, erreurs: list[str]) -> object:
    """La valeur d'une case non vide, traduite (codes → ids, nombres, oui/non)."""
    libelle, champ, genre = colonne
    texte = str(v).strip()
    if genre == "nombre":
        return _nombre(libelle, champ, v, erreurs)
    if genre == "actif":
        if texte.lower() in ACTIF:
            return ACTIF[texte.lower()]
        erreurs.append(f"{libelle} : oui ou non (« {texte} »)")
        return None
    if genre == "unite_prix":
        return texte.lower()
    if genre == "texte":
        if len(texte) > LONGUEURS.get(champ, L.DESIGNATION):
            erreurs.append(f"{libelle} : {LONGUEURS.get(champ, L.DESIGNATION)} caractères au plus")
        return texte
    trouve = r.chercher(genre, texte)
    if trouve is None:
        erreurs.append(r.inconnu(genre, texte))
    return trouve


def _lire_ligne(
    r: Referentiels, index: dict[str, int], n: int, valeurs: tuple[Any, ...]
) -> LigneLue:
    """Une ligne : ses cases traduites, ses colonnes à vider, sa version."""

    def case(nom: str) -> object:
        """La case d'une colonne (ou rien)."""
        i = index.get(nom)
        return valeurs[i] if i is not None and i < len(valeurs) else None

    lue = LigneLue(n, str(case("Référence") or "").strip())
    donnees: dict[str, Any] = {}
    for colonne in COLONNES[1:]:
        v = case(colonne[0])
        if not _vide(v):
            donnees[colonne[1]] = _cellule(r, colonne, v, lue.erreurs)
    for nom in [x.strip() for x in str(case(VIDER) or "").split(";") if x.strip()]:
        champ = LIBELLES.get(nom)
        if champ is None:
            lue.erreurs.append(f"Vider : colonne inconnue « {nom} »")
        elif champ in NE_SE_VIDE_PAS:
            lue.erreurs.append(f"Vider : « {nom} » ne se vide pas")
        else:
            donnees[champ] = VIDE[champ]
    version = case(VERSION)
    if not _vide(version):
        try:
            lue.version = int(str(version).strip().split(".")[0])
        except ValueError:
            lue.erreurs.append(f"Version : pas un nombre (« {version} »)")
    lue.ecrire = donnees
    return lue


NOUVELLE: dict[str, Any] = {
    "designation": "",
    "role_id": None,
    "composition_id": None,
    "type_fil_id": None,
    "titrage_id": None,
    "couleur_id": None,
    "couleur_code": "",
    "unite_prix": "kg",
    "densite": Decimal(0),
    "fournisseur_id": None,
    "prix": Decimal(0),
    "devise": "MAD",
    "unite_achat": "kg",
    "contenance_achat": Decimal(1),
    "minimum_commande": Decimal(0),
    "poids_bobine": Decimal(0),
    "bobines_palette": Decimal(0),
    "stock_min": Decimal(0),
    "couv_min_mois": Decimal(0),
    "groupe_equiv": "",
    "priorite_equiv": None,
    "code_douanier_id": None,
    "actif": True,
}
"""Les valeurs d'une matière nouvelle avant les cases du tableur."""


def _preparer(
    s: Session, r: Referentiels, lue: LigneLue, a: Matiere | None, noms: dict[int, str]
) -> None:
    """La fiche telle qu'elle serait après l'import, passée par les règles de la fiche."""
    if a is None and "role_id" not in lue.ecrire:
        lue.erreurs.append("Une matière nouvelle a besoin de son rôle")
    if a is not None and lue.version is not None and lue.version != a.version:
        par = noms.get(a.modifie_par_id or 0, "un autre poste")
        lue.erreurs.append(
            f"Modifiée entre-temps par {par} depuis le modèle : retélécharger le modèle"
        )
    if lue.erreurs:
        return
    avant = {k: getattr(a, k) for k in NOUVELLE} if a is not None else dict(NOUVELLE)
    fiche = {**avant, **lue.ecrire, "reference": a.reference if a is not None else lue.reference}
    try:
        propre = verifier(s, fiche, a)
        if a is None:
            reference_libre(s, propre["reference"], None)
    except Refus as refus:
        lue.erreurs.append(refus.detail)
        return
    propre.pop("reference")
    change = {k: v for k, v in propre.items() if a is None or avant[k] != v}
    lue.matiere_id, lue.version = (a.id, a.version) if a is not None else (None, None)
    lue.ecrire = change
    lue.etat = "nouvelle" if a is None else ("modifiee" if change else "inchangee")
    if a is not None:
        lue.changements = [
            {
                "champ": k,
                "libelle": INTITULES.get(k, k),
                "avant": r.nom(genre(k), avant[k]),
                "apres": r.nom(genre(k), v),
            }
            for k, v in sorted(change.items(), key=lambda x: list(NOUVELLE).index(x[0]))
        ]


def apercu(s: Session, contenu: bytes) -> Apercu:
    """Lit le tableur et dit ce qui serait créé, modifié, ou ce qui ne va pas (route 79)."""
    rangees = _feuille(contenu)
    index = _entetes(next(rangees, ()))
    r = referentiels(s)
    existantes = {m.reference.lower(): m for m in s.scalars(select(Matiere))}
    noms = dict(s.execute(select(Utilisateur.id, Utilisateur.nom)).all())
    vues: dict[str, int] = {}
    lignes: list[LigneLue] = []
    for n, valeurs in enumerate(rangees, start=2):
        if n - 1 > LIGNES_MAX:
            raise Refus.saisie(
                f"Tableur trop long : {LIGNES_MAX} lignes au plus.", code="tableur_invalide"
            )
        lue = _lire_ligne(r, index, n, valeurs)
        if not lue.reference:
            continue
        cle = lue.reference.lower()
        if cle in vues:
            lue.erreurs.append(
                f"La référence « {lue.reference} » figure déjà ligne {vues[cle]} : une seule "
                "ligne par matière"
            )
        else:
            vues[cle] = n
            if len(lue.reference) > L.REFERENCE:
                lue.erreurs.append(f"Référence : {L.REFERENCE} caractères au plus")
            _preparer(s, r, lue, existantes.get(cle), noms)
        if lue.erreurs:
            lue.etat, lue.ecrire = "erreur", {}
        lignes.append(lue)
    return Apercu(lignes, _empreinte(lignes))


def _empreinte(lignes: list[LigneLue]) -> str:
    """L'empreinte du plan d'écriture : ce que l'application écrira, exactement."""
    plan = [
        [x.reference, x.matiere_id, x.version, {k: str(v) for k, v in sorted(x.ecrire.items())}]
        for x in lignes
        if x.ecrire or x.erreurs
    ]
    texte = json.dumps(plan, ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha256(texte.encode()).hexdigest()[:32]


def appliquer(s: Session, qui: audit.Qui, contenu: bytes, empreinte: str) -> Apercu:
    """Écrit le tableur, tout ou rien, s'il dit encore ce que l'aperçu a montré (route 80)."""
    resultat = apercu(s, contenu)
    erreurs = resultat.compte("erreur")
    if erreurs:
        raise Refus.saisie(
            f"{erreurs} ligne(s) en erreur : corriger le tableur d'abord, rien n'a été écrit.",
            code="tableur_invalide",
        )
    if empreinte != resultat.empreinte:
        raise Refus.etat(
            "Des matières du tableur ont été modifiées sur un autre poste depuis l'aperçu : "
            "relancer l'aperçu, puis appliquer.",
            code="conflit_version",
        )
    for lue in resultat.lignes:
        if not lue.ecrire:
            continue
        if lue.matiere_id is None:
            m = Matiere(reference=lue.reference, **lue.ecrire, modifie_par_id=qui.id)
            s.add(m)
            audit.creation(s, qui, "import des matières", m, detail=m.reference)
            continue
        m = s.get_one(Matiere, lue.matiere_id, with_for_update=True)
        with audit.suivre(s, qui, "import des matières", m, detail=m.reference):
            for k, v in lue.ecrire.items():
                setattr(m, k, v)
            m.modifie_par_id = qui.id
    audit.journal(
        s,
        qui,
        "import des matières",
        ObjetJournal.matiere,
        detail=f"{resultat.compte('nouvelle')} créée(s), {resultat.compte('modifiee')} modifiée(s)",
    )
    s.flush()
    return resultat
