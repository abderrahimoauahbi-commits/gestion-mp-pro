"""Qualités, recettes et versions (routes 82–89, 92, 93 ; lot D2).

Héritage REGLES A § 1.7–1.8 (R-REF-55 à R-REF-69, X-REF-39 à X-REF-52) ; défauts de Gestion MP
D-01, D-03, D-04, D-05, D-18, D-19 ; ECRANS-PRO partie 7 (P7-5, P7-8).
"""

from __future__ import annotations

import io
import json

import pytest
from fabrique import *  # noqa: F403
from fabriques.qualites import (
    corps_fiche,
    creer_qualite,
    lignes,
    lire_qualite,
    ok,
    ordre_valide,
    poids,
    refus,
)
from fabriques.referentiels import matiere, role_id, unique
from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _versions(client: TestClient, qid: int) -> list[dict]:
    return ok(client.get(f"/api/qualites/{qid}/versions"))  # type: ignore[no-any-return]


def test_qualite_empreinte_doublon_suppression(direction: TestClient, achats: TestClient) -> None:
    """Hérité : test_referentiels_complets::test_qualite_empreinte_doublon_suppression."""
    poil, trame = matiere(prix="2.5"), matiere(role="Trame", prix="4")
    q = creer_qualite(direction, poids_m2={"Poil": 1.2, "Trame": 0.4})
    autre = creer_qualite(direction)
    lue = q["version"]
    r = ok(
        direction.put(
            f"/api/qualites/{q['id']}/recette",
            json={"lignes": lignes({poil: 100, trame: 100}), "version": lue},
        )
    )
    assert r["version"] > lue
    # Une modification partie de l'ancienne lecture est refusée, en nommant l'auteur.
    perime = direction.put(f"/api/qualites/{q['id']}", json=corps_fiche(q, version=lue))
    assert refus(perime, 409, "Direction")["code"] == "conflit_version"
    doublon = direction.put(f"/api/qualites/{q['id']}", json=corps_fiche(r, code=autre["code"]))
    assert refus(doublon, 409, "existe déjà")["code"] == "doublon"
    liste = {x["code"]: x for x in ok(achats.get("/api/qualites"))}
    lu = liste[q["code"]]
    assert lu["kg_m2"] == pytest.approx(1.6) and lu["fabrications"] == 0
    f = lu["facteurs_cout"]
    attendu = sum(
        x["kg_m2"] * x["prix_kg"] * f["cours"][x["devise_prix"]] * f["perte"] for x in lu["recette"]
    )
    assert (
        lu["cout_m2_mad"]
        == pytest.approx(attendu)
        == pytest.approx(1.2 * 2.5 * 1.02 + 0.4 * 4 * 1.02)
    )
    assert achats.delete(f"/api/qualites/{q['id']}?version={r['version']}").status_code == 403
    assert direction.delete(f"/api/qualites/{q['id']}?version={r['version']}").status_code == 204
    assert all(x["code"] != q["code"] for x in ok(achats.get("/api/qualites")))


def test_une_qualite_creee_a_ses_premieres_versions(direction: TestClient) -> None:
    """D-18 : la première version de recette (ses poids, sans ligne) et le plan « d'avant »."""
    q = creer_qualite(direction, poids_m2={"Poil": "0.7", "Trame": 0})
    assert [p["role"] for p in q["poids"]] == ["Poil"]  # un poids nul n'est pas enregistré
    v = _versions(direction, q["id"])
    assert len(v) == 1 and v[0]["lignes"] == [] and v[0]["numero"] == 1
    assert v[0]["poids"] == [{"role_id": role_id("Poil"), "role": "Poil", "poids": 0.7}]
    assert v[0]["par"]["nom"] == "Direction" and v[0]["en_vigueur"]
    plans = ok(direction.get(f"/api/qualites/{q['id']}/plan/versions"))
    assert len(plans) == 1 and plans[0]["depuis_toujours"] and plans[0]["base_mensuelle"] == 0
    assert plans[0]["saisonnalite"] == [1.0] * 12 and plans[0]["en_service"]


def test_d01_poids_et_plan_refusent_nan(direction: TestClient) -> None:
    """D-01 (Hérité : test_defauts_d::test_d01_plan_et_poids_refusent_nan) : NaN et infini
    refusés en 422 ; ensuite le stock et les besoins se calculent."""
    q = creer_qualite(direction, poids_m2={"Poil": 1})
    entete = {"Content-Type": "application/json"}
    for valeur in ("NaN", "Infinity"):
        brut = json.dumps(corps_fiche(q)).replace('"poids": 1.0', f'"poids": {valeur}')
        brut = brut.replace('"poids": 1', f'"poids": {valeur}')
        assert (
            direction.put(f"/api/qualites/{q['id']}", content=brut, headers=entete).status_code
            == 422
        )
        plan = '{"base_mensuelle": 100, "coefficients": [%s, 1,1,1,1,1,1,1,1,1,1,1], "version": %d}'
        r = direction.put(
            f"/api/qualites/{q['id']}/plan", content=plan % (valeur, q["version"]), headers=entete
        )
        assert r.status_code == 422
    assert direction.get("/api/stock").status_code == 200
    assert direction.get("/api/production/besoins").status_code == 200


def test_une_qualite_fabriquee_ne_se_supprime_pas(direction: TestClient) -> None:
    """X-REF-44 : une qualité qu'un ordre a tissée appartient à l'historique."""
    q = creer_qualite(direction)
    ordre_valide(q["id"], 10)
    q = lire_qualite(direction, q["id"])
    assert q["fabrications"] == 1
    r = direction.delete(f"/api/qualites/{q['id']}?version={q['version']}")
    assert refus(r, 409, "a été fabriquée par 1 ordre(s)")["code"] == "fiche_utilisee"


def test_poids_refuses(direction: TestClient) -> None:
    """X-REF-42, X-REF-43 ; un rôle cité deux fois."""
    url = "/api/qualites"
    refus(
        direction.post(url, json={"code": unique("Q"), "poids": [{"role_id": 999999, "poids": 1}]}),
        422,
        "Rôle inconnu : « 999999 ».",
    )
    refus(
        direction.post(url, json={"code": unique("Q"), "poids": poids(Poil=-1)}),
        422,
        "Un poids au m² ne peut pas être négatif.",
    )
    deux = poids(Poil=1) * 2
    refus(direction.post(url, json={"code": unique("Q"), "poids": deux}), 422, "deux fois")
    refus(direction.get("/api/qualites/999999"), 404, "Qualité introuvable.")


def test_l_alimentation_lit_les_qualites_sans_prix(
    direction: TestClient, alimentation: TestClient
) -> None:
    """R-REF-56 : ni prix, ni coût, ni unité du prix, ni cours pour l'alimentation."""
    from test_prix_masques import fuites

    a = matiere(prix="2.5", devise="MAD")
    q = creer_qualite(direction, poids_m2={"Poil": 1}, parts={a: 100})
    for chemin in ("/api/qualites", f"/api/qualites/{q['id']}", "/api/nomenclature"):
        corps = ok(alimentation.get(chemin))
        assert fuites(corps) == [], chemin
    ligne = lire_qualite(alimentation, q["id"])["recette"][0]
    assert ligne["prix"] is None and ligne["unite_prix"] is None and ligne["cout_m2_mad"] is None
    assert ligne["kg_m2"] == pytest.approx(1)
    assert alimentation.get("/api/export/qualites").status_code == 403


def test_d19_export_des_qualites(direction: TestClient, achats: TestClient) -> None:
    """D-19 : « En service » et « Fabriquée » (vraie) en deux colonnes ; la partie qualités de
    Hérité : test_referentiels_complets::test_les_exports_des_referentiels ; et
    Hérité : test_defauts_d::test_d19_export_des_qualites
    """
    q = creer_qualite(direction, poids_m2={"Poil": 1}, parts={matiere(): 100})
    ordre_valide(q["id"], 5)
    r = achats.get("/api/export/qualites")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    feuille = load_workbook(io.BytesIO(r.content)).active
    assert feuille is not None
    lignes_lues = list(feuille.iter_rows(values_only=True))
    entete = lignes_lues[0]
    assert "En service" in entete and "Fabriquée" in entete
    ligne = next(x for x in lignes_lues if x[0] == q["code"])
    assert ligne[entete.index("Fabriquée")] == "oui" and ligne[entete.index("En service")] == "oui"


def test_journal_avant_apres_d_une_qualite(direction: TestClient) -> None:
    """Chaque geste écrit au journal ce qui a changé (avant / après)."""
    q = creer_qualite(direction, poids_m2={"Poil": 1})
    ok(direction.put(f"/api/qualites/{q['id']}", json=corps_fiche(q, nom="Autre nom")))
    histo = ok(direction.get(f"/api/historique/qualite/{q['id']}"))
    lignes_histo = histo if isinstance(histo, list) else histo["elements"]
    derniere = next(x for x in reversed(lignes_histo) if x["action"] == "modification qualité")
    assert derniere["avant"] == {"nom": "Essai D2"} and derniere["apres"] == {"nom": "Autre nom"}
    # Renvoyer la fiche telle quelle n'écrit rien : ni journal, ni version.
    lue = lire_qualite(direction, q["id"])
    assert (
        ok(direction.put(f"/api/qualites/{q['id']}", json=corps_fiche(lue)))["version"]
        == (lue["version"])
    )
    assert len(ok(direction.get(f"/api/historique/qualite/{q['id']}"))) == len(lignes_histo)
