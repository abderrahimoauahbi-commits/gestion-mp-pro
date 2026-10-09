"""Routes du stock (lot S4) : état du stock (103), fiche matière (104), journal (105) et son
classeur (106), classeur des lots (236) ; prix, montants et conseil d'achat jamais montrés à
l'alimentation. Les données sont validées en base, à des noms uniques.
"""

from __future__ import annotations

from collections.abc import Iterator
from io import BytesIO

import pytest
from fabrique import *  # noqa: F403
from fabriques.stock import JOUR, entree, jour, lot, matiere, sortie, unique
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from gestion_mp_pro.db import fabrique
from gestion_mp_pro.modeles import Lot


@pytest.fixture
def base() -> Iterator[Session]:
    """Une session dont les écritures sont validées (les routes ont leurs propres sessions)."""
    with fabrique()() as s:
        yield s


def _journal_valorise(s: Session) -> tuple[int, str]:
    """500 kg à 20 MAD (J−2), 100 kg à 30 MAD (hier), 100 kg sortis aujourd'hui."""
    ref = unique("Fil Frisé à valoriser")
    m = matiere(s, ref, prix=25)
    x = lot(s, m, entre_le=jour(-2))
    entree(s, x, 500, jour(-2), prix=20)
    entree(s, x, 100, jour(-1), prix=30)
    sortie(s, x, 100)
    s.commit()
    return m.id, ref


def test_etat_du_stock_et_masques_de_l_alimentation(
    base: Session, direction: TestClient, alimentation: TestClient
) -> None:
    """Route 103 : une ligne par matière, à la date ; l'alimentation n'a ni prix, ni montant,
    ni conseil d'achat (E-01), mais bien le stock."""
    matiere_id, ref = _journal_valorise(base)
    r = direction.get("/api/stock", params={"q": ref.replace("é", "e").replace("à", "a")})
    assert r.status_code == 200, r.text
    (ligne,) = r.json()["lignes"]
    assert (ligne["reference"], ligne["matiere_id"], ligne["stock_net"]) == (ref, matiere_id, 500)
    assert ligne["prix_unite"] == 25
    hier = direction.get("/api/stock", params={"q": ref, "date": jour(-1).isoformat()}).json()
    assert hier["date_reference"] == jour(-1).isoformat()
    assert hier["lignes"][0]["stock_net"] == 600
    vu = alimentation.get("/api/stock", params={"q": ref}).json()["lignes"][0]
    assert vu["stock_net"] == 500
    masques = ("prix_unite", "prix_kg", "montant_a_commander_mad", "valeur_machine_mad",
               "quantite_a_commander", "classe_abc", "action")  # fmt: skip
    assert {k: vu[k] for k in masques} == dict.fromkeys(masques)
    assert direction.get("/api/stock", params={"date": "2999-01-01"}).status_code == 422


def test_fiche_de_stock_d_une_matiere(base: Session, direction: TestClient) -> None:
    """Route 104 : sa ligne, ses lots (à la date) ; 404 pour une matière inconnue."""
    matiere_id, ref = _journal_valorise(base)
    r = direction.get(f"/api/stock/{matiere_id}")
    assert r.status_code == 200, r.text
    f = r.json()
    assert f["ligne"]["reference"] == ref
    assert [(x["quantite"], x["statut"]) for x in f["lots"]] == [(500, "disponible")]
    assert f["besoins_par_qualite"] == [] and f["sur_les_metiers"] == []
    assert direction.get("/api/stock/999999").status_code == 404


def test_journal_pagine_et_valorise(
    base: Session, direction: TestClient, alimentation: TestClient
) -> None:
    """Route 105. Hérité : test_stock_atelier::test_journal_pagine_et_valorise — page de 2 et
    son curseur ; 500 × 20 = 10 000 ; la sortie au coût moyen du moment (20 < c < 22)."""
    matiere_id, _ = _journal_valorise(base)
    params = {"matiere": matiere_id, "limite": 2, "avec_valeur": "true"}
    page = direction.get("/api/mouvements", params=params).json()
    assert page["total"] == 3 and page["suivant"]
    assert [x["quantite"] for x in page["elements"]] == [-100, 100]
    sortie_ = page["elements"][0]
    assert sortie_["libelle"] == "Sortie vers les machines"
    assert 20 < sortie_["cout_mad"] < 22
    assert sortie_["valeur_mad"] == pytest.approx(-100 * sortie_["cout_mad"])
    suite = direction.get("/api/mouvements", params={**params, "apres": page["suivant"]}).json()
    (ouverture,) = suite["elements"]
    assert (ouverture["valeur_mad"], suite["suivant"]) == (10_000, None)
    vu = alimentation.get("/api/mouvements", params=params).json()["elements"][0]
    assert (vu["prix"], vu["cout_mad"], vu["valeur_mad"], vu["cours"]) == (None, None, None, None)


def test_recherche_du_journal_sans_accents_et_jokers_litteraux(
    base: Session, direction: TestClient
) -> None:
    """C-14 : « frise valoriser » trouve « Frisé à valoriser » ; « % » ne ramène pas tout."""
    _, ref = _journal_valorise(base)
    numero = ref.rsplit("-", 1)[1]
    trouve = direction.get("/api/mouvements", params={"q": f"FRISE valoriser {numero}"}).json()
    assert trouve["total"] == 3 and {x["reference"] for x in trouve["elements"]} == {ref}
    assert direction.get("/api/mouvements", params={"q": "%"}).json()["total"] == 0
    illisible = direction.get("/api/mouvements", params={"apres": "pas-un-curseur"})
    assert (illisible.status_code, illisible.json()["code"]) == (422, "curseur_invalide")


def _feuille(contenu: bytes) -> list[tuple[object, ...]]:
    ws = load_workbook(BytesIO(contenu)).active
    assert ws is not None
    return list(ws.iter_rows(values_only=True))


def test_classeurs_des_mouvements_et_des_lots(
    base: Session, direction: TestClient, alimentation: TestClient
) -> None:
    """Routes 106 et 236 : mêmes filtres que l'écran, vraies dates ; aucun prix ni valeur dans
    le classeur de l'alimentation (barrière 4)."""
    matiere_id, ref = _journal_valorise(base)
    r = direction.get("/api/export/mouvements", params={"matiere": matiere_id})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    entete, *lignes = _feuille(r.content)
    assert "Valeur (MAD)" in entete and len(lignes) == 3
    assert lignes[-1][0].date() == jour(-2)  # type: ignore[union-attr]
    sans_prix = _feuille(
        alimentation.get("/api/export/mouvements", params={"matiere": matiere_id}).content
    )
    assert not {"Prix / kg", "Valeur (MAD)", "Cours", "Devise"} & set(sans_prix[0])
    lots_ = _feuille(direction.get("/api/export/lots", params={"matiere": matiere_id}).content)
    assert lots_[0][:3] == ("Lot", "Statut", "Référence")
    assert [(x[2], x[10]) for x in lots_[1:]] == [(ref, 500)]


def test_ecritures_fermees_et_texte_dangereux(base: Session, direction: TestClient) -> None:
    """Une cellule qui commence par « = » reste du texte (jamais une formule)."""
    m = matiere(base, unique("=SOMME(A1)"))
    x = lot(base, m)
    entree(base, x, 1)
    base.commit()
    contenu = direction.get("/api/export/mouvements", params={"matiere": m.id}).content
    ws = load_workbook(BytesIO(contenu)).active
    assert ws is not None
    cellule = ws.cell(row=2, column=5)
    assert cellule.data_type == "s" and str(cellule.value).startswith("=SOMME")
    assert isinstance(base.get(Lot, x.id), Lot)
    assert JOUR == jour(0)
