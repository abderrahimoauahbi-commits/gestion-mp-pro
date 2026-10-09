"""Fournisseurs (routes 64–69) — lot D1.

Héritage R-REF-16…20, X-REF-14…19 (REGLES § 1.5) ; défauts A-14, A-23, A-25 ; essais
`test_referentiels_complets` (port, banque, empreinte devenue `version` ; délai mesuré ; pièces
jointes) et `test_referentiels::test_un_pays_s_ecrit_d_une_seule_facon`.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any

from fabrique import *  # noqa: F403
from fabriques.referentiels import devise, fournisseur, matiere, unique
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import update

from gestion_mp_pro.db import fabrique as sessions
from gestion_mp_pro.modeles import (
    Commande,
    CommandeLigne,
    PieceJointe,
    Reception,
    ReceptionLigne,
)
from gestion_mp_pro.modeles.enums import ObjetPiece, StatutDocument
from gestion_mp_pro.services import horloge, lots
from gestion_mp_pro.services.fournisseurs import est_local


def ok(r: Any, statut: int = 200) -> Any:
    """La réponse JSON d'un appel qui doit réussir."""
    assert r.status_code == statut, r.text
    return r.json()


def fiche(client: TestClient, ident: int) -> Any:
    """La fiche d'un fournisseur."""
    return ok(client.get(f"/api/fournisseurs/{ident}"))


def test_fournisseur_port_banque_version(achats: TestClient) -> None:
    """Hérité : test_referentiels_complets::test_fournisseur_port_banque_empreinte — champs port
    et banque, compteurs par devise, verrou optimiste (`version` remplace l'empreinte),
    suppression refusée tant qu'il a des matières."""
    usd = devise(cours="10")
    nom = unique("Filature")
    f = ok(achats.post("/api/fournisseurs", json={
        "nom": f" {nom} ", "devise": usd.lower(), "port_chargement": "Mersin ",
        "banque": "Ziraat", "delai_jours": 45,
    }), 201)  # fmt: skip
    assert (f["nom"], f["devise"], f["port_chargement"], f["banque"]) == (
        nom,
        usd,
        "Mersin",
        "Ziraat",
    )
    matiere(fournisseur_id=f["id"], devise=usd)
    matiere(fournisseur_id=f["id"])
    matiere(fournisseur_id=f["id"])
    lu = fiche(achats, f["id"])
    assert (lu["nb_matieres"], lu["devises"]) == (3, {usd: 1, "MAD": 2})
    corps = {
        **{k: lu[k] for k in ("nom", "devise", "delai_jours")},
        "banque": "BMCE",
        "version": lu["version"],
    }
    apres = ok(achats.put(f"/api/fournisseurs/{f['id']}", json=corps))
    assert apres["banque"] == "BMCE" and apres["version"] == lu["version"] + 1
    perime = achats.put(f"/api/fournisseurs/{f['id']}", json={**corps, "banque": "CIH"})
    assert perime.status_code == 409
    assert perime.json()["detail"].startswith(
        "Ce fournisseur a été modifié entre-temps par Achats ("
    )
    r = achats.delete(f"/api/fournisseurs/{f['id']}")
    assert (r.status_code, r.json()["detail"]) == (
        409,
        f"« {nom} » a des matières du catalogue : le désactiver plutôt que le supprimer.",
    )


def test_un_pays_s_ecrit_d_une_seule_facon(achats: TestClient) -> None:
    """Hérité : test_referentiels::test_un_pays_s_ecrit_d_une_seule_facon (R-REF-16)."""
    f = ok(achats.post("/api/fournisseurs", json={"nom": unique("T"), "pays": " turkey  "}), 201)
    assert f["pays"] == "TURKEY"
    g = ok(
        achats.post("/api/fournisseurs", json={"nom": unique("T"), "pays": "côte  d'ivoire"}), 201
    )
    assert g["pays"] == "CÔTE D'IVOIRE"


def test_refus_de_saisie(achats: TestClient) -> None:
    """Hérité : X-REF-14, X-REF-15 (casse ignorée), X-REF-16."""
    nom = unique("Doublon")
    ok(achats.post("/api/fournisseurs", json={"nom": nom}), 201)
    r = achats.post("/api/fournisseurs", json={"nom": nom.upper()})
    assert (r.status_code, r.json()["detail"]) == (
        409,
        f"Le fournisseur « {nom.upper()} » existe déjà.",
    )
    r = achats.post("/api/fournisseurs", json={"nom": unique("D"), "devise": "XYZ"})
    assert (
        r.json()["detail"] == "Devise inconnue : « XYZ ». L'ajouter d'abord (Réglages › Devises)."
    )
    assert achats.post("/api/fournisseurs", json={"nom": "x", "delai_jours": 0}).status_code == 422
    assert achats.get("/api/fournisseurs/999999").json()["detail"] == "Fournisseur introuvable."


def test_l_alimentation_ne_voit_pas_la_banque_du_fournisseur(alimentation: TestClient) -> None:
    """Hérité : test_defauts_a::test_l_alimentation_ne_voit_pas_la_banque_du_fournisseur (A-14)."""
    f = fournisseur(banque="Attijari", conditions_paiement="60 j", notes="remise 2 %")
    lus = ok(alimentation.get("/api/fournisseurs"))
    vu = next(x for x in lus if x["id"] == f)
    assert set(vu) == {"id", "nom", "pays", "actif"}
    assert alimentation.get(f"/api/fournisseurs/{f}").status_code == 403


def test_un_fournisseur_sans_pays_en_usd_n_est_local_nulle_part(achats: TestClient) -> None:
    """Hérité : test_defauts_a::test_un_fournisseur_sans_pays_en_usd_n_est_local_nulle_part
    (A-23) — une seule définition, que les importations lisent aussi."""
    assert not est_local("", "USD") and est_local("", "MAD") and est_local("ma", "USD")
    assert est_local("Morocco", "EUR") and not est_local("TURKEY", "MAD")
    f = fournisseur(devise=devise())
    assert fiche(achats, f)["local"] is False


def test_le_fournisseur_dit_tout_ce_qui_empeche_sa_suppression(achats: TestClient) -> None:
    """Hérité :
    test_defauts_a::test_le_fournisseur_dit_tout_ce_qui_empeche_sa_suppression (A-25)."""
    f = fournisseur()
    m = matiere()
    with sessions()() as s, s.begin():
        lots.lot_ou_creer(s, m, unique("L"), horloge.aujourd_hui(), fournisseur_id=f)
    lu = fiche(achats, f)
    assert (lu["nb_lots"], lu["empechements"], lu["utilise"]) == (1, ["des lots"], True)
    r = achats.delete(f"/api/fournisseurs/{f}")
    assert r.json()["detail"].endswith("a des lots : le désactiver plutôt que le supprimer.")
    libre = fournisseur()
    assert fiche(achats, libre)["empechements"] == []
    assert achats.delete(f"/api/fournisseurs/{libre}").status_code == 204


def test_une_piece_jointe_garde_sa_fiche(achats: TestClient, comptes: dict[str, int]) -> None:
    """Hérité : test_referentiels_complets::test_une_piece_jointe_garde_sa_fiche — la pièce
    retient sa matière et son fournisseur ; retirée, elle garde l'historique."""
    f = fournisseur()
    m = matiere(fournisseur_id=f, actif=False)
    with sessions()() as s, s.begin():
        for objet, champ, ident in (
            (ObjetPiece.matiere, "matiere_id", m),
            (ObjetPiece.fournisseur, "fournisseur_id", f),
        ):
            piece = PieceJointe(objet=objet, nom="contrat.pdf", type_mime="application/pdf")
            piece.taille, piece.empreinte, piece.chemin = 12, "0" * 64, f"essai/{ident}.pdf"
            setattr(piece, champ, ident)
            s.add(piece)
    r = achats.delete(f"/api/matieres/{m}")
    assert (
        r.status_code == 409
        and "figure dans des pièces jointes : l'en retirer d'abord" in r.json()["detail"]
    )
    assert fiche(achats, f)["nb_pieces"] == 1
    r = achats.delete(f"/api/fournisseurs/{f}")
    assert "a des matières du catalogue" in r.json()["detail"]
    with sessions()() as s, s.begin():
        s.execute(update(PieceJointe).where(PieceJointe.fournisseur_id == f).values(
            retiree_le=horloge.maintenant(), retiree_par_id=comptes["achats"]))  # fmt: skip


def test_delai_mesure(achats: TestClient) -> None:
    """Hérité : test_referentiels_complets::test_delai_mesure (essai faible, complété) — sans
    commande rien n'est mesuré ; commande du 1er septembre reçue le 11 : 10 jours."""
    f = fournisseur()
    assert (fiche(achats, f)["delai_mesure_jours"], fiche(achats, f)["livraisons_mesurees"]) == (
        None,
        0,
    )
    m = matiere(fournisseur_id=f)
    with sessions()() as s, s.begin():
        c = Commande(fournisseur_id=f, date_commande=date(2026, 9, 1), devise="MAD")
        s.add(c)
        s.flush()
        ligne = CommandeLigne(commande_id=c.id, matiere_id=m, quantite=Decimal(100))
        r = Reception(date=date(2026, 9, 11), fournisseur_id=f, devise="MAD", numero=unique("REC"))
        r.statut, r.validee_le = StatutDocument.validee, horloge.maintenant()
        s.add_all([ligne, r])
        s.flush()
        s.add(ReceptionLigne(reception_id=r.id, commande_ligne_id=ligne.id, matiere_id=m))
    lu = fiche(achats, f)
    assert (lu["delai_mesure_jours"], lu["livraisons_mesurees"], lu["nb_receptions"]) == (
        10.0,
        1,
        1,
    )


def test_recherche_et_export(achats: TestClient, alimentation: TestClient) -> None:
    """Hérité : test_referentiels_complets::test_les_exports_des_referentiels (fiches) ;
    recherche « tous les mots, sans accents »."""
    nom = f"Société Générale Fils {unique('G')}"
    f = fournisseur(nom, pays="FRANCE")
    trouves = ok(
        achats.get("/api/fournisseurs", params={"q": f"generale {nom.rsplit(maxsplit=1)[-1]}"})
    )
    assert [x["id"] for x in trouves] == [f]
    assert ok(achats.get("/api/fournisseurs", params={"actif": False, "q": nom})) == []
    r = achats.get("/api/export/fournisseurs/fiches")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    feuille = load_workbook(BytesIO(r.content)).active
    assert feuille["A1"].value.startswith("Fournisseurs — exporté le 08/10/2026")
    assert feuille["A3"].value == "Fournisseur"
    assert any(row[0] == nom for row in feuille.iter_rows(min_row=4, values_only=True))
    assert alimentation.get("/api/export/fournisseurs/fiches").status_code == 403
