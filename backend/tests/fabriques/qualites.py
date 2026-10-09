"""Fabriques du lot D2 : qualités et métiers par l'API, ordres validés et sorties en décor.

Les référentiels viennent de `fabriques.referentiels` (lot D1). Les ordres de fabrication
(lot D6) et les sorties vers les machines sont écrits **directement en base** (décor d'essai,
validé : l'API les voit) ; le stock passe par `services/stock` (S4), seul chemin.

    from fabriques.qualites import ok, creer_qualite, creer_metier, ordre_valide, sortir
"""

from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

from fabriques.referentiels import _ecrire, magasin_seme, role_id, unique
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from gestion_mp_pro.modeles import Lot, Ordre, OrdreQualite
from gestion_mp_pro.modeles.enums import DocumentMouvement, StatutOrdre, TypeMouvement
from gestion_mp_pro.services import horloge, stock

UN = [1] * 12


def classeur_lu(contenu: bytes) -> tuple[tuple[Any, ...], list[tuple[Any, ...]]]:
    """L'en-tête (la première ligne d'au moins deux cases) et les lignes d'un classeur."""
    import io

    from openpyxl import load_workbook

    feuille = load_workbook(io.BytesIO(contenu)).active
    assert feuille is not None
    lignes = list(feuille.iter_rows(values_only=True))
    rang = next(i for i, x in enumerate(lignes) if sum(v is not None for v in x) >= 2)
    return lignes[rang], lignes[rang + 1 :]


def ok(r: Any, statut: int = 200) -> Any:
    """La réponse JSON d'un appel qui doit réussir."""
    assert r.status_code == statut, r.text
    return r.json()


def refus(r: Any, statut: int, mot: str = "") -> dict[str, Any]:
    """Un refus attendu : son statut et un mot de son message."""
    assert r.status_code == statut, r.text
    corps: dict[str, Any] = r.json()
    assert mot in corps["detail"], corps
    return corps


def poids(**par_role: float | str) -> list[dict[str, Any]]:
    """`poids(Poil=1.2, Trame="0.4")` → la liste des poids au m² par id de rôle."""
    return [{"role_id": role_id(r), "poids": p} for r, p in par_role.items()]


def lignes(parts: dict[int, float | str]) -> list[dict[str, Any]]:
    """{matière : part} → les lignes d'une recette."""
    return [{"matiere_id": m, "pourcentage": p} for m, p in parts.items()]


def creer_qualite(
    direction: TestClient,
    code: str | None = None,
    *,
    poids_m2: dict[str, float | str] | None = None,
    parts: dict[int, float | str] | None = None,
    actif: bool = True,
) -> dict[str, Any]:
    """Une qualité créée par l'API (et sa recette, si `parts`) ; rend sa fiche."""
    fiche = ok(
        direction.post(
            "/api/qualites",
            json={
                "code": code or unique("Q"),
                "nom": "Essai D2",
                "actif": actif,
                "poids": poids(**(poids_m2 or {})),
            },
        ),
        201,
    )
    if parts is not None:
        fiche = ok(
            direction.put(
                f"/api/qualites/{fiche['id']}/recette",
                json={"lignes": lignes(parts), "version": fiche["version"]},
            )
        )
    return fiche


def lire_qualite(client: TestClient, qid: int) -> dict[str, Any]:
    """La fiche d'une qualité."""
    return ok(client.get(f"/api/qualites/{qid}"))  # type: ignore[no-any-return]


def corps_fiche(fiche: dict[str, Any], **changements: Any) -> dict[str, Any]:
    """Le corps d'un `PUT /qualites/{id}` repris de la fiche lue (fiche entière)."""
    corps = {
        "code": fiche["code"],
        "nom": fiche["nom"],
        "actif": fiche["actif"],
        "poids": [{"role_id": p["role_id"], "poids": p["poids"]} for p in fiche["poids"]],
        "version": fiche["version"],
    }
    return corps | changements


def creer_metier(direction: TestClient, **champs: Any) -> dict[str, Any]:
    """Un métier créé par l'API (code unique) ; rend sa fiche."""
    corps = {"code": unique("M"), "laize": 4, "bobines_par_metre": 10, "etages": 3} | champs
    return ok(direction.post("/api/metiers", json=corps), 201)  # type: ignore[no-any-return]


def lire_metier(client: TestClient, mid: int) -> dict[str, Any]:
    """La fiche d'un métier."""
    return ok(client.get(f"/api/metiers/{mid}"))  # type: ignore[no-any-return]


def version_etage(fiche: dict[str, Any], etage: int) -> int:
    """La version d'un étage lue dans la fiche du métier (1 pour un étage qui n'existe pas)."""
    return next((e["version"] for e in fiche["etages_detail"] if e["etage"] == etage), 1)  # type: ignore[no-any-return]


def ordre_valide(
    qualite_id: int,
    m2: float | str,
    *,
    jour: date | None = None,
    metier_id: int | None = None,
    statut: StatutOrdre = StatutOrdre.valide,
    s: Session | None = None,
) -> int:
    """Un ordre de fabrication (décor : écrit en base, sans sortie) qui tisse `m2` d'une
    qualité ; validé par défaut."""

    def faire(x: Session) -> int:
        le = jour or horloge.aujourd_hui()
        valide = statut == StatutOrdre.valide
        o = Ordre(
            numero=unique("OF") if valide else None,
            statut=statut,
            date=le,
            magasin_id=magasin_seme(s=x),
            metier_id=metier_id,
            valide_le=datetime.combine(le, time(12), horloge.FUSEAU) if valide else None,
        )
        x.add(o)
        x.flush()
        x.add(OrdreQualite(ordre_id=o.id, qualite_id=qualite_id, m2=Decimal(str(m2))))
        return o.id

    return _ecrire(s, faire)


def sortir(lot_id: int, kg: float | str, *, jour: date | None = None) -> None:
    """Une sortie vers les machines (décor : document « reprise »), par le seul chemin."""

    def faire(x: Session) -> None:
        lot = x.get(Lot, lot_id)
        assert lot is not None
        stock.ecrire(
            x,
            type=TypeMouvement.SORTIE,
            lot=lot,
            magasin_id=magasin_seme(s=x),
            quantite=-Decimal(str(kg)),
            jour=jour or horloge.aujourd_hui(),
            piece=stock.Piece(DocumentMouvement.reprise, None, ""),
            auteur_id=None,
        )

    _ecrire(None, faire)


def peser_lot(lot_id: int, poids_bobine: float | str) -> None:
    """Le poids d'une bobine mesuré sur un lot (M43)."""

    def faire(x: Session) -> None:
        lot = x.get(Lot, lot_id)
        assert lot is not None
        lot.poids_bobine = Decimal(str(poids_bobine))

    _ecrire(None, faire)
