"""Routes des opérations et de « À valider » (lot D5 ; API-PRO 115–123, 241). Les données sont
validées en base, à des noms uniques (les routes ont leurs propres sessions).
"""

from __future__ import annotations

from collections.abc import Iterator
from io import BytesIO
from typing import Any

import pytest
from fabrique import *  # noqa: F403
from fabriques.operations import Monde, monde, motif, ordre_valide
from fabriques.stock import JOUR, jour
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from test_prix_masques import fuites

from gestion_mp_pro.db import fabrique


@pytest.fixture
def base() -> Iterator[Session]:
    """Une session dont les écritures sont validées."""
    with fabrique()() as s:
        yield s


@pytest.fixture
def m(base: Session) -> Monde:
    """Le décor des essais hérités, validé en base."""
    decor = monde(base)
    base.commit()
    return decor


def _ok(r: Any, code: int = 200) -> Any:
    assert r.status_code == code, r.text
    return r.json() if r.content else None


def _ligne(m: Monde, lot: str, q: float, **champs: Any) -> dict[str, Any]:
    x = m.lots[lot]
    g = m.quarantaine if lot == "E-Q1" else m.courant
    return {"matiere_id": x.matiere_id, "lot_id": x.id, "magasin_id": g.id, "quantite": q, **champs}


def test_ouverture_valider_d_un_coup_et_prix_masques(
    m: Monde, achats: TestClient, alimentation: TestClient
) -> None:
    """`valider: true` (une transaction) ; l'alimentation lit l'opération sans prix ni valeur
    (balayage des clés monétaires) ; 404 « Opération introuvable. »."""
    corps = {"type": "ouverture", "date": JOUR.isoformat(), "valider": True, "lignes": [
        {"matiere_id": m.a2.id, "lot_numero": "OUV-API-1", "magasin_id": m.courant.id,
         "quantite": 12.5, "prix": 8}]}  # fmt: skip
    o = _ok(achats.post("/api/operations", json=corps), 201)
    assert o["numero"].startswith("OUV-") and o["statut"] == "validee"
    assert o["valeurs"] == [{"devise": "MAD", "montant": 100.0}] and o["lignes"][0]["prix"] == 8
    vu = _ok(alimentation.get(f"/api/operations/{o['id']}"))
    assert fuites(vu) == [] and vu["valeurs"] is None and vu["lignes"][0]["quantite"] == 12.5
    page = _ok(alimentation.get("/api/operations", params={"q": "OUV-API-1"}))
    assert page["total"] == 1 and fuites(page) == []
    r = achats.get("/api/operations/999999")
    assert r.status_code == 404 and r.json()["detail"] == "Opération introuvable."
    refus = achats.post("/api/operations", json={**corps, "lignes": [
        {**corps["lignes"][0], "devise": "XXX"}]})  # fmt: skip
    assert refus.status_code == 422


def test_liste_paginee_cherchee_sans_accents(m: Monde, base: Session, achats: TestClient) -> None:
    """C-13 : la liste se pagine (total, curseur) au lieu d'être coupée en silence ; la
    recherche ignore accents et casse."""
    mise = motif(base, "transfert", "Mise en quarantaine")
    ids = []
    for q in (1, 2, 3):
        corps = {
            "type": "transfert",
            "date": JOUR.isoformat(),
            "motif_id": mise,
            "commentaire": "Rangé près du métier Désiré",
            "lignes": [_ligne(m, "E-L1", q, magasin_destination_id=m.quarantaine.id)],
        }
        ids.append(_ok(achats.post("/api/operations", json=corps), 201)["id"])
    page1 = _ok(achats.get("/api/operations", params={"q": "metier desire", "limite": 2}))
    assert page1["total"] == 3 and len(page1["elements"]) == 2 and page1["suivant"]
    page2 = _ok(
        achats.get("/api/operations", params={"q": "metier desire", "apres": page1["suivant"]})
    )
    assert [x["id"] for x in page1["elements"] + page2["elements"]] == sorted(ids, reverse=True)
    assert achats.get("/api/operations", params={"apres": "%%"}).status_code == 422
    for i in ids:
        _ok(achats.delete(f"/api/operations/{i}", params={"version": 1}), 204)


def test_version_perimee_sur_chaque_ecriture(m: Monde, base: Session, achats: TestClient) -> None:
    """API-PRO § 0.1 : une version périmée → 409 `conflit_version` (PUT, DELETE, valider,
    annuler) ; rien n'a bougé."""
    corps = {"type": "ajustement", "date": JOUR.isoformat(), "commentaire": "palette retrouvée",
             "motif_id": motif(base, "ajustement", "Matière retrouvée"),
             "lignes": [_ligne(m, "E-T1", 5)]}  # fmt: skip
    o = _ok(achats.post("/api/operations", json=corps), 201)
    _ok(achats.put(f"/api/operations/{o['id']}", json={**corps, "version": 1}))
    for r in (
        achats.put(f"/api/operations/{o['id']}", json={**corps, "version": 1}),
        achats.delete(f"/api/operations/{o['id']}", params={"version": 1}),
        achats.post(f"/api/operations/{o['id']}/valider", json={"version": 1}),
    ):
        assert r.status_code == 409 and r.json()["code"] == "conflit_version", r.text
    v = _ok(achats.post(f"/api/operations/{o['id']}/valider", json={"version": 2}))
    r = achats.post(f"/api/operations/{o['id']}/annuler", json={"version": 2})
    assert r.status_code == 409 and r.json()["code"] == "conflit_version"
    a = _ok(achats.post(f"/api/operations/{o['id']}/annuler", json={"version": v["version"]}))
    assert a["statut"] == "annulee" and a["annulee_par"] == "Achats"


def test_brouillons_a_valider_et_tout_valider(
    m: Monde, base: Session, achats: TestClient, alimentation: TestClient, direction: TestClient
) -> None:
    """Hérité : test_stock_atelier::test_brouillons_a_valider_et_tout_valider (partie
    opérations) et test_defauts_c::test_tout_valider_achats_signale_ordres_non_permis_et_liens
    (C-17) : liste triée (date, rang) ; les achats voient le retour des machines sans pouvoir le
    valider, et « Tout valider » le dit (`ignores` + raison) ; un refus reste brouillon, intact."""
    o = ordre_valide(base, [(m.lots["E-L2"], 40)])
    base.commit()
    reliquat = motif(base, "retour_machine", "Reliquat de fin d'ordre")
    retour = _ok(alimentation.post("/api/operations", json={
        "type": "retour_machine", "date": JOUR.isoformat(), "ordre_id": o.id,
        "motif_id": reliquat, "lignes": [_ligne(m, "E-L2", 3)]}), 201)  # fmt: skip
    ajustement = _ok(achats.post("/api/operations", json={
        "type": "ajustement", "date": jour(-1).isoformat(), "commentaire": "palette derrière",
        "motif_id": motif(base, "ajustement", "Matière retrouvée"),
        "lignes": [_ligne(m, "E-T1", 5)]}), 201)  # fmt: skip
    trop = _ok(achats.post("/api/operations", json={
        "type": "retour_fournisseur", "date": JOUR.isoformat(),
        "motif_id": motif(base, "retour_fournisseur", "Bobines endommagées"),
        "lignes": [_ligne(m, "E-T1", 100000)]}), 201)  # fmt: skip
    assert any("demandés" in x for x in trop["blocages"])
    liste = _ok(achats.get("/api/brouillons"))
    cles = [(x["date"], x["id"]) for x in liste["lignes"]]
    assert cles == sorted(cles, key=lambda c: c[0])
    vue = next(x for x in liste["operations"] if x["id"] == retour["id"])
    assert not vue["peut_valider"] and "alimentation" in vue["raison"]
    assert vue["lien"] == f"/operations/{retour['id']}"
    choix = [{"genre": "operation", "id": x["id"]} for x in (retour, ajustement, trop)]
    r = _ok(achats.post("/api/brouillons/valider", json={"documents": choix}))
    assert [x["id"] for x in r["valides"]] == [ajustement["id"]]
    assert r["valides"][0]["numero"].startswith("AJU-")
    assert [(x["id"], "réservé" in x["raison"]) for x in r["ignores"]] == [(retour["id"], True)]
    assert [x["id"] for x in r["refuses"]] == [trop["id"]] and "ne porte que" in r["refuses"][0][
        "raison"
    ]
    assert _ok(achats.get(f"/api/operations/{trop['id']}"))["statut"] == "brouillon"
    tout = _ok(alimentation.post("/api/brouillons/valider", json={"tous": True}))
    assert retour["id"] in [x["id"] for x in tout["valides"]]
    assert direction.post("/api/brouillons/valider", json={}).status_code == 422
    _ok(achats.delete(f"/api/operations/{trop['id']}", params={"version": trop["version"]}), 204)


def test_classeur_des_operations(m: Monde, achats: TestClient, alimentation: TestClient) -> None:
    """Route 241 : une ligne par ligne d'opération ; pour l'alimentation, ni prix ni devise, et
    ses deux types seulement (transferts, retours des machines)."""
    corps = {"type": "ouverture", "date": JOUR.isoformat(), "valider": True, "lignes": [
        {"matiere_id": m.a2.id, "lot_numero": "XLS-1", "magasin_id": m.courant.id, "quantite": 2,
         "prix": 8}]}  # fmt: skip
    _ok(achats.post("/api/operations", json=corps), 201)

    def feuille(client: TestClient) -> list[tuple[Any, ...]]:
        r = client.get("/api/export/operations", params={"q": "XLS-1"})
        assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd")
        ws = load_workbook(BytesIO(r.content)).active
        assert ws is not None
        rangs = list(ws.iter_rows(values_only=True))
        return rangs[next(i for i, x in enumerate(rangs) if x[0] == "Numéro") :]

    entete, *rangs = feuille(achats)
    assert "Prix / kg" in entete and len(rangs) == 1 and rangs[0][entete.index("Lot")] == "XLS-1"
    entete, *rangs = feuille(alimentation)
    assert "Prix / kg" not in entete and "Devise" not in entete and rangs == []
