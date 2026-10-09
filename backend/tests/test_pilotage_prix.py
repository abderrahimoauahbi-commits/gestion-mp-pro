"""Prix payés, matrice, opportunités, note des fournisseurs et classeurs de l'analyse (lot D8 ;
routes 211, 213–220 ; défauts E-04, E-05, E-12, P8-23 ; domaine 12 de ESSAIS.md).
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

import pytest
from fabrique import *  # noqa: F403
from fabriques.pilotage import Usine, fixture_usine, jour, ok  # noqa: F401
from fastapi.testclient import TestClient
from openpyxl import load_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def recevoir(u: Usine, f: int, m: int, prix: float, devise: str, lot: str, le: int = 0) -> None:  # noqa: PLR0917
    """Une réception validée hors commande, dans la devise donnée."""
    r = ok(u.achats.post("/api/receptions", json={
        "date": jour(le).isoformat(), "fournisseur_id": f, "lignes": [
            {"matiere_id": m, "lot_numero": lot, "quantite_pesee": 100, "magasin_id": u.courant,
             "prix": prix, "devise": devise}]}), 201)[0]  # fmt: skip
    ok(u.achats.post(f"/api/receptions/{r['id']}/valider", json={"version": r["version"]}))


def lignes_du_classeur(contenu: bytes, feuille: str) -> list[tuple[Any, ...]]:
    """Les lignes d'une feuille d'un classeur."""
    return list(load_workbook(BytesIO(contenu))[feuille].iter_rows(values_only=True))


def test_matrice_ne_melange_pas_les_devises(usine: Usine) -> None:
    """Hérité : test_defauts_e::test_matrice_ne_melange_pas_les_devises (E-04) ;
    test_direction::test_M80_la_matrice_s_exporte_dans_la_devise_de_la_facture — une ligne
    par devise, prix dans la monnaie de la facture, à l'écran comme dans le classeur."""
    u = usine
    f = u.fournisseur("MAT FILATURE", devise="USD")
    m = u.matiere("MAT FIL", fournisseur_id=f, devise="USD", prix="2")
    recevoir(u, f, m, 2, "USD", "MAT-1")
    recevoir(u, f, m, 3, "EUR", "MAT-2")
    matrice = u.get(u.achats, "/api/prix/matrice")
    par_devise = {
        x["devise"]: x["prix"][-1] for x in matrice["lignes"] if x["reference"] == "MAT FIL"
    }
    assert par_devise == {"USD": 2, "EUR": 3} and matrice["mois"] == ["2026-10"]
    assert all(x["matiere_id"] == m for x in matrice["lignes"])
    assert u.get(u.achats, "/api/prix/matrice", q="introuvable")["lignes"] == []
    lignes = lignes_du_classeur(u.achats.get("/api/export/prix").content, "Matrice des prix")
    usd = next(x for x in lignes if x[0] == "MAT FIL" and x[6] == "USD")
    assert usd[7] == pytest.approx(2.0)


def test_l_historique_des_prix_se_cherche_et_compte_ses_hausses(usine: Usine) -> None:
    """Prix payés : le précédent du même fournisseur et de la même devise, l'évolution, les
    hausses (fortes au-delà du réglage) ; recherche et dates au serveur ; vraies dates à
    l'export (E-12, hérité : test_defauts_e::test_export_prix_dates_sont_des_dates)."""
    u = usine
    f = u.fournisseur("HIST FILATURE")
    m = u.matiere("HIST FIL", fournisseur_id=f)
    recevoir(u, f, m, 10, "MAD", "H-1", le=-10)
    recevoir(u, f, m, 12, "MAD", "H-2")
    h = u.get(u.achats, "/api/prix/historique")
    dernier = h["lignes"][0]
    assert (dernier["prix_achat"], dernier["prix_precedent"]) == (12, 10)
    assert dernier["evolution"] == pytest.approx(0.2)
    assert (h["hausses"], h["fortes_hausses"], h["baisses"]) == (1, 1, 0)
    assert h["seuils"]["forte"] == pytest.approx(0.05) and dernier["matiere_id"] == m
    assert len(u.get(u.achats, "/api/prix/historique", du=jour(-1).isoformat())["lignes"]) == 1
    assert u.get(u.achats, "/api/prix/historique", q="filature hist")["total_achats_mad"] == 2200
    lignes = lignes_du_classeur(u.achats.get("/api/export/prix").content, "Historique des achats")
    debut = next(i for i, x in enumerate(lignes) if x and x[0] == "Date") + 1
    assert lignes[debut:] and all(hasattr(x[0], "year") for x in lignes[debut:] if x[0])


def test_M82_les_opportunites_disent_leur_seuil_et_leurs_criteres(usine: Usine) -> None:
    """Hérité : test_direction::test_M82_les_opportunites_disent_leur_seuil ;
    test_direction::test_recette_les_opportunites_comparent_la_meme_couleur_par_defaut ;
    test_defauts_e::test_export_opportunites_fenetre_negative_422 (E-05)."""
    u = usine
    o = u.get(u.achats, "/api/opportunites")
    assert o["seuils"]["forte"] == pytest.approx(0.05)
    assert {"sans_date", "hors_fenetre"} <= set(o["ecartees"])
    assert "coût rendu" in o["avertissement"] and o["criteres"]["meme_couleur"] is True
    assert o["fenetre_jours"] == 90
    assert u.get(u.achats, "/api/opportunites", fenetre_jours=30)["fenetre_jours"] == 30
    for fenetre in (-5, 0, 3651):
        for url in ("/api/opportunites", "/api/export/opportunites"):
            assert u.achats.get(url, params={"fenetre_jours": fenetre}).status_code == 422
    r = u.achats.get("/api/export/opportunites", params={"date": jour(-5).isoformat()})
    titre = lignes_du_classeur(r.content, "Opportunités")[0][0]
    assert "03/10/2026" in str(titre)


def test_la_note_des_fournisseurs_se_lit_par_id(usine: Usine, alimentation: TestClient) -> None:
    """Route 215 : servie malgré la fiche `/fournisseurs/{id}` (l'entier départage) ; chaque
    ligne porte l'id de sa fiche, le classement sans émoji (P8-23) ; période bornée ; classeur à
    deux feuilles ; l'alimentation ne la lit pas."""
    u = usine
    f = u.fournisseur("NOTE FILATURE")
    m = u.matiere(fournisseur_id=f)
    recevoir(u, f, m, 10, "MAD", "N-1", le=-3)
    p = u.get(u.achats, "/api/fournisseurs/performance")
    ligne = next(x for x in p["lignes"] if x["fournisseur"] == "NOTE FILATURE")
    assert ligne["fournisseur_id"] == f and ligne["receptions"] == 1
    assert not any(x["classement"][:1] in "🟢🟡🟠🔴" for x in p["lignes"])
    assert p["receptions"][0]["reception_id"] and p["receptions"][0]["fournisseur_id"] == f
    vide = u.get(u.achats, "/api/fournisseurs/performance", du=jour(-1).isoformat())
    assert vide["receptions"] == [] and vide["du"] == jour(-1).isoformat()
    assert u.get(u.achats, f"/api/fournisseurs/{f}")["id"] == f
    r = u.achats.get("/api/export/fournisseurs")
    assert load_workbook(BytesIO(r.content)).sheetnames == ["Fournisseurs", "Réceptions notées"]
    assert alimentation.get("/api/fournisseurs/performance").status_code == 403


@pytest.mark.parametrize(
    ("chemin", "feuille"),
    [
        ("/api/export/controles", "Contrôles"),
        ("/api/export/valorisation", "Valorisation"),
        ("/api/export/prix", "Historique des achats"),
        ("/api/export/opportunites", "Opportunités"),
        ("/api/export/fournisseurs", "Fournisseurs"),
        ("/api/export/echeancier", "Échéances"),
    ],
)
def test_M80_l_analyse_s_exporte(
    achats: TestClient, alimentation: TestClient, chemin: str, feuille: str
) -> None:
    """Hérité : test_direction::test_M80_l_analyse_s_exporte — un classeur par onglet de
    l'analyse ; l'alimentation n'y a pas accès."""
    r = achats.get(chemin)
    assert r.status_code == 200, r.text[:300]
    assert r.headers["content-type"].startswith(XLSX)
    assert r.headers["content-disposition"].startswith("attachment")
    assert feuille in load_workbook(BytesIO(r.content)).sheetnames
    assert alimentation.get(chemin).status_code == 403


def test_export_des_controles_gravite_en_clair(usine: Usine) -> None:
    """Hérité : test_defauts_e::test_export_pilotage_controles_gravite_en_clair (E-12) —
    gravités et états en clair ; une feuille par contrôle en anomalie, écart en pourcentage."""
    u = usine
    u.ouvrir(u.matiere("EXPORT C07"), 1, 1)
    u.matiere("EXPORT JAMAIS")
    contenu = u.achats.get("/api/export/controles").content
    valeurs = {str(x) for ligne in lignes_du_classeur(contenu, "Contrôles") for x in ligne}
    assert "bloquant" not in valeurs and "Bloquant" in valeurs
    assert any(v.startswith("garanti") for v in valeurs)
    noms = load_workbook(BytesIO(contenu)).sheetnames
    assert any(n.startswith("C07") for n in noms)
