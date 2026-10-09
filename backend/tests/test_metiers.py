"""Métiers à tisser et montage (routes 94–102, 238 ; lot D2).

Héritage REGLES A § 1.11 (R-REF-74 à R-REF-84, X-REF-59 à X-REF-73) ; défaut D-12 (pas de
fil monté sur un métier arrêté, ni d'une matière désactivée) ; décision de Mohamed : le fil
monté est du stock « sur les métiers » (compté dans le stock global, le disponible et le
projeté) et monter n'écrit aucun mouvement.
"""

from __future__ import annotations

import io
from typing import Any

import pytest
from fabrique import *  # noqa: F403
from fabriques.qualites import (
    creer_metier,
    lire_metier,
    ok,
    ordre_valide,
    peser_lot,
    refus,
    sortir,
    version_etage,
)
from fabriques.referentiels import matiere, qualite, stock_initial
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from gestion_mp_pro.calcul.stock import stock_en_machine
from gestion_mp_pro.db import fabrique as sessions
from gestion_mp_pro.modeles import Journal, Matiere, Mouvement
from gestion_mp_pro.services import donnees_moteur


def _fil(matiere_id: int, metres: float, part: float = 100, commentaire: str = "") -> dict:
    return {
        "matiere_id": matiere_id,
        "metres": metres,
        "pourcentage": part,
        "commentaire": commentaire,
    }


def _etage(client: TestClient, x: dict, n: int, fils: list[dict], effet: float) -> Any:
    """Monte un étage (route 100) à la version lue dans `x`."""
    return client.put(
        f"/api/metiers/{x['id']}/etages/{n}",
        json={"version": version_etage(x, n), "fils": fils, "confirme_effet_kg": effet},
    )


def _reglage(x: dict, **changements: Any) -> dict:
    """Le corps d'un `PUT /metiers/{id}` repris de la fiche lue."""
    cles = (
        "code",
        "nom",
        "laize",
        "bobines_par_metre",
        "etages",
        "actif",
        "commentaire",
        "version",
    )
    return {k: x[k] for k in cles} | changements


def test_monter_un_etage_au_telephone(
    direction: TestClient, alimentation: TestClient, achats: TestClient
) -> None:
    """Hérité : test_referentiels_complets::test_monter_un_etage_au_telephone."""
    poil = matiere(poids_bobine="2")
    cuir = matiere(role="Cuir")
    x = creer_metier(direction, commentaire="Van de Wiele")
    assert x["commentaire"] == "Van de Wiele" and x["etages_garnis"] == 0
    r = ok(_etage(alimentation, x, 1, [_fil(poil, 2, commentaire="bobines neuves")], 40))
    # Le montage dit ce qu'il fait au stock, et le journal garde qui a monté quoi.
    assert (r["effet_kg"], r["avant_kg"], r["apres_kg"]) == (40, 0, 40)
    assert r["message"].startswith("Monter n'écrit aucun mouvement : ce fil est déjà sorti")
    assert "(+40,0 kg)" in r["message"]
    with sessions()() as s:
        detail = s.scalars(
            select(Journal.detail)
            .where(Journal.action == "montage", Journal.objet_id == x["id"])
            .order_by(Journal.id.desc())
        ).first()
    assert detail is not None and " 2 m" in detail and "(+40,0 kg)" in detail
    x = r["metier"]
    ok(_etage(alimentation, x, 2, [_fil(poil, 1.5, 50)], 15))
    x = lire_metier(alimentation, x["id"])
    assert x["etages_garnis"] == 2 and x["fils"][0]["commentaire"] == "bobines neuves"
    assert x["poids"] == pytest.approx(2 * 10 * 2.0 + 1.5 * 10 * 2.0 * 0.5)
    # Les bobines d'un fil qui partage ses mètres (50 %) : 1,5 m × 10 × 50 % = 7,5.
    assert x["bobines"] == pytest.approx(20 + 7.5)
    assert x["etages_detail"][0]["poids"] == pytest.approx(40)
    assert x["dernier_montage"]["par"] == "Alimentation" and x["dernier_montage"]["etage"] == 2
    refs = {m["matiere_id"]: m for m in ok(alimentation.get("/api/montage/matieres"))}
    assert (
        refs[poil]["deja_monte"] == pytest.approx(55) and refs[poil]["origine_poids"] == "catalogue"
    )
    assert cuir not in refs
    # La laize est tenue, les mètres sont exigés, une matière au mètre ne se monte pas.
    refus(_etage(alimentation, x, 3, [_fil(poil, 5)], 0), 422, "laize")
    refus(_etage(alimentation, x, 3, [_fil(poil, 0)], 0), 422, "saisir les mètres montés")
    refus(_etage(alimentation, x, 3, [_fil(cuir, 1)], 0), 422, "se pose au mètre linéaire")
    refus(_etage(alimentation, x, 4, [], 0), 404, "Étage 4 : ce métier n'en a que 3.")
    # Les achats lisent, n'écrivent pas.
    assert _etage(achats, x, 3, [], 0).status_code == 403


def test_deux_telephones_sur_le_meme_etage(direction: TestClient, alimentation: TestClient) -> None:
    """Le montage d'un étage est protégé par sa version : un seul des deux passe."""
    poil = matiere(poids_bobine="1")
    x = creer_metier(direction)
    ok(_etage(alimentation, x, 1, [_fil(poil, 1)], 10))
    r = _etage(direction, x, 1, [_fil(poil, 2)], 20)
    assert refus(r, 409, "L'étage 1 a été modifié entre-temps par Alimentation")["code"] == (
        "conflit_version"
    )
    # Un autre étage, lu à la même version, passe : les étages ne se gênent pas.
    ok(_etage(direction, x, 2, [_fil(poil, 1)], 10))


def test_l_effet_annonce_doit_etre_le_bon(direction: TestClient, alimentation: TestClient) -> None:
    """`confirme_effet_kg` : un effet qui a changé depuis son calcul est refusé (409), rien
    n'est écrit."""
    poil = matiere(poids_bobine="2")
    x = creer_metier(direction)
    r = _etage(alimentation, x, 1, [_fil(poil, 2)], 30)
    assert refus(r, 409, "+40,0 kg au lieu de +30,0 kg")["code"] == "effet_change"
    assert lire_metier(direction, x["id"])["fils"] == []
    ok(_etage(alimentation, x, 1, [_fil(poil, 2)], "40.04"))


def test_deplacer_un_fil_d_etage_au_bureau(direction: TestClient, alimentation: TestClient) -> None:
    """Hérité : test_referentiels_complets::test_deplacer_un_fil_d_etage_au_bureau."""
    poil = matiere(poids_bobine="2")
    x = creer_metier(direction)
    x = ok(_etage(alimentation, x, 1, [_fil(poil, 2)], 40))["metier"]
    x = ok(_etage(alimentation, x, 2, [_fil(poil, 1)], 20))["metier"]
    corps = {
        "etages": [
            {"etage": 2, "version": version_etage(x, 2), "fils": []},
            {"etage": 3, "version": version_etage(x, 3), "fils": [_fil(poil, 1)]},
        ],
        "confirme_effet_kg": 0,
    }
    r = ok(alimentation.put(f"/api/metiers/{x['id']}/montage", json=corps))
    assert sorted(f["etage"] for f in r["metier"]["fils"]) == [1, 3]
    assert r["message"] == "Le stock sur les métiers ne change pas."
    assert alimentation.put(f"/api/metiers/{x['id']}/montage", json=corps).status_code == 409
    historique = ok(alimentation.get(f"/api/metiers/{x['id']}/montages"))
    assert [h["origine"] for h in historique] == ["bureau", "téléphone", "téléphone"]
    assert historique[0]["etage"] is None and historique[0]["avant_kg"] == 20
    assert historique[0]["apres_kg"] == 20 and historique[0]["effet_kg"] == 0


def test_poids_de_bobine_mesure(direction: TestClient, alimentation: TestClient) -> None:
    """Hérité : test_referentiels_complets::test_poids_de_bobine_mesure (lot pesé : 2,3 kg)."""
    poil = matiere(poids_bobine="2")
    x = creer_metier(direction)
    ok(_etage(alimentation, x, 1, [_fil(poil, 2)], 40))
    lot = stock_initial(poil, 230)
    peser_lot(lot, "2.3")
    f = lire_metier(alimentation, x["id"])["fils"][0]
    # Le fil monté se compte au poids pesé du lot (M43), pas au catalogue…
    assert f["poids_bobine_utilise"] == pytest.approx(2.3) and "(reçu)" in f["source_poids"]
    assert f["poids"] == pytest.approx(f["bobines"] * 2.3)
    assert f["ecart_mesure"] == pytest.approx(0.15) and f["ecart_hors_tolerance"] is True
    # … et le stock « sur les métiers » du moteur compte le même poids.
    with sessions()() as s:
        en_machine = stock_en_machine(donnees_moteur.charger(s, avec_journal=False))
        reference = s.get(Matiere, poil).reference  # type: ignore[union-attr]
    assert en_machine[reference] == pytest.approx(f["poids"])


def test_regler_et_supprimer_un_metier(direction: TestClient, alimentation: TestClient) -> None:
    """Hérité : test_referentiels_complets::test_regler_et_supprimer_un_metier."""
    poil = matiere(poids_bobine="1")
    x = creer_metier(direction)
    x = ok(_etage(alimentation, x, 3, [_fil(poil, 3)], 30))["metier"]
    url = f"/api/metiers/{x['id']}"
    assert alimentation.put(url, json=_reglage(x)).status_code == 403
    refus(direction.put(url, json=_reglage(x, etages=2)), 422, "montés à l'étage 3")
    refus(direction.put(url, json=_reglage(x, laize=1)), 422, "plus que la laize de 1 m demandée")
    r = ok(direction.put(url, json=_reglage(x, commentaire="révisé")))
    assert r["commentaire"] == "révisé" and r["version"] == x["version"] + 1
    refus(direction.put(url, json=_reglage(x)), 409, "modifié entre-temps par Direction")
    refus(direction.delete(f"{url}?version={r['version']}"), 409, "porte encore 1 fil(s)")
    ok(_etage(alimentation, r, 3, [], -30))
    autre = creer_metier(direction)
    refus(
        direction.put(f"/api/metiers/{autre['id']}", json=_reglage(autre, code=x["code"])),
        409,
        "existe déjà",
    )
    assert direction.delete(f"{url}?version={r['version']}").status_code == 204
    refus(direction.get(url), 404, "Métier introuvable.")


def test_un_metier_nomme_par_un_ordre_ne_se_supprime_pas(direction: TestClient) -> None:
    """X-REF-65."""
    x = creer_metier(direction)
    ordre_valide(qualite(), 5, metier_id=x["id"])
    x = lire_metier(direction, x["id"])
    assert x["ordres"] == 1
    refus(
        direction.delete(f"/api/metiers/{x['id']}?version={x['version']}"),
        409,
        "nommé par 1 ordre(s) de fabrication",
    )


def test_d12_pas_de_montage_sur_un_metier_arrete(
    direction: TestClient, alimentation: TestClient
) -> None:
    """Hérité : test_defauts_d::test_d12_montage_refuse_sur_un_metier_arrete ; et une matière
    désactivée ne se monte pas non plus (D-12)."""
    poil = matiere(poids_bobine="1")
    x = creer_metier(direction, etages=2)
    x = ok(_etage(alimentation, x, 1, [_fil(poil, 3)], 30))["metier"]
    x = ok(direction.put(f"/api/metiers/{x['id']}", json=_reglage(x, actif=False)))
    assert x["sur_les_metiers"] == 0 and x["poids"] == 30
    refus(
        _etage(alimentation, x, 1, [_fil(poil, 4)], 0),
        409,
        f"Le métier {x['code']} est arrêté : le remettre en service avant d'y monter du fil.",
    )
    refus(_etage(alimentation, x, 2, [_fil(poil, 1)], 0), 409, "arrêté")
    r = ok(_etage(alimentation, x, 1, [], 0))  # le démonter reste possible
    assert r["message"] == "Métier arrêté : rien n'est compté sur lui."
    inactive = matiere(poids_bobine="1", actif=False)
    y = creer_metier(direction)
    refus(_etage(alimentation, y, 1, [_fil(inactive, 1)], 10), 409, "est désactivée")


def test_alerte_au_dela_de_ce_qui_est_parti_en_30_jours(
    direction: TestClient, alimentation: TestClient
) -> None:
    """1,1 × sorti en 30 jours + 1 kg : au-delà, le montage avertit (et l'historique le garde)."""
    poil = matiere(poids_bobine="1")
    lot = stock_initial(poil, 500)
    sortir(lot, 100)
    x = creer_metier(direction)
    r = ok(_etage(alimentation, x, 1, [_fil(poil, 4)], 40))  # 40 kg ≤ 111 : rien à dire
    assert r["alertes"] == []
    r = ok(_etage(alimentation, r["metier"], 2, [_fil(poil, 4)], 40))
    assert r["alertes"] == []
    y = creer_metier(direction)
    r = ok(_etage(alimentation, y, 1, [_fil(poil, 4)], 40))  # 120 kg sur les métiers
    assert len(r["alertes"]) == 1
    assert (
        "120,0 kg sur les métiers pour 100,0 kg partis aux machines en 30 jours" in r["alertes"][0]
    )
    assert ok(alimentation.get(f"/api/metiers/{y['id']}/montages"))[0]["alerte"] == r["alertes"][0]
    with sessions()() as s:  # monter n'a écrit aucun mouvement
        n = s.scalar(
            select(func.count()).select_from(Mouvement).where(Mouvement.matiere_id == poil)
        )
    assert n == 2


def test_reglage_et_fils_en_une_transaction(direction: TestClient) -> None:
    """Route 97 : réglage et montage ensemble, tout ou rien (an. 14)."""
    poil = matiere(poids_bobine="1")
    x = creer_metier(direction, etages=2)
    montage = {
        "etages": [{"etage": 3, "version": 1, "fils": [_fil(poil, 2)]}],
        "confirme_effet_kg": 20,
    }
    r = direction.put(f"/api/metiers/{x['id']}", json=_reglage(x, montage=montage, nom="Refusé"))
    refus(r, 422, "Des fils sont montés à l'étage 3")
    assert lire_metier(direction, x["id"])["nom"] == ""
    r = ok(
        direction.put(
            f"/api/metiers/{x['id']}",
            json=_reglage(x, etages=3, nom="Trois étages", montage=montage),
        )
    )
    assert r["nom"] == "Trois étages" and r["montage"]["effet_kg"] == 20
    assert [e["etage"] for e in r["etages_detail"]] == [1, 2, 3] and r["etages_garnis"] == 1


def test_le_stock_est_en_magasin_ou_sur_les_metiers(
    direction: TestClient, alimentation: TestClient, achats: TestClient
) -> None:
    """Hérité : test_audit2_atelier::test_le_stock_est_en_magasin_ou_en_machine (partie stock ;
    la valorisation et le cockpit sont au lot D8)."""
    monte, autre = matiere(poids_bobine="2", prix="10"), matiere(poids_bobine="1")
    x = creer_metier(direction, laize="4.2", etages=2)
    # Deux fils qui se partagent les mêmes 4 m à 60 et 40 % tiennent dans 4,2 m.
    r = ok(_etage(alimentation, x, 1, [_fil(monte, 4, 60), _fil(autre, 4, 40)], 64))
    assert r["metier"]["etages_detail"][0]["metres"] == pytest.approx(4)
    assert not r["metier"]["etages_trop_larges"]
    refus(_etage(alimentation, r["metier"], 2, [_fil(monte, 3), _fil(autre, 2)], 0), 422, "laize")
    ligne = next(x for x in ok(achats.get("/api/stock"))["lignes"] if x["matiere_id"] == monte)
    assert ligne["stock_machine"] == pytest.approx(48)
    assert ligne["stock_disponible"] == pytest.approx(48)
    assert ligne["stock_projete"] == pytest.approx(48)


def test_valeur_sur_les_metiers_et_prix_masques(
    direction: TestClient, achats: TestClient, alimentation: TestClient
) -> None:
    """La valeur (poids × coût moyen) pour qui voit les prix ; rien pour l'alimentation."""
    from test_prix_masques import fuites

    poil = matiere(poids_bobine="1", prix="10")
    stock_initial(poil, 100, prix="10")
    x = creer_metier(direction)
    ok(_etage(alimentation, x, 1, [_fil(poil, 2)], 20))
    assert lire_metier(achats, x["id"])["valeur_mad"] == pytest.approx(200)
    lu = lire_metier(alimentation, x["id"])
    assert lu["valeur_mad"] is None and lu["fils"][0]["valeur_mad"] is None
    for chemin in (
        "/api/metiers",
        f"/api/metiers/{x['id']}",
        f"/api/metiers/{x['id']}/montages",
        "/api/montage/matieres",
    ):
        assert fuites(ok(alimentation.get(chemin))) == [], chemin
    entetes = []
    for client in (achats, alimentation):
        r = client.get("/api/export/metiers")
        assert r.status_code == 200 and r.content[:2] == b"PK"
        feuille = load_workbook(io.BytesIO(r.content)).active
        assert feuille is not None
        entetes.append(next(feuille.iter_rows(values_only=True)))
    assert "Valeur (MAD)" in entetes[0] and "Valeur (MAD)" not in entetes[1]


def test_creer_un_metier(direction: TestClient, achats: TestClient) -> None:
    """Route 96 : direction seule ; code unique ; une ligne d'étage par étage (version 1)."""
    x = creer_metier(direction, etages=8)
    assert [e["version"] for e in x["etages_detail"]] == [1] * 8 and x["densite_a_saisir"] is False
    assert achats.post("/api/metiers", json={"code": "X"}).status_code == 403
    refus(
        direction.post("/api/metiers", json={"code": x["code"]}),
        409,
        f"Le métier « {x['code']} » existe déjà.",
    )
    refus(direction.post("/api/metiers", json={"code": "Z", "etages": 41}), 422)
