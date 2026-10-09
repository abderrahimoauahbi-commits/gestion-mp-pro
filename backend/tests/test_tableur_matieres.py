"""Le catalogue par tableur : modèle, aperçu, application (routes 78–80) — lot D1.

Héritage R-REF-49…54, X-REF-36…38 (REGLES § 1.6.8) ; défaut A-03 (aperçu = application, sans
500, doublons, nombres non finis, « En service » mal écrit, casse des codes, empreinte) ; P7-9…12
(colonnes « Priorité dans le groupe », « Vider », « Version »).
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from fabrique import *  # noqa: F403
from fabriques.referentiels import devise, fournisseur, matiere, qualite, unique
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from gestion_mp_pro.db import fabrique as sessions
from gestion_mp_pro.modeles import Fournisseur, Journal, Matiere

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def classeur(lignes: list[dict[str, Any]], feuille: str = "Matières", active: bool = True) -> bytes:
    """Un tableur dont la première ligne porte les intitulés des colonnes utilisées."""
    wb = Workbook()
    if not active:
        wb.active.title = "Notes"
        f = wb.create_sheet(feuille)
    else:
        f = wb.active
        f.title = feuille
    entetes = list(dict.fromkeys(k for x in lignes for k in x))
    f.append(entetes)
    for x in lignes:
        f.append([x.get(k) for k in entetes])
    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def apercu(client: TestClient, contenu: bytes) -> Any:
    """`POST /matieres/tableur/apercu`."""
    return client.post("/api/matieres/tableur/apercu", files={"fichier": ("m.xlsx", contenu, XLSX)})


def appliquer(client: TestClient, contenu: bytes, empreinte: str) -> Any:
    """`POST /matieres/tableur/appliquer`."""
    return client.post(
        "/api/matieres/tableur/appliquer",
        files={"fichier": ("m.xlsx", contenu, XLSX)},
        data={"empreinte_apercu": empreinte},
    )


def ok(r: Any, statut: int = 200) -> Any:
    """La réponse JSON d'un appel qui doit réussir."""
    assert r.status_code == statut, r.text
    return r.json()


def lue(ident: int) -> Matiere:
    """Une matière relue en base."""
    with sessions()() as s:
        return s.get_one(Matiere, ident)


def test_modele_puis_import_apercu_egal_application(achats: TestClient) -> None:
    """Hérité : test_suivi::test_bilan_echeancier_et_import_des_matieres (import) ; R-REF-49,
    50, 54 — le modèle est le catalogue, rien n'est écrit à l'aperçu, l'application écrit ce qui
    a été vu, une matière absente n'est pas supprimée, chaque fiche a son journal."""
    m = matiere(prix="2")
    r = achats.get("/api/matieres/tableur")
    modele = load_workbook(BytesIO(r.content))
    f = modele["Matières"]
    entetes = [c.value for c in f[1]]
    assert entetes[:3] == ["Référence", "Désignation", "Rôle"]
    assert {"Priorité dans le groupe", "Vider", "Version", "Se commande en (lecture seule)"} <= set(
        entetes
    )
    reference = lue(m).reference
    nom = unique("Nouvelle")
    f_nom = fournisseur(nom=unique("F"))
    with sessions()() as s:
        fournisseur_nom = s.get_one(Fournisseur, f_nom).nom
    contenu = classeur([
        {"Référence": reference, "Prix": "2,5", "Devise": "mad", "En service": "OUI"},
        {"Référence": nom, "Rôle": "poil", "Fournisseur": fournisseur_nom.upper(), "Prix": 3},
    ])  # fmt: skip
    vu = ok(apercu(achats, contenu))
    assert (vu["creees"], vu["modifiees"], vu["erreurs"], vu["appliquable"]) == (1, 1, 0, True)
    modif = next(x for x in vu["lignes"] if x["reference"] == reference)
    assert modif["changements"] == [
        {"champ": "prix", "libelle": "Prix", "avant": "2", "apres": "2,5"}
    ]
    assert float(lue(m).prix) == 2
    fait = ok(appliquer(achats, contenu, vu["empreinte_apercu"]))
    assert fait["empreinte_apercu"] == vu["empreinte_apercu"]
    assert float(lue(m).prix) == 2.5 and lue(m).version == 2
    with sessions()() as s:
        nouvelle = s.scalars(select(Matiere).where(Matiere.reference == nom)).one()
        assert nouvelle.fournisseur_id == f_nom and nouvelle.devise == "MAD"
        actions = s.scalars(
            select(Journal.action).where(Journal.objet_id == m, Journal.objet == "matiere")
        ).all()
        assert "import des matières" in actions


def test_l_import_refuse_une_reference_en_double_sans_500(achats: TestClient) -> None:
    """Hérité : test_defauts_a::test_l_import_refuse_une_reference_en_double_sans_500 (A-03)."""
    ref = unique("Double")
    contenu = classeur(
        [{"Référence": ref, "Rôle": "Poil", "Fournisseur": "x"}, {"Référence": ref.lower()}]
    )
    vu = ok(apercu(achats, contenu))
    erreurs = {x["ligne"]: x["erreurs"] for x in vu["lignes"] if x["etat"] == "erreur"}
    assert erreurs[3] == [
        f"La référence « {ref.lower()} » figure déjà ligne 2 : une seule ligne par matière"
    ]
    r = appliquer(achats, contenu, vu["empreinte_apercu"])
    assert (r.status_code, r.json()["detail"]) == (
        422,
        "2 ligne(s) en erreur : corriger le tableur d'abord, rien n'a été écrit.",
    )


def test_l_import_refuse_un_prix_negatif_a_l_apercu(achats: TestClient) -> None:
    """Hérité : test_defauts_a::test_l_import_refuse_un_prix_negatif_a_l_apercu (A-03) — bornes,
    « nan », longueurs, « En service » mal écrit, casse des codes."""
    m = matiere()
    ref = lue(m).reference
    cas = {
        "Prix": (-1, "Prix : entre 0 et 10 000 000 — reçu -1"),
        "Densité": ("nan", "Densité : pas un nombre (« nan »)"),
        "Couverture minimum (mois)": (61, "Couverture minimum (mois) : entre 0 et 60 — reçu 61"),
        "Groupe d'équivalence": ("g" * 65, "Groupe d'équivalence : 64 caractères au plus"),
        "En service": ("ouii", "En service : oui ou non (« ouii »)"),
        "Devise": ("xyz", "Devise « xyz » inconnue : l'ajouter d'abord (Réglages › Devises)"),
        "Composition": (
            "ZZ9",
            "Composition « ZZ9 » inconnue : l'ajouter d'abord à la classification",
        ),
    }
    for colonne, (valeur, message) in cas.items():
        vu = ok(apercu(achats, classeur([{"Référence": ref, colonne: valeur}])))
        assert vu["lignes"][0]["erreurs"] == [message], colonne
    usd = devise()
    vu = ok(apercu(achats, classeur([{"Référence": ref, "Devise": usd.lower()}])))
    assert vu["lignes"][0]["etat"] == "modifiee"


def test_l_import_refuse_un_changement_de_role_en_recette(achats: TestClient) -> None:
    """Hérité : test_defauts_a::test_l_import_refuse_un_changement_de_role_en_recette (A-02)."""
    m = matiere()
    qualite(recette={m: 100}, poids={"Poil": "1"})
    vu = ok(apercu(achats, classeur([{"Référence": lue(m).reference, "Rôle": "Trame"}])))
    assert "son rôle ne change pas tant qu'elle y est" in vu["lignes"][0]["erreurs"][0]


def test_vider_une_case_et_la_version_du_modele(achats: TestClient) -> None:
    """P7-12 : une case vide ne change rien ; « Vider » vide ; « Version » refuse une fiche
    modifiée depuis le modèle ; une matière nouvelle a besoin de son rôle."""
    m = matiere(groupe_equiv="G1", couleur_code="7612")
    ref, version = lue(m).reference, lue(m).version
    contenu = classeur(
        [
            {
                "Référence": ref,
                "Code couleur": "",
                "Vider": "Groupe d'équivalence",
                "Version": version,
            }
        ]
    )
    vu = ok(apercu(achats, contenu))
    assert [c["champ"] for c in vu["lignes"][0]["changements"]] == ["groupe_equiv"]
    ok(appliquer(achats, contenu, vu["empreinte_apercu"]))
    assert (lue(m).groupe_equiv, lue(m).couleur_code) == ("", "7612")
    vu = ok(apercu(achats, contenu))
    assert vu["lignes"][0]["erreurs"][0].startswith(
        "Modifiée entre-temps par Achats depuis le modèle"
    )
    vu = ok(apercu(achats, classeur([{"Référence": ref, "Vider": "Rôle; Couleurs"}])))
    assert vu["lignes"][0]["erreurs"] == [
        "Vider : « Rôle » ne se vide pas",
        "Vider : colonne inconnue « Couleurs »",
    ]
    vu = ok(apercu(achats, classeur([{"Référence": unique("Sans role")}])))
    assert vu["lignes"][0]["erreurs"] == ["Une matière nouvelle a besoin de son rôle"]


def test_l_empreinte_refuse_une_fiche_changee_depuis_l_apercu(achats: TestClient) -> None:
    """Défaut A-03 : on écrit exactement ce qu'on a vu (409 sinon)."""
    m = matiere(prix="5")
    contenu = classeur([{"Référence": lue(m).reference, "Prix": 6}])
    vu = ok(apercu(achats, contenu))
    x = lue(m)
    corps = {"reference": x.reference, "role_id": x.role_id, "fournisseur_id": x.fournisseur_id}
    corps |= {"prix": 5.5, "version": x.version}
    ok(achats.put(f"/api/matieres/{m}", json=corps))
    r = appliquer(achats, contenu, vu["empreinte_apercu"])
    assert (r.status_code, r.json()["code"]) == (409, "conflit_version")
    assert float(lue(m).prix) == 5.5


def test_fichiers_refuses(achats: TestClient, alimentation: TestClient) -> None:
    """Hérité : X-REF-36, X-REF-37 ; la feuille « Matières » est lue, active ou non."""
    r = apercu(achats, b"pas un classeur")
    assert (r.status_code, r.json()["code"]) == (422, "tableur_invalide")
    r = apercu(achats, classeur([{"Nom": "x"}]))
    assert r.json()["detail"] == "La colonne « Référence » est obligatoire (télécharger le modèle)."
    m = matiere()
    vu = ok(apercu(achats, classeur([{"Référence": lue(m).reference}], active=False)))
    assert vu["inchangees"] == 1 and not vu["appliquable"]
    assert apercu(alimentation, classeur([{"Référence": "x"}])).status_code == 403
    with sessions()() as s:
        assert s.scalar(select(func.count()).where(Matiere.reference == "x")) == 0
