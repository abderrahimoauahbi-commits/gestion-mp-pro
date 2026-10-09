"""Le coût de revient des matières importées (routes 206–208) et le classeur des dossiers (205),
lot D7. Preuves héritées de `test_importations`, `test_importations_second_audit` et
`test_defauts_b` (domaine 8) : recherche sans accents, combinaisons, qualité, kilos, frais non
répartis qui suivent les filtres, statut sans casse, classeurs.
"""

from __future__ import annotations

import io
import json
from typing import Any

import pytest
from fabrique import *  # noqa: F403
from fabrique import client_de
from fabriques.importations import Import, a_la_main, fiche, frais, geste, ligne, ouvrir
from fabriques.receptions import menage, ok  # noqa: F401 — fixture du module
from fabriques.referentiels import poste_frais, qualite, unique, valeur_classification
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

pytestmark = pytest.mark.usefixtures("menage")
URL = "/api/analyse/revient"


@pytest.fixture(scope="module")
def monde(application: FastAPI) -> dict[str, Any]:
    """Trois fils importés (rouge 1 000 kg × 2, noir 500 kg × 4, jute 2 000 kg × 1, au cours
    10), un fret de 3 500 à la valeur, dossier clos ; une qualité 60 % rouge, 40 % noir."""
    achats = client_de(application, "achats")
    imp = Import.nouveau(cours="10")
    code_pes = unique("PES")[:12]
    pes = valeur_classification("composition", code_pes)
    imp.matiere("FRISÉ ROUGE", composition_id=pes)
    imp.matiere("FRISÉ NOIR", composition_id=pes)
    lignes = [("FRISÉ ROUGE", "AR", 1000, 2), ("FRISÉ NOIR", "AN", 500, 4), ("JUTE", "AJ", 2000, 1)]
    imp.recevoir(achats, "ANA-1", lignes)
    d = frais(achats, ouvrir(achats, imp.f("ANA-1")), "FRET", 3500)
    d = geste(achats, d, "cloturer", alertes_lues=True)
    rouge, noir = imp.matieres["FRISÉ ROUGE"], imp.matieres["FRISÉ NOIR"]
    q = qualite(recette={rouge: 60, noir: 40}, poids={"Poil": "1"})
    return {"imp": imp, "dossier": d, "qualite": q, "pes": code_pes}


def test_analyse_combinaisons_qualite_et_recherche(
    monde: dict[str, Any], achats: TestClient, alimentation: TestClient
) -> None:
    """Hérité : test_importations::test_analyse_combinaisons_qualite_et_recherche — « frise
    rouge » trouve le frisé rouge (sans accents, mots dans le désordre) ; le poids total = le
    détail ; des matières choisies ; une qualité et des matières se cumulent (la recette) ; le
    classeur ; l'alimentation ne lit ni l'analyse ni les dossiers."""
    imp: Import = monde["imp"]
    rouge = ligne(monde["dossier"], "IMP FRISÉ ROUGE")["reference"]
    suffixe = rouge.split("-", 1)[1].lower()
    r = ok(achats.get(URL, params={"q": f"{suffixe} rouge frise", "detail": True}))
    assert {x["valeur"] for x in r["lignes"]} == {rouge}
    assert r["ids_dossiers"] == {monde["dossier"]["numero"]: monde["dossier"]["id"]}
    assert r["total_poids"] == pytest.approx(sum(x["poids"] for x in r["detail_lignes"]))
    jute, noir = imp.matieres["JUTE"], imp.matieres["FRISÉ NOIR"]
    r = ok(achats.get(URL, params={"matieres": [jute, noir]}))
    assert len(r["lignes"]) == 2 and r["total_poids"] == pytest.approx(2500)
    r = ok(achats.get(URL, params={"matieres": [jute, noir], "qualite": monde["qualite"]}))
    assert [x["valeur"] for x in r["lignes"]] == [
        ligne(monde["dossier"], "IMP FRISÉ NOIR")["reference"]
    ]
    r = ok(achats.get(f"{URL}/qualites/{monde['qualite']}"))
    assert r["references_importees"] == 2
    classeur = load_workbook(
        io.BytesIO(achats.get("/api/export/revient", params={"q": "frise"}).content)
    )
    assert classeur.sheetnames == ["Par référence", "Lot par lot"]
    assert alimentation.get(URL).status_code == 403
    assert alimentation.get("/api/dossiers").status_code == 403


def test_les_combinaisons_se_reunissent(monde: dict[str, Any], achats: TestClient) -> None:
    """Une combinaison cumule ses critères, la sélection réunit ses combinaisons (M54)."""
    selection = json.dumps([{"composition": monde["pes"]}])
    r = ok(achats.get(URL, params={"selection": selection}))
    assert r["total_poids"] == pytest.approx(1500)
    refus = achats.get(URL, params={"selection": "[1"})
    assert refus.status_code == 422 and refus.json()["detail"].startswith("Sélection illisible")
    refus = achats.get(URL, params={"selection": json.dumps([{"teinte": "x"}])})
    assert refus.json()["detail"] == "Critère inconnu dans la sélection : teinte."
    refus = achats.get(URL, params={"axe": "teinte"})
    assert refus.status_code == 422 and refus.json()["detail"].startswith(
        "Axe inconnu : « teinte »"
    )


def test_cout_d_une_qualite_pondere_par_la_recette(
    monde: dict[str, Any], achats: TestClient
) -> None:
    """Hérité : test_importations::test_cout_d_une_qualite_pondere_par_la_recette — 1 kg de
    poil au m², 60 % de rouge et 40 % de noir : coût au m² = 0,6 × rouge + 0,4 × noir ;
    recette couverte ; Σ des parts = 100 ; qualité inconnue → 404."""
    r = ok(achats.get(f"{URL}/qualites/{monde['qualite']}"))
    cout = {x["valeur"]: x["revient_unite"] for x in r["lignes"]}
    rouge = ligne(monde["dossier"], "IMP FRISÉ ROUGE")["reference"]
    noir = ligne(monde["dossier"], "IMP FRISÉ NOIR")["reference"]
    attendu = 0.6 * cout[rouge] + 0.4 * cout[noir]
    assert r["cout_m2"] == pytest.approx(attendu) and r["couverture"] == pytest.approx(1)
    assert r["cout_kg_tapis"] == pytest.approx(attendu)
    assert sum(x["part_recette"] for x in r["lignes"]) == pytest.approx(100)
    # Fret 3 500 à la valeur sur 60 000 : le rouge (20 000 MAD, 1 000 kg) coûte 21,1667 MAD/kg.
    assert cout[rouge] == pytest.approx((20_000 + 3500 / 3) / 1000)
    refus = achats.get(f"{URL}/qualites/999999")
    assert refus.status_code == 404 and refus.json()["detail"] == "Qualité introuvable."


def test_l_analyse_additionne_des_kilos(achats: TestClient) -> None:
    """Hérité : test_importations_second_audit::test_l_analyse_additionne_des_kilos — par
    fournisseur : 200 kg (le fil et la bande), 30 MAD/kg ; les classeurs disent la même chose."""
    imp = Import.nouveau(cours="10")
    imp.matiere("BANDE", role="Ruban", unite_prix="ml", densite="0.05")
    manut = unique("MAN")[-16:].replace("-", "_")
    poste_frais(code=manut, cle="poids")
    imp.recevoir(achats, "A2-KG", [("FIL", "KF", 100, 2), ("BANDE", "KB", 100, 2)])
    d = frais(achats, ouvrir(achats, imp.f("A2-KG")), manut, 2000)
    nom = d["fournisseurs"][0]
    r = ok(achats.get(URL, params={"axe": "fournisseur", "fournisseur": nom, "detail": True}))
    (g,) = r["lignes"]
    assert g["poids"] == pytest.approx(200) and r["total_poids"] == pytest.approx(200)
    assert g["quantite"] == pytest.approx(200) and g["unite"] == "kg" and g["unites"] == ["kg"]
    assert g["revient_kg"] == pytest.approx(6000 / 200)
    feuille = _feuille(
        achats, "/api/export/revient", "Par fournisseur", axe="fournisseur", fournisseur=nom
    )
    assert _valeur(feuille, "Fournisseur", nom, "Poids (kg)") == pytest.approx(200)
    assert _valeur(feuille, "Fournisseur", nom, "Coût hors TVA / kg") == pytest.approx(30)
    bande = ligne(d, "IMP BANDE")["reference"]
    feuille = _feuille(achats, "/api/export/dossiers", "Marchandises", q=d["numero"])
    assert _valeur(feuille, "Référence", bande, "Quantité (kg)") == pytest.approx(100)
    assert _valeur(feuille, "Référence", bande, "Coût hors TVA / kg") == pytest.approx(30)


def _feuille(client: TestClient, url: str, nom: str, **params: Any) -> list[tuple[Any, ...]]:
    r = client.get(url, params=params)
    assert r.status_code == 200, r.text
    return list(load_workbook(io.BytesIO(r.content))[nom].iter_rows(values_only=True))


def _valeur(lignes: list[tuple[Any, ...]], colonne: str, cle: str, voulue: str) -> Any:
    """La cellule `voulue` de la ligne dont la colonne `colonne` vaut `cle`."""
    entetes = next(x for x in lignes if colonne in x)
    rang = next(x for x in lignes if x[entetes.index(colonne)] == cle)
    return rang[entetes.index(voulue)]


def test_les_frais_non_repartis_suivent_les_filtres(achats: TestClient) -> None:
    """Hérité : test_importations_second_audit::test_les_frais_non_repartis_suivent_les_filtres
    — un dossier vide est annoncé sans filtre, pas sous un filtre de fournisseur, et par une
    recherche de son numéro."""
    d = frais(achats, a_la_main(achats, date_arrivee="2019-04-01"), "DD", 1234)
    assert d["numero"] in {x["dossier"] for x in ok(achats.get(URL))["dossiers_non_repartis"]}
    r = ok(achats.get(URL, params={"fournisseur": "personne"}))
    assert d["numero"] not in {x["dossier"] for x in r["dossiers_non_repartis"]}
    r = ok(achats.get(URL, params={"q": d["numero"]}))
    assert d["numero"] in {x["dossier"] for x in r["dossiers_non_repartis"]}


def test_analyse_statut_insensible_a_la_casse(monde: dict[str, Any], achats: TestClient) -> None:
    """Hérité : test_defauts_b::test_analyse_statut_insensible_a_la_casse (B-09) — « clos » et
    « Clos » disent la même chose ; un statut inconnu → 422, plus jamais une analyse vide."""
    clos = ok(achats.get(URL, params={"statut": "Clos"}))
    assert ok(achats.get(URL, params={"statut": "clos"}))["total_revient_mad"] == pytest.approx(
        clos["total_revient_mad"]
    )
    assert monde["dossier"]["numero"] in clos["ids_dossiers"]
    r = achats.get(URL, params={"statut": "fermé"})
    assert r.status_code == 422 and "Statut de dossier inconnu" in r.json()["detail"]


def test_export_dossiers_titre_clos(achats: TestClient, direction: TestClient) -> None:
    """Hérité : test_importations::test_export_des_dossiers et
    test_defauts_b::test_export_dossiers_titre_clos (B-09) — quatre feuilles ; « — clos » ;
    statut en clair ; répartition « figée (clé d'avant son changement) » ; un frais passé à un
    autre poste perd sa répartition figée."""
    imp = Import.nouveau(cours="10")
    code = unique("DBMAN")[-16:].replace("-", "_")
    poste_frais(code=code, cle="valeur")
    imp.recevoir(achats, "DB-FIG", [("FIL", "DBF-1", 900, 2), ("FIL", "DBF-2", 100, 2)])
    d = frais(achats, ouvrir(achats, imp.f("DB-FIG")), code, 1000)
    cid = d["frais"][0]["id"]
    geste(achats, d, "cloturer", alertes_lues=True)
    p = next(x for x in ok(direction.get("/api/postes-frais")) if x["code"] == code)
    champs = {
        k: p[k]
        for k in ("libelle", "recuperable", "nature", "entre_valeur_douane", "ordre", "version")
    }
    ok(direction.put(f"/api/postes-frais/{p['id']}", json={**champs, "cle": "poids"}))
    r = achats.get("/api/export/dossiers", params={"statut": "clos"})
    classeur = load_workbook(io.BytesIO(r.content))
    assert classeur.sheetnames == ["Dossiers", "Marchandises", "Frais", "Échéancier"]
    assert str(classeur["Dossiers"].cell(1, 1).value).startswith("Dossiers d'importation — clos")
    lignes = list(classeur["Frais"].iter_rows(values_only=True))
    assert (
        _valeur(lignes, "Dossier", d["numero"], "Répartition")
        == "figée (clé d'avant son changement)"
    )
    assert (
        _valeur(
            list(classeur["Dossiers"].iter_rows(values_only=True)), "Dossier", d["numero"], "Statut"
        )
        == "Clos"
    )
    titre = load_workbook(
        io.BytesIO(achats.get("/api/export/dossiers", params={"statut": "ouvert"}).content)
    )
    assert str(titre["Dossiers"].cell(1, 1).value).startswith("Dossiers d'importation — ouverts")
    assert achats.get("/api/export/dossiers", params={"statut": "fermé"}).status_code == 422
    d = geste(achats, fiche(achats, d["id"]), "rouvrir")
    corps = {"version": d["frais"][0]["version"], "poste_id": _poste_id("TRANSIT"), "montant": 1000}
    d = ok(achats.put(f"/api/dossiers/{d['id']}/frais/{cid}", json=corps))
    assert d["frais"][0]["parts"] == []


def _poste_id(code: str) -> int:
    from fabriques.importations import poste

    return poste(code)
