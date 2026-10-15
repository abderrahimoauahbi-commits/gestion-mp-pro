"""L'échéancier par l'API (lot D8 ; domaine 7 de ESSAIS.md, routes 228–230, P8-21, B-01, I20).

Les factures viennent de vraies réceptions (lot D4) et les frais de dossiers d'importation :
posés en base dans une usine isolée pour les totaux exacts, ou ouverts par les vrais services du
lot D7 sur la base partagée (lignes retrouvées par leur pièce).
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

import pytest
from fabrique import *  # noqa: F403
from fabriques import importations as imp
from fabriques.pilotage import Usine, fixture_usine, jour, ok  # noqa: F401
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from gestion_mp_pro.calcul import suivi
from gestion_mp_pro.modeles import DossierFrais
from gestion_mp_pro.modeles.enums import StatutDossier


def camion(u: Usine, f: int, kg: float, prix: float, **facture: Any) -> dict[str, Any]:
    """Une réception validée hors commande, sa facture saisie (en dirhams)."""
    m = u.matiere(fournisseur_id=f)
    corps: dict[str, Any] = {
        "date": jour().isoformat(), "fournisseur_id": f,
        "lignes": [{"matiere_id": m, "lot_numero": f"L-{facture['numero']}", "quantite_pesee": kg,
                    "magasin_id": u.courant, "prix": prix}],
        "factures": {"MAD": facture},
    }  # fmt: skip
    r = ok(u.achats.post("/api/receptions", json=corps), 201)[0]
    return ok(u.achats.post(f"/api/receptions/{r['id']}/valider", json={"version": r["version"]}))  # type: ignore[no-any-return]


def echeancier(u: Usine, **params: Any) -> dict[str, Any]:
    """L'échéancier lu par les achats."""
    return u.get(u.achats, "/api/echeancier", **params)  # type: ignore[no-any-return]


def piece(liste: list[dict[str, Any]], numero: str) -> dict[str, Any]:
    """La ligne d'une pièce."""
    return next(x for x in liste if x["piece"] == numero)


def payer(
    client: TestClient, x: dict[str, Any], le: str | None, code: int = 200, **corps: Any
) -> Any:
    """`PUT /echeancier/paiement` avec la version de la ligne."""
    genre = "facture" if x["objet"] == "facture" else "frais"
    corps = {"genre": genre, "id": x["id"], "payee_le": le, "version": x["version"], **corps}
    return ok(client.put("/api/echeancier/paiement", json=corps), code)


def test_l_echeancier_suit_les_paiements(usine: Usine) -> None:
    """Hérité : test_commandes_receptions::test_l_echeancier_suit_les_paiements (partie
    échéancier) — sans montant : 200 kg × 10 = 2 000 estimé, en retard ; payée demain → 422 ;
    payée aujourd'hui : sort des lignes, l'échu baisse de 2 000, elle est dans les réglées ;
    l'annulation (direction, motif) la rend due ; les achats ne l'annulent pas (I20)."""
    u = usine
    f = u.fournisseur("ECH FILATURE")
    camion(u, f, 200, 10, numero="CR-F1", date=jour(-40).isoformat(),
           echeance=jour(-10).isoformat())  # fmt: skip
    e = echeancier(u)
    x = piece(e["lignes"], "CR-F1")
    assert (x["montant"], x["estime"], x["en_retard"]) == (2000, True, True)
    assert e["echu_mad"] == 2000 and e["echues"] == 1
    refus = payer(u.achats, x, jour(1).isoformat(), 422)
    assert refus["detail"] == "Un règlement ne se date pas dans le futur : c'est un fait constaté."
    regle = payer(u.achats, x, jour().isoformat())
    assert regle["etat"] == "reglee" and regle["payee_le"] == jour().isoformat()
    e = echeancier(u)
    assert "CR-F1" not in {y["piece"] for y in e["lignes"]} and e["echu_mad"] == 0
    assert piece(e["reglees"], "CR-F1")["payee_le"] == jour().isoformat()
    payer(u.direction, x, None, 409, motif="réglée par erreur")  # version périmée : 409
    regle = piece(echeancier(u)["reglees"], "CR-F1")
    refus = payer(u.achats, regle, None, 403, motif="réglée par erreur")
    assert "ne s'annule que par la direction" in refus["detail"]
    assert payer(u.direction, regle, None, 422)["code"] == "motif_requis"
    assert payer(u.direction, regle, None, motif="réglée par erreur")["etat"] == "a_payer"
    assert "CR-F1" in {y["piece"] for y in echeancier(u)["lignes"]}


def test_une_facture_sans_date_est_a_dater_et_s_exporte(usine: Usine) -> None:
    """Hérité : test_commandes_receptions::test_une_facture_sans_date_est_a_dater_et_s_exporte ;
    test_direction::test_recette_l_echeancier_liste_ce_qui_est_a_dater — 30 × 10 = 300, « à
    dater », estimé ; la feuille « À dater » du classeur la porte (pièce en 3e colonne)."""
    u = usine
    camion(u, u.fournisseur(), 30, 10, numero="CR-F-AD")
    e = echeancier(u)
    x = piece(e["a_dater"], "CR-F-AD")
    assert x["montant"] == 300 and x["estime"] and e["a_dater_mad"] == 300
    r = u.achats.get("/api/export/echeancier")
    classeur = load_workbook(BytesIO(r.content))
    assert classeur.sheetnames == ["Échéances", "À dater", "Par mois", "Réglées", "Écartées"]
    assert "CR-F-AD" in {ligne[2] for ligne in classeur["À dater"].iter_rows(values_only=True)}


def test_la_facture_d_une_reception_annulee_n_est_plus_due(usine: Usine) -> None:
    """Hérité : test_commandes_receptions::test_la_facture_d_une_reception_annulee_n_est_plus_due
    — la réception annulée, sa facture sort des lignes et entre dans les écartées."""
    u = usine
    r = camion(u, u.fournisseur(), 40, 10, numero="CR-F-AN", date=jour().isoformat(),
               montant=400, echeance=jour(30).isoformat())  # fmt: skip
    assert "CR-F-AN" in {x["piece"] for x in echeancier(u)["lignes"]}
    ok(u.achats.post(f"/api/receptions/{r['id']}/annuler",
                     json={"version": r["version"], "motif": "saisie en double"}))  # fmt: skip
    e = echeancier(u)
    assert "CR-F-AN" not in {x["piece"] for x in e["lignes"]}
    assert piece(e["ecartees"], "CR-F-AN")["raison"] == "réception annulée : plus due"


def test_filtres_et_indicateurs_suivent_l_ecran(usine: Usine) -> None:
    """P8-21 : recherche et nature au serveur, indicateurs qui les suivent ; la période ne
    cache pas l'échu ; par mois sur les lignes filtrées ; dans les 30 jours."""
    u = usine
    f = u.fournisseur("FILTRE FILATURE")
    camion(u, f, 10, 10, numero="FX-1", date=jour(-20).isoformat(), montant=100,
           echeance=jour(-5).isoformat())  # fmt: skip
    camion(u, f, 10, 10, numero="FX-2", date=jour().isoformat(), montant=200,
           echeance=jour(10).isoformat())  # fmt: skip
    d = u.dossier(frais={"TRANSIT": 300})
    ch = u.s.query(DossierFrais).filter_by(dossier_id=d.id).one()
    ch.echeance, ch.prestataire = jour(40), "TRANSIT DU PORT"
    tout = echeancier(u)
    assert tout["total_mad"] == 600 and tout["echu_mad"] == 100
    assert tout["dans_30_jours_mad"] == 200
    assert tout["par_mois"] == {"2026-10": 300, "2026-11": 300}
    frais = echeancier(u, nature="frais")
    assert [x["tiers"] for x in frais["lignes"]] == ["TRANSIT DU PORT"]
    assert frais["echu_mad"] == 0 and frais["lignes"][0]["dossier_id"] == d.id
    fx2 = echeancier(u, q="fx 2")
    assert [x["piece"] for x in fx2["lignes"]] == ["FX-2"] and fx2["lignes"][0]["reception_id"]
    futur = echeancier(u, du=jour(1).isoformat())
    assert futur["echu_mad"] == 100 and {x["piece"] for x in futur["lignes"]} == {"FX-2", d.numero}
    assert u.achats.get("/api/echeancier", params={"du": jour(2), "au": jour(1)}).status_code == 422
    assert u.alimentation.get("/api/echeancier").status_code == 403


def test_les_reglees_plafonnees_sans_periode_toutes_avec(
    usine: Usine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Sans période : les `REGLEES_AU_PLUS` règlements les plus récents parmi les lignes gardées
    (ici 2), et la réponse le dit ; une période `du`/`au` (date de règlement, bornes incluses)
    les rend toutes, factures et frais, sans plafond ; le classeur suit et le dit ; régler une
    pièce au-delà du plafond rend quand même sa ligne ; l'alimentation : 403."""
    u = usine
    monkeypatch.setattr(suivi, "REGLEES_AU_PLUS", 1)
    f = u.fournisseur("REGLEE FILATURE")
    for n in ("RG-1", "RG-2"):
        camion(u, f, 10, 10, numero=n, date=jour(-30).isoformat(), montant=100,
               echeance=jour(-5).isoformat())  # fmt: skip
    d = u.dossier(frais={"TRANSIT": 300})
    ch = u.s.query(DossierFrais).filter_by(dossier_id=d.id).one()
    ch.payee_le = jour(-10)
    e = echeancier(u)
    payer(u.achats, piece(e["lignes"], "RG-2"), jour(-1).isoformat())
    ancienne = payer(u.achats, piece(e["lignes"], "RG-1"), jour(-20).isoformat())
    assert ancienne["etat"] == "reglee" and ancienne["payee_le"] == jour(-20).isoformat()
    monkeypatch.setattr(suivi, "REGLEES_AU_PLUS", 2)
    sans = echeancier(u)
    assert [x["piece"] for x in sans["reglees"]] == ["RG-2", d.numero]
    assert (sans["reglees_total"], sans["reglees_plafond"]) == (3, 2)
    avec = echeancier(u, du=jour(-20).isoformat(), au=jour(-10).isoformat())
    assert [x["piece"] for x in avec["reglees"]] == [d.numero, "RG-1"]
    assert (avec["reglees_total"], avec["reglees_plafond"]) == (2, None)
    assert len(echeancier(u, du=jour(-20).isoformat())["reglees"]) == 3
    factures = echeancier(u, nature="facture")
    assert [x["piece"] for x in factures["reglees"]] == ["RG-2", "RG-1"]
    assert (factures["reglees_total"], factures["reglees_plafond"]) == (2, 2)
    assert u.alimentation.get("/api/echeancier", params={"du": jour(-20)}).status_code == 403
    r = u.achats.get("/api/export/echeancier")
    lignes = list(load_workbook(BytesIO(r.content))["Réglées"].iter_rows(values_only=True))
    assert lignes[1][0] == (
        "Réglées : les 2 dernières sur 3. Pour toutes les voir, poser une période (du / au) : "
        "elle s'applique à la date de règlement."
    )
    assert "RG-1" not in {x[4] for x in lignes}
    r = u.achats.get("/api/export/echeancier", params={"du": jour(-20), "au": jour(-10)})
    lignes = list(load_workbook(BytesIO(r.content))["Réglées"].iter_rows(values_only=True))
    attendu = f"du {jour(-20):%d/%m/%Y} au {jour(-10):%d/%m/%Y} (2, date de règlement)."
    assert lignes[1][0].startswith("Filtres — du :")
    assert lignes[2][0] == f"Réglées : toutes les pièces réglées {attendu}"
    assert {x[4] for x in lignes} >= {"RG-1", d.numero}


def test_un_frais_se_regle_au_journal_de_son_dossier(usine: Usine) -> None:
    """Un frais se règle aux mêmes règles qu'une facture (ni futur, ni avant sa pièce) ; le
    règlement s'écrit au journal de son dossier ; la version périmée est refusée."""
    u = usine
    d = u.dossier(frais={"TRANSIT": 300})
    ch = u.s.query(DossierFrais).filter_by(dossier_id=d.id).one()
    ch.echeance, ch.date_piece, ch.piece = jour(5), jour(-3), "T-77"
    x = piece(echeancier(u)["lignes"], f"{d.numero} · T-77")
    refus = payer(u.achats, x, jour(-4).isoformat(), 422)
    assert "précède la pièce" in refus["detail"]
    regle = payer(u.achats, x, jour().isoformat())
    assert regle["etat"] == "reglee"
    assert payer(u.achats, x, jour().isoformat(), 409)["code"] == "conflit_version"
    journal = u.get(u.direction, f"/api/historique/dossier/{d.id}")
    assert any(y["action"] == "règlement" for y in journal)


def test_les_frais_dus_d_un_dossier_annule_ne_le_sont_plus(usine: Usine) -> None:
    """Lot D7 (B-01) : un dossier annulé n'a plus de frais dus (écartés) ; un frais réglé reste
    réglé."""
    u = usine
    d = u.dossier(frais={"TRANSIT": 300, "FRET": 500})
    t, fret = sorted(u.s.query(DossierFrais).filter_by(dossier_id=d.id), key=lambda c: c.montant)
    t.echeance, fret.echeance, fret.payee_le = jour(10), jour(5), jour(-1)
    d.statut, d.annule_le = StatutDossier.annule, imp_maintenant()
    d.annulation_motif = "doublon"
    e = echeancier(u)
    assert e["lignes"] == []
    assert [x["raison"] for x in e["ecartees"]] == ["dossier annulé : plus dû"]
    assert [x["montant"] for x in e["reglees"]] == [500]


def imp_maintenant() -> Any:
    """L'instant de l'usine (annulation posée en base)."""
    from gestion_mp_pro.services import horloge

    return horloge.maintenant()


def test_echeanciers_regle_estime_et_clos_a_dater(achats: TestClient) -> None:
    """Hérité : test_importations_second_audit::test_echeanciers_regle_estime_et_clos_a_dater
    (partie « échéancier général », laissée par D7) — un frais non réglé d'un dossier **clos**,
    sans échéance, reste « à dater » : clore fige le coût, pas le paiement du transitaire."""
    i = imp.Import.nouveau(code="USD")
    i.recevoir(achats, "A2-ECH", [("FIL", "A2F-10", 50, 2.0)])
    d = imp.ouvrir(achats, i.f("A2-ECH"), local_confirme=True)
    d = imp.frais(achats, d, "TRANSIT", 700, prestataire="TRANSIT OUBLIÉ D8")
    imp.geste(achats, d, "cloturer", alertes_lues=True)
    e = ok(achats.get("/api/echeancier", params={"q": "transit oublie d8"}))
    assert [x["montant"] for x in e["a_dater"] if x["tiers"] == "TRANSIT OUBLIÉ D8"] == [700]
    assert e["a_dater_mad"] == pytest.approx(700)
