"""Valorisation et Cockpit par l'API (lot D8 ; domaine 10 de ESSAIS.md, défauts E-06, H-07).

Une usine isolée (`fabriques.pilotage.usine`) : les totaux sont exacts. Deux pages, un chiffre
(V1, E-06) : la valeur du Cockpit est celle de la Valorisation. Le fil monté sur les métiers est
du stock (décision de Mohamed, U8), à toute date : l'historique des montages le rejoue.
"""

from __future__ import annotations

from datetime import datetime, time
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from fabrique import *  # noqa: F403
from fabriques.pilotage import Usine, fixture_usine, jour, ok  # noqa: F401
from openpyxl import load_workbook
from sqlalchemy import text

from gestion_mp_pro.modeles import Devise, Matiere, MontageEvenement
from gestion_mp_pro.modeles.enums import TypeMouvement
from gestion_mp_pro.services import horloge


def valeur(u: Usine, reference: str, **params: Any) -> dict[str, Any]:
    """La ligne de valorisation d'une matière."""
    lignes = u.get(u.achats, "/api/valorisation", **params)["lignes"]
    return next(x for x in lignes if x["reference"] == reference)  # type: ignore[no-any-return]


def test_cockpit_egale_total_valorisation(usine: Usine) -> None:
    """Hérité : test_defauts_e::test_cockpit_egale_total_valorisation (E-06) — un coût nul est un
    coût (pas le catalogue) ; un stock négatif ne vaut rien sur les deux pages ; le Cockpit, son
    indicateur et la Valorisation disent le même total."""
    u = usine
    gratuit = u.matiere("E6 GRATUIT", prix="0")
    u.ouvrir(gratuit, 50, 0)
    fiche = u.s.get(Matiere, gratuit)
    assert fiche is not None
    fiche.prix = Decimal(10)
    u.ouvrir(u.matiere("E6 PAYANT"), 100, 12)
    negatif = u.ouvrir(u.matiere("E6 NEGATIF"), 10, 7)
    u.s.execute(
        text(
            "INSERT INTO mouvement (date_effet, type, matiere_id, lot_id, magasin_id, quantite,"
            " document, auteur_id) VALUES (:j, 'AJUSTEMENT', :m, :l, :g, -30, 'reprise', :a)"
        ),
        {"j": jour(), "m": negatif.matiere_id, "l": negatif.id, "g": u.courant, "a": u.auteur},
    )
    v = u.get(u.achats, "/api/valorisation")
    assert valeur(u, "E6 GRATUIT")["cmup_mad"] == 0
    assert v["total_mad"] == pytest.approx(1200) and v["nb_en_stock"] == 2
    c = u.get(u.direction, "/api/cockpit")
    assert c["stock"]["valeur_mad"] == pytest.approx(v["total_mad"])
    kpi = {k["cle"]: k for k in c["kpi"]}
    assert kpi["valeur"]["valeur"] == pytest.approx(1200)


def test_api_annuler_une_reception_rend_le_cmup_d_avant(usine: Usine) -> None:
    """Hérité : test_valorisation::test_api_annuler_une_reception_rend_le_cmup_d_avant —
    125,7 kg à 26,89 puis 1 000 kg à 200 : CMUP mêlé ; la réception annulée, 26,89."""
    u = usine
    f = u.fournisseur("VALO FILATURE")
    m = u.matiere("VALO PLASTIQUE-50", fournisseur_id=f, prix="26.89")
    u.ouvrir(m, 125.7, 26.89, le=jour(-1))
    r = ok(u.achats.post("/api/receptions", json={
        "date": jour().isoformat(), "fournisseur_id": f, "lignes": [
            {"matiere_id": m, "lot_numero": "VP-2", "quantite_pesee": 1000,
             "magasin_id": u.courant, "prix": 200}]}), 201)[0]  # fmt: skip
    r = ok(u.achats.post(f"/api/receptions/{r['id']}/valider", json={"version": r["version"]}))
    attendu = (125.7 * 26.89 + 200_000) / 1125.7
    assert valeur(u, "VALO PLASTIQUE-50")["cmup_mad"] == pytest.approx(attendu)
    ok(u.achats.post(f"/api/receptions/{r['id']}/annuler",
                     json={"version": r["version"], "motif": "erreur de saisie"}))  # fmt: skip
    v = valeur(u, "VALO PLASTIQUE-50")
    assert v["cmup_mad"] == pytest.approx(26.89)
    assert v["valeur_mad"] == pytest.approx(125.7 * 26.89)


def test_api_retour_fournisseur_au_cout_du_lot(usine: Usine) -> None:
    """Hérité : test_valorisation::test_api_retour_fournisseur_au_cout_du_lot — A à 10, B à 50,
    B renvoyé : stock 100, CMUP 10."""
    u = usine
    m = u.matiere("VALO RETOUR")
    u.ouvrir(m, 100, 10, le=jour(-1))
    b = u.ouvrir(m, 100, 50, le=jour(-1))
    u.ecrire(b, -100, type_=TypeMouvement.RETOUR_FOURNISSEUR)
    v = valeur(u, "VALO RETOUR")
    assert v["stock"] == pytest.approx(100) and v["cmup_mad"] == pytest.approx(10)


def test_api_le_cours_d_une_entree_est_fige_a_sa_validation(usine: Usine) -> None:
    """Hérité : test_valorisation::test_api_le_cours_d_une_entree_est_fige_a_sa_validation —
    2 USD au cours du jour (10) : 20 MAD ; le cours change, ce que l'entrée a coûté non."""
    u = usine
    m = u.matiere("VALO EN DOLLARS", devise="USD", prix="2")
    x = u.lot(m)
    u.ecrire(x, 100, prix=2, devise="USD")
    assert valeur(u, "VALO EN DOLLARS")["cmup_mad"] == pytest.approx(20)
    usd = u.s.get(Devise, "USD")
    assert usd is not None
    usd.cours = Decimal(15)
    assert valeur(u, "VALO EN DOLLARS")["cmup_mad"] == pytest.approx(20)


def test_api_une_date_passee_rejoue_stock_et_valeur(usine: Usine) -> None:
    """Hérité : test_valorisation::test_api_une_date_passee_rejoue_stock_et_valeur — au J−5 :
    100 kg, 1 000 MAD ; aujourd'hui 150 kg ; le stock passé dit 100."""
    u = usine
    m = u.matiere("VALO HISTOIRE")
    u.ouvrir(m, 100, 10, le=jour(-10))
    u.ouvrir(m, 50, 40)
    passe = jour(-5).isoformat()
    v = valeur(u, "VALO HISTOIRE", date=passe)
    assert v["stock"] == pytest.approx(100) and v["valeur_mad"] == pytest.approx(1000)
    assert valeur(u, "VALO HISTOIRE")["stock"] == pytest.approx(150)
    stock = u.get(u.achats, "/api/stock", date=passe)["lignes"]
    assert next(x for x in stock if x["reference"] == "VALO HISTOIRE")["stock_net"] == 100
    assert u.get(u.achats, "/api/valorisation", date=passe)["date_reference"] == passe


def test_le_stock_est_en_magasin_ou_en_machine(usine: Usine) -> None:
    """Hérité : test_audit2_atelier::test_le_stock_est_en_magasin_ou_en_machine (moitié
    valorisation, laissée par D2) — 100 kg entrés à 10, 20 kg sortis et montés (1 m × 10
    bobines/m × 2 kg) : magasins 80, métiers 20, 100 kg et 1 000 MAD sur la Valorisation **et**
    le Cockpit (U8)."""
    u = usine
    m = u.matiere("MACHINE FIL", poids_bobine=Decimal(2))
    x = u.ouvrir(m, 100, 10, le=jour(-3))
    u.sortir(x, 20, le=jour(-3))
    u.metier(fils=((1, m, 1),))
    v = valeur(u, "MACHINE FIL")
    assert (v["stock_magasin"], v["stock_machine"], v["stock"]) == (80, 20, 100)
    assert v["valeur_mad"] == pytest.approx(1000) and v["valeur_machine_mad"] == pytest.approx(200)
    total = u.get(u.achats, "/api/valorisation")
    assert total["total_kg_machine"] == 20 and total["total_mad"] == pytest.approx(1000)
    c = u.get(u.direction, "/api/cockpit")
    assert c["stock"]["valeur_mad"] == pytest.approx(1000)
    assert c["stock"]["valeur_machine_mad"] == pytest.approx(200)
    assert "dont 200 MAD sur les métiers" in c["kpi"][0]["detail"].replace(" ", " ").replace(
        "\xa0", " "
    )
    # Un montage repris sans histoire compte aussi à une date passée.
    assert valeur(u, "MACHINE FIL", date=jour(-1).isoformat())["stock_machine"] == 20


def test_a_une_date_passee_le_montage_de_ce_soir_la(usine: Usine) -> None:
    """Le fil monté à une date passée est celui de ce soir-là : un montage fait aujourd'hui
    (événement de l'historique) ne compte pas hier ; défait, l'étage d'hier revient."""
    u = usine
    m = u.matiere("HISTOIRE FIL", poids_bobine=Decimal(2))
    x = u.ouvrir(m, 100, 10, le=jour(-3))
    u.sortir(x, 40, le=jour(-3))
    metier = u.metier(fils=((1, m, 2),))
    ligne = {"etage": 1, "matiere_id": m, "reference": "HISTOIRE FIL", "metres": 1.0,
             "pourcentage": 100.0, "commentaire": "", "kg": 20.0}  # fmt: skip
    deux = {**ligne, "metres": 2.0, "kg": 40.0}
    u.s.add(MontageEvenement(
        metier_id=metier.id, etage=1, avant=[ligne], apres=[deux],
        quand=datetime.combine(jour(), time(9), tzinfo=horloge.FUSEAU),
    ))  # fmt: skip
    u.s.flush()
    assert valeur(u, "HISTOIRE FIL")["stock_machine"] == 40
    assert valeur(u, "HISTOIRE FIL", date=jour(-1).isoformat())["stock_machine"] == 20


def test_le_cockpit_compte_ses_alertes_et_dit_ses_formules(usine: Usine) -> None:
    """Hérité : test_direction::test_le_cockpit_compte_toutes_ses_alertes — chaque alerte dit
    son statut de couverture, son action, son titrage ; six indicateurs, chacun sa formule ;
    « N en rupture ou critiques » : un seul compte, au Cockpit et dans le menu (H-07)."""
    u = usine
    from fabriques import referentiels as ref

    m = u.matiere("ROUGE 2900")
    ref.qualite(recette={m: 100}, poids={"Poil": "1"}, base_mensuelle="1000", s=u.s)
    u.ouvrir(u.matiere("BLEU 2900"), 10, 10)
    c = u.get(u.direction, "/api/cockpit")
    assert c["nb_alertes"] >= len(c["alertes"]) >= 1
    alerte = next(x for x in c["alertes"] if x["reference"] == "ROUGE 2900")
    assert {"statut_couverture", "action", "titrage"} <= set(alerte)
    assert alerte["matiere_id"] == m
    assert [k["cle"] for k in c["kpi"]] == [
        "valeur", "couverture", "a_commander", "commande", "rotation", "controles",
    ]  # fmt: skip
    assert all(k["formule"] and k["lien"].startswith("/") for k in c["kpi"])
    assert c["alerte_rouge"] == c["stock"]["ruptures"] + c["stock"]["critiques"] >= 1
    for client, voit in ((u.achats, True), (u.alimentation, True), (u.administrateur, False)):
        compteurs = u.get(client, "/api/contexte")["compteurs"]
        assert (compteurs.get("alerte_rouge") == c["alerte_rouge"]) is voit
    assert (
        u.get(u.achats, "/api/contexte")["compteurs"]["anomalies"] == (c["controles"]["a_corriger"])
    )
    assert "anomalies" not in u.get(u.alimentation, "/api/contexte")["compteurs"]


def test_le_cockpit_d_une_date_passee_repond_et_borne_ses_dates(usine: Usine) -> None:
    """Hérité : test_audit2_pilotage::test_le_cockpit_d_une_date_passee_repond ;
    test_exports::test_une_date_hors_periode_est_refusee (422 en PRO) ; le pilotage est
    réservé à la direction et aux achats."""
    u = usine
    u.ouvrir(u.matiere(), 10, 10, le=jour(-40))
    passe = jour(-30).isoformat()
    assert u.get(u.direction, "/api/cockpit", date=passe)["date_reference"] == passe
    for date in ("1990-01-01", jour(1).isoformat()):
        for url in ("/api/cockpit", "/api/valorisation", "/api/controles", "/api/opportunites"):
            assert u.direction.get(url, params={"date": date}).status_code == 422, (url, date)
    for client in (u.alimentation, u.administrateur):
        assert client.get("/api/cockpit").status_code == 403
        assert client.get("/api/valorisation").status_code == 403


def test_la_valorisation_s_exporte_a_une_fin_de_mois(usine: Usine) -> None:
    """Hérité : test_valeur_et_dates::test_la_valorisation_s_exporte_a_une_fin_de_mois —
    « valorisation-2026-09-30.xlsx », titre « au 30/09/2026 »."""
    u = usine
    u.ouvrir(u.matiere("FIN DE MOIS"), 10, 10, le=jour(-30))
    r = u.achats.get("/api/export/valorisation", params={"date": "2026-09-30"})
    assert r.status_code == 200
    assert "valorisation-2026-09-30.xlsx" in r.headers["content-disposition"]
    feuille = load_workbook(BytesIO(r.content))["Valorisation"]
    assert "au 30/09/2026" in str(feuille["A1"].value)
    lignes = [x for x in feuille.iter_rows(values_only=True) if x and x[0] == "FIN DE MOIS"]
    assert lignes and lignes[0][5] == 10


VUES = (
    "/api/cockpit", "/api/valorisation", "/api/opportunites", "/api/controles",
    "/api/controles/C07", "/api/prix/historique", "/api/prix/matrice",
    "/api/fournisseurs/performance", "/api/echeancier", "/api/indicateurs/historique",
    "/api/bilan?mois=2026-10",
)  # fmt: skip


def test_les_vues_de_pilotage_repondent(usine: Usine) -> None:
    """Hérité : test_parcours::test_les_vues_de_pilotage_repondent — chaque vue répond à la
    direction et aux achats, jamais à l'alimentation ni à l'administrateur."""
    u = usine
    x = u.ouvrir(u.matiere(), 100, 10, le=jour(-5))
    u.sortir(x, 10)
    for url in VUES:
        for client in (u.direction, u.achats):
            assert client.get(url).status_code == 200, url
        for client in (u.alimentation, u.administrateur):
            assert client.get(url).status_code == 403, url
