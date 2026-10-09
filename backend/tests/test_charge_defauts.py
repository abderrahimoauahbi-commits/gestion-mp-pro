"""Défauts de PRO trouvés par l'épreuve de charge (lot V4, PLAN-DE-RECETTE-PRO § 10, A9).

Chaque essai échoue sans sa correction : il mesure la **forme** du travail (lectures de la
feuille, requêtes, calculs refaits), jamais une durée, pour ne pas dépendre de la machine.
Les durées elles-mêmes sont dans `essais/resultat-charge.txt` (`essais/charge.py`).
"""

from __future__ import annotations

from typing import Any

import pytest
from openpyxl.worksheet.worksheet import Worksheet

from gestion_mp_pro.exports.classeur import Colonne, Feuille, classeur


def test_un_classeur_ne_relit_pas_sa_feuille_a_chaque_ligne(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """V4-1 : `_ajouter` relisait la dernière ligne par `ws[ws.max_row]`, qui parcourt toute
    la feuille : le journal du stock (10 411 lignes) s'exportait en 40 s au volume réel. Le
    nombre de lectures de `max_row` ne dépend plus du nombre de lignes."""
    lectures = 0
    vrai = Worksheet.max_row.fget  # type: ignore[attr-defined]

    def compter(ws: Worksheet) -> Any:
        nonlocal lectures
        lectures += 1
        return vrai(ws)

    monkeypatch.setattr(Worksheet, "max_row", property(compter))
    colonnes = (Colonne("N", lambda x: x, "kg"), Colonne("Texte", lambda x: f"={x}"))
    contenu = classeur([Feuille("Essai", colonnes, list(range(2_000)), ["Titre"])], True)
    assert contenu[:2] == b"PK"
    assert lectures < 20  # noqa: PLR2004
