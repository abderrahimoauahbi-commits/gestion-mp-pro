"""Le taux d'un frais de dossier (lot TAUX-FRAIS, révision 0006) : gardé pour mémoire avec le
montant saisi, **jamais** utilisé pour recalculer ce montant — ni à l'écriture (la douane arrondit),
ni à la lecture (même quand le cours payé d'une marchandise change ensuite). Bornes : plus de 0,
au plus 100 ; un frais sans taux reste valable. Le classeur des dossiers dit le taux.
"""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

import pytest
from fabrique import *  # noqa: F403
from fabriques.importations import Import, fiche, frais, ouvrir
from fabriques.receptions import ok
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from gestion_mp_pro.db import fabrique as sessions
from gestion_mp_pro.modeles import DossierFrais, Journal


def _dossier(achats: TestClient, nom: str) -> dict[str, Any]:
    """Un dossier de deux fils (2 000 et 4 000 en devise) au cours payé de 10 : 60 000 MAD."""
    imp = Import.nouveau(cours="10")
    imp.recevoir(achats, nom, [("ROUGE", f"{nom}-R", 1000, 2.0), ("NOIR", f"{nom}-N", 1000, 4.0)])
    d = ouvrir(achats, imp.f(nom))
    return ok(  # type: ignore[no-any-return]
        achats.put(f"/api/dossiers/{d['id']}/cours", json={"version": d["version"], "cours": 10})
    )


def _le_frais(d: dict[str, Any], code: str) -> dict[str, Any]:
    """Le frais du poste `code` dans la fiche."""
    return next(x for x in d["frais"] if x["poste"] == code)  # type: ignore[no-any-return]


def _corriger(achats: TestClient, d: dict[str, Any], x: dict[str, Any], **champs: Any) -> Any:
    """`PUT` du frais entier (comme l'écran), avec les champs changés ; la réponse brute."""
    corps = {
        k: x[k]
        for k in ("poste_id", "montant", "taux", "ligne_id", "piece", "prestataire", "commentaire")
    }
    corps.update(version=x["version"], **champs)
    return achats.put(f"/api/dossiers/{d['id']}/frais/{x['id']}", json=corps)


def test_taux_et_montant_enregistres_tels_quels(achats: TestClient) -> None:
    """TVA à 20 % d'une base de 60 000 : 12 000 calculés, la douane en a pris 12 001,37 —
    le montant reçu est gardé tel quel, le taux aussi ; le journal dit les deux."""
    d = frais(achats, _dossier(achats, "TX-TVA"), "TVA", 12_001.37, taux=20)
    tva = _le_frais(d, "TVA")
    assert (tva["montant"], tva["taux"]) == (12_001.37, 20)
    assert _le_frais(fiche(achats, d["id"]), "TVA")["montant"] == 12_001.37
    with sessions()() as s:
        x = s.get(DossierFrais, tva["id"])
        assert x is not None and (x.montant, x.taux) == (Decimal("12001.37"), Decimal("20"))
        detail = s.scalar(
            select(Journal.detail).where(Journal.action == "frais", Journal.objet_id == d["id"])
        )
    assert detail == f"{d['numero']} · TVA 12 001,37 MAD (20 %)"
    d = frais(achats, d, "DD", 1_500.4, taux="2.5")
    assert _le_frais(d, "DD")["taux"] == 2.5


def test_modifier_garde_le_montant(achats: TestClient, direction: TestClient) -> None:
    """Changer le taux ou un texte ne touche pas au montant ; le serveur n'en recalcule aucun."""
    d = frais(achats, _dossier(achats, "TX-MOD"), "DD", 1_499.99, taux="2.5")
    x = _le_frais(d, "DD")
    d = ok(_corriger(achats, d, x, taux=3))
    x = _le_frais(d, "DD")
    assert (x["montant"], x["taux"]) == (1_499.99, 3)
    d = ok(_corriger(achats, d, x, commentaire="quittance reçue"))
    x = _le_frais(d, "DD")
    assert (x["montant"], x["taux"], x["commentaire"]) == (1_499.99, 3, "quittance reçue")
    d = ok(_corriger(achats, d, x, montant=1_800, taux=None))
    x = _le_frais(d, "DD")
    assert (x["montant"], x["taux"]) == (1_800, None)
    historique = ok(direction.get(f"/api/historique/dossier/{d['id']}"))
    modifs = [h for h in historique if h["action"] == "frais modifié"]
    assert any(h["avant"] == {"taux": 2.5} and h["apres"] == {"taux": 3} for h in modifs)


@pytest.mark.parametrize("taux", [0, 101, -2.5, "2.12345"])
def test_taux_hors_bornes_refuse(achats: TestClient, taux: Any) -> None:
    """0, plus de 100, négatif, ou plus de quatre décimales : 422, rien n'est écrit."""
    d = _dossier(achats, "TX-BORNES")
    r = frais(achats, d, "DD", 1_000, code=422, taux=taux)
    assert "taux" in str(r)
    assert fiche(achats, d["id"])["frais"] == []


def test_taux_hors_bornes_refuse_par_la_base(achats: TestClient) -> None:
    """Le CHECK `ck_dossier_frais_taux_borne` tient, même sans passer par l'API."""
    d = frais(achats, _dossier(achats, "TX-CK"), "DD", 1_000, taux=2)
    for faux in (Decimal(0), Decimal("100.0001")):
        with sessions()() as s, s.begin():
            x = s.get(DossierFrais, d["frais"][0]["id"])
            assert x is not None
            x.taux = faux
            with pytest.raises(IntegrityError, match="ck_dossier_frais_taux_borne"):
                s.flush()
            s.rollback()


def test_frais_sans_taux_toujours_accepte(achats: TestClient) -> None:
    """Comme avant le lot : sans `taux`, ou `taux` vide, le frais s'enregistre ; 100 % permis."""
    d = frais(achats, _dossier(achats, "TX-SANS"), "FRET", 3_500, prestataire="CMA CGM")
    assert _le_frais(d, "FRET")["taux"] is None
    d = frais(achats, d, "TRANSIT", 800, taux=None)
    assert _le_frais(d, "TRANSIT")["taux"] is None
    d = frais(achats, d, "PORT", 250, taux=100)
    assert _le_frais(d, "PORT")["taux"] == 100


def test_montant_jamais_recalcule_quand_le_cours_change(achats: TestClient) -> None:
    """Droits à 2,5 % de la valeur en douane (60 000 → 1 500) ; le cours payé passe à 12 : la
    valeur en douane suit le cours (72 000), le montant du frais reste 1 500 et son taux 2,5."""
    d = frais(achats, _dossier(achats, "TX-COURS"), "DD", 1_500, taux="2.5")
    assert sum(x["valeur_douane_mad"] for x in d["lignes"]) == pytest.approx(60_000)
    d = ok(
        achats.put(f"/api/dossiers/{d['id']}/cours", json={"version": d["version"], "cours": 12})
    )
    assert sum(x["valeur_douane_mad"] for x in d["lignes"]) == pytest.approx(72_000)
    assert [x["cours_applique"] for x in d["lignes"]] == [12, 12]
    for lu in (d, fiche(achats, d["id"])):
        dd = _le_frais(lu, "DD")
        assert (dd["montant"], dd["taux"]) == (1_500, 2.5)
    assert d["droits_mad"] == 1_500


def test_le_classeur_des_dossiers_dit_le_taux(achats: TestClient) -> None:
    """Feuille « Frais » : une colonne « Taux » (fraction, format %), vide sans taux."""
    d = frais(achats, _dossier(achats, "TX-XLS"), "DD", 1_500, taux="2.5")
    d = frais(achats, d, "FRET", 3_500)
    r = achats.get("/api/export/dossiers", params={"q": d["numero"]})
    lignes = list(load_workbook(io.BytesIO(r.content))["Frais"].iter_rows(values_only=True))
    entetes = next(x for x in lignes if "Taux" in x)
    rang = entetes.index("Taux")
    assert entetes[rang + 1] == "Montant MAD"
    taux = {x[entetes.index("Poste")]: x[rang] for x in lignes if x[0] == d["numero"]}
    assert taux == {"Droits de douane": pytest.approx(0.025), "Fret": None}
