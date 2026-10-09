"""Plan de production, versions du plan, réalisé, besoins détaillés (routes 90, 91, 145–148, 237).

Héritage REGLES A § 1.9 (R-REF-70, X-REF-53), Gestion MP `api/production_plan.py` ; défauts
D-01 (plan d'une qualité arrêtée), D-17 (tissé hors plan, total hors croissance). Le passé est
jugé sur le plan de son époque (décision de Mohamed).
"""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta

import pytest
from fabrique import *  # noqa: F403
from fabriques.qualites import UN, corps_fiche, creer_qualite, lire_qualite, ok, ordre_valide, refus
from fabriques.referentiels import matiere, stock_initial
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select, update

from gestion_mp_pro.calcul.besoins import besoins_mensuels
from gestion_mp_pro.calcul.plan import plan_mensuel
from gestion_mp_pro.db import fabrique as sessions
from gestion_mp_pro.modeles import PlanVersion
from gestion_mp_pro.services import donnees_moteur, horloge

JOUR = date(2026, 10, 8)


def _ligne(qualite: dict, base: float, coefficients: list[float] | None = None) -> dict:
    """Une ligne du plan global, à la version lue."""
    return {
        "qualite_id": qualite["id"],
        "base_mensuelle": base,
        "coefficients": coefficients or UN,
        "version": qualite["version"],
    }


def _plan(client: TestClient, jour: str = "") -> dict:
    return ok(client.get("/api/production/plan" + (f"?date={jour}" if jour else "")))  # type: ignore[no-any-return]


def test_le_plan_s_enregistre_d_un_bloc(
    direction: TestClient, achats: TestClient, alimentation: TestClient
) -> None:
    """Hérité : test_production_plan::test_le_plan_s_enregistre_d_un_bloc."""
    a = matiere(prix="2")
    q1 = creer_qualite(direction, poids_m2={"Poil": 2.0}, parts={a: 100})
    q2 = creer_qualite(direction, poids_m2={"Poil": 1.0}, parts={a: 100})
    saison = [1.0] * 11 + [2.0]
    corps = {"qualites": [_ligne(q1, 1000, saison), _ligne(q2, 500)]}
    assert achats.put("/api/production/plan", json=corps).status_code == 403
    assert alimentation.put("/api/production/plan", json=corps).status_code == 403
    vue = ok(direction.put("/api/production/plan", json=corps))
    assert vue["modifiees"] == 2
    versions = {x["id"]: x["version"] for x in vue["qualites"]}
    # Renvoyer le même plan ne change rien, et n'écrit ni version ni journal.
    corps = {
        "qualites": [
            _ligne(q1 | {"version": versions[q1["id"]]}, 1000, saison),
            _ligne(q2 | {"version": versions[q2["id"]]}, 500),
        ]
    }
    assert ok(direction.put("/api/production/plan", json=corps))["modifiees"] == 0
    assert len(ok(direction.get(f"/api/qualites/{q1['id']}/plan/versions"))) == 2
    vue = _plan(alimentation, "2026-01-15")
    pla = next(q for q in vue["qualites"] if q["id"] == q1["id"])
    assert pla["total_annuel"] == 13000
    assert len(pla["glissant"]) == 12 and pla["glissant"][11] == 2000  # décembre
    assert pla["kg_m2"] == pytest.approx(2.0 * 1.02)
    # Tout passe, ou rien : un coefficient faux bloque aussi la qualité juste.
    faux = {
        "qualites": [
            _ligne(q1 | {"version": versions[q1["id"]]}, 1),
            _ligne(q2 | {"version": versions[q2["id"]]}, 1, [1.0] * 11 + [99]),
        ]
    }
    r = direction.put("/api/production/plan", json=faux)
    refus(r, 422, "un coefficient de saisonnalité reste entre 0 et 20. Rien n'a été enregistré.")
    assert lire_qualite(direction, q1["id"])["base_mensuelle"] == 1000


def test_plan_global_refus(direction: TestClient) -> None:
    """Doublon (422), qualité inconnue (404), version périmée (409 qui nomme la qualité)."""
    q = creer_qualite(direction)
    deux = {"qualites": [_ligne(q, 1), _ligne(q, 2)]}
    refus(
        direction.put("/api/production/plan", json=deux),
        422,
        "Une qualité apparaît deux fois dans le plan envoyé.",
    )
    inconnue = {"qualites": [_ligne(q, 1), _ligne({"id": 999999, "version": 1}, 1)]}
    refus(
        direction.put("/api/production/plan", json=inconnue),
        404,
        "Qualité introuvable : le plan n'a pas été enregistré.",
    )
    ok(direction.put("/api/production/plan", json={"qualites": [_ligne(q, 10)]}))
    perime = direction.put("/api/production/plan", json={"qualites": [_ligne(q, 20)]})
    assert refus(perime, 409, f"La qualité {q['code']}")["code"] == "conflit_version"


def test_plan_d_une_qualite(direction: TestClient, achats: TestClient) -> None:
    """R-REF-70, X-REF-53 : coefficient hors 0–20 refusé ; version tracée, avec motif."""
    q = creer_qualite(direction)
    url = f"/api/qualites/{q['id']}/plan"
    assert (
        achats.put(url, json={"base_mensuelle": 1, "coefficients": UN, "version": 1}).status_code
        == 403
    )
    refus(
        direction.put(
            url, json={"base_mensuelle": 1, "coefficients": [21] + UN[1:], "version": q["version"]}
        ),
        422,
        "Un coefficient de saisonnalité reste entre 0 et 20.",
    )
    r = ok(
        direction.put(
            url,
            json={
                "base_mensuelle": 1200,
                "coefficients": UN,
                "motif": "salon",
                "version": q["version"],
            },
        )
    )
    assert r["base_mensuelle"] == 1200 and r["version"] > q["version"]
    v = ok(direction.get(f"/api/qualites/{q['id']}/plan/versions"))
    assert [x["base_mensuelle"] for x in v] == [1200, 0] and v[0]["motif"] == "salon"
    assert v[0]["en_vigueur"] and v[1]["depuis_toujours"] and v[0]["par"]["nom"] == "Direction"
    refus(
        direction.put(url, json={"base_mensuelle": 1, "coefficients": UN, "version": q["version"]}),
        409,
        "modifiée entre-temps",
    )


def test_tout_chemin_qui_change_le_plan_en_garde_la_version(direction: TestClient) -> None:
    """Hérité : test_plan_besoins::test_tout_chemin_qui_change_le_plan_en_garde_la_version."""
    q = creer_qualite(direction, poids_m2={"Poil": 1}, parts={matiere(): 100})
    q = ok(
        direction.put(
            f"/api/qualites/{q['id']}/plan",
            json={"base_mensuelle": 800, "coefficients": UN, "version": q["version"]},
        )
    )
    vue = ok(direction.put("/api/production/plan", json={"qualites": [_ligne(q, 900)]}))
    q = lire_qualite(direction, q["id"])
    assert vue["modifiees"] == 1
    q = ok(direction.put(f"/api/qualites/{q['id']}", json=corps_fiche(q, actif=False)))
    versions = list(reversed(ok(direction.get(f"/api/qualites/{q['id']}/plan/versions"))))
    assert [x["base_mensuelle"] for x in versions] == [0, 800, 900, 900]
    assert versions[-1]["en_service"] is False and all(x["en_service"] for x in versions[:-1])
    with sessions()() as s:
        d = donnees_moteur.charger(s)
    plan = next(x for x in d.plan if x.qualite == q["code"])
    assert plan.base_mensuelle == 900 and not plan.actif  # la base reste, pour le passé
    assert plan_mensuel(d, JOUR)[q["code"]] == [0.0] * 12
    assert d.versions_plan[q["code"]][-1].actif is False
    q = ok(direction.put(f"/api/qualites/{q['id']}", json=corps_fiche(q, actif=True)))
    assert ok(direction.get(f"/api/qualites/{q['id']}/plan/versions"))[0]["en_service"]


def test_d01_le_plan_d_une_qualite_arretee_ne_change_pas(direction: TestClient) -> None:
    """Hérité : test_defauts_d::test_d01_le_plan_d_une_qualite_arretee_ne_change_pas."""
    q = creer_qualite(direction, actif=False)
    r = direction.put("/api/production/plan", json={"qualites": [_ligne(q, 500)]})
    refus(r, 409, f"La qualité {q['code']} est arrêtée : son plan ne compte plus.")
    assert "Rien n'a été enregistré." in r.json()["detail"]
    url = f"/api/qualites/{q['id']}/plan"
    refus(
        direction.put(
            url, json={"base_mensuelle": 500, "coefficients": UN, "version": q["version"]}
        ),
        409,
        "arrêtée",
    )
    # Son plan renvoyé tel quel passe (l'écran envoie ce qu'il a).
    ok(direction.put(url, json={"base_mensuelle": 0, "coefficients": UN, "version": q["version"]}))


def test_d17_tisse_hors_plan_une_qualite_a_base_zero(direction: TestClient) -> None:
    """Hérité : test_defauts_d::test_d17_tisse_hors_plan_une_qualite_a_base_zero ; et le total
    annuel est la saisie hors croissance, nul pour une arrêtée."""
    q = creer_qualite(direction, poids_m2={"Poil": 1}, parts={matiere(): 100})
    ordre_valide(q["id"], 7)
    vue = _plan(direction)
    assert {"qualite_id": q["id"], "qualite": q["code"], "m2": 7} in vue["tisse_hors_plan"]
    ligne = next(x for x in vue["qualites"] if x["id"] == q["id"])
    assert ligne["tisse_mois_en_cours"] == 7 and ligne["total_annuel"] == 0


def test_plan_des_mois_clos_fige(direction: TestClient) -> None:
    """Hérité : test_stock_atelier::test_plan_des_mois_clos_fige (assertion inconditionnelle,
    ESSAIS § 15) : changer le plan aujourd'hui ne réécrit pas l'écart des mois clos."""
    a = matiere()
    stock_initial(a, 100, le=date(2025, 1, 10))  # l'usine tourne depuis janvier 2025
    q = creer_qualite(direction, poids_m2={"Poil": 1}, parts={a: 100})
    q = ok(
        direction.put(
            f"/api/qualites/{q['id']}/plan",
            json={"base_mensuelle": 1000, "coefficients": UN, "version": q["version"]},
        )
    )
    with sessions()() as s, s.begin():  # ce plan vaut depuis le 1er janvier 2025
        s.execute(
            update(PlanVersion)
            .where(PlanVersion.qualite_id == q["id"], PlanVersion.base_mensuelle == 1000)
            .values(valide_depuis=datetime(2025, 1, 1, tzinfo=horloge.FUSEAU))
        )
    ordre_valide(q["id"], 300, jour=date(2026, 9, 5))
    avant = ok(direction.get("/api/production/realise"))
    ligne = next(x for x in avant["lignes"] if x["qualite_id"] == q["id"])
    clos = avant["mois_clos"]
    assert avant["plan_fige"] and clos == 11
    assert ligne["plan"][:clos] == pytest.approx([1000.0] * clos) and ligne["realise"][10] == 300
    ok(direction.put("/api/production/plan", json={"qualites": [_ligne(q, 2000)]}))
    apres = ok(direction.get("/api/production/realise"))
    ligne2 = next(x for x in apres["lignes"] if x["qualite_id"] == q["id"])
    assert ligne2["plan"][:clos] == pytest.approx(ligne["plan"][:clos])
    assert ligne2["plan"][clos] != ligne["plan"][clos]


def test_le_mrp_detaille_rejoint_les_besoins(direction: TestClient, achats: TestClient) -> None:
    """Hérité : test_production_plan::test_le_mrp_detaille_rejoint_les_besoins (les besoins
    simples sont lus au moteur : la route `/besoins` est au lot D8)."""
    a = matiere(prix="2")
    q1 = creer_qualite(direction, poids_m2={"Poil": 2.0}, parts={a: 100})
    q2 = creer_qualite(direction, poids_m2={"Poil": 1.0}, parts={a: 100})
    ok(
        direction.put(
            "/api/production/plan", json={"qualites": [_ligne(q1, 1000), _ligne(q2, 500)]}
        )
    )
    detail = ok(achats.get(f"/api/production/besoins?date=2026-01-15&matiere={a}"))
    assert len(detail["lignes"]) == 1
    ligne = detail["lignes"][0]
    with sessions()() as s:
        d = donnees_moteur.charger_au(s, date(2026, 1, 15))
    attendu = besoins_mensuels(d, date(2026, 1, 15))[ligne["reference"]]
    assert ligne["total"] == pytest.approx(sum(attendu)) == pytest.approx(12 * 2500 * 1.02)
    assert {x["qualite_id"] for x in ligne["qualites"]} == {q1["id"], q2["id"]}
    for i in range(12):
        assert sum(x["kg"][i] for x in ligne["qualites"]) == pytest.approx(ligne["kg"][i])
    assert ligne["role"] == "Poil" and len(ligne["retenu"]) == 12 and ligne["matiere_id"] == a


def test_les_besoins_detailles_a_une_date_passee_lisent_le_stock_de_cette_date(
    achats: TestClient, alimentation: TestClient
) -> None:
    """Hérité : test_audit2_achat::test_les_besoins_detailles_a_une_date_passee_lisent_le_stock
    _de_cette_date ; et l'alimentation ne reçoit pas le conseil d'achat (E-01)."""
    hier = (JOUR - timedelta(days=1)).isoformat()
    besoins = ok(achats.get(f"/api/production/besoins?date={hier}"))
    stock = {x["reference"]: x for x in ok(achats.get(f"/api/stock?date={hier}"))["lignes"]}
    assert besoins["lignes"]
    for x in besoins["lignes"]:
        assert x["stock_disponible"] == pytest.approx(stock[x["reference"]]["stock_disponible"])
        assert x["stock_projete"] == pytest.approx(stock[x["reference"]]["stock_projete"])
    lu = ok(alimentation.get(f"/api/production/besoins?date={hier}"))
    assert all(x["quantite_a_commander"] is None for x in lu["lignes"])
    assert all(x["quantite_a_commander_achat"] is None for x in lu["lignes"])


def test_exports_de_la_production(achats: TestClient, alimentation: TestClient) -> None:
    """Route 237 : trois onglets ; sans quantité à commander pour l'alimentation."""
    for onglet in ("glissant", "realise", "besoins"):
        r = achats.get(f"/api/export/production?onglet={onglet}")
        assert r.status_code == 200 and r.content[:2] == b"PK", onglet
    entetes = []
    for client in (achats, alimentation):
        r = client.get("/api/export/production?onglet=besoins")
        feuille = load_workbook(io.BytesIO(r.content)).active
        assert feuille is not None
        entetes.append(next(feuille.iter_rows(values_only=True)))
    assert "À commander (kg)" in entetes[0] and "À commander (kg)" not in entetes[1]
    assert achats.get("/api/export/production?onglet=autre").status_code == 422


def test_le_plan_se_lit_par_les_trois_roles(alimentation: TestClient) -> None:
    """Routes 145, 147, 148 : lecture C ; date future refusée."""
    for chemin in ("/api/production/plan", "/api/production/realise", "/api/production/besoins"):
        assert alimentation.get(chemin).status_code == 200
        assert alimentation.get(chemin + "?date=2026-10-09").status_code == 422
    with sessions()() as s:
        assert s.scalar(select(PlanVersion.id).limit(1)) is not None
