"""Export du journal du stock **en flux** (route 106, lot PAGINATION ; V4 § 6).

Preuve que le classeur rendu en flux (`exports/flux.py`, `mouvements.tous_en_flux`) a le **même
contenu** que l'ancien classeur bâti en une fois (mêmes lignes, même en-tête, mêmes valeurs,
mêmes prix masqués), et que les lignes sont consommées **une à une** (générateur, pas de liste de
tout l'historique en mémoire).
"""

from __future__ import annotations

import types
from collections.abc import Iterator
from io import BytesIO

import pytest
from fabrique import *  # noqa: F403
from fabriques.stock import entree, jour, lot, matiere, sortie, unique
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from gestion_mp_pro.db import fabrique
from gestion_mp_pro.exports import stock as classeurs
from gestion_mp_pro.services import mouvements


@pytest.fixture
def base() -> Iterator[Session]:
    """Une session validée (les routes ont leurs propres sessions)."""
    with fabrique()() as s:
        yield s


def _journal(s: Session) -> tuple[int, str]:
    """Un petit journal valorisé : 500 kg (J−2), 100 kg (hier), 100 kg sortis aujourd'hui."""
    ref = unique("Fil en flux")
    m = matiere(s, ref, prix=25)
    x = lot(s, m, entre_le=jour(-2))
    entree(s, x, 500, jour(-2), prix=20)
    entree(s, x, 100, jour(-1), prix=30)
    sortie(s, x, 100)
    s.commit()
    return m.id, ref


def _lignes(contenu: bytes) -> list[tuple[object, ...]]:
    ws = load_workbook(BytesIO(contenu)).active
    assert ws is not None
    return list(ws.iter_rows(values_only=True))


def test_le_flux_rend_le_meme_classeur_que_la_version_en_une_fois(
    base: Session, direction: TestClient
) -> None:
    """Le classeur en flux (route) = l'ancien classeur (mêmes valeurs, en-tête compris)."""
    matiere_id, _ = _journal(base)
    f = mouvements.FiltresMouvements(matiere_id=matiere_id)
    attendu = classeurs.mouvements(mouvements.tous(base, f, avec_valeur=True), voit_prix=True)
    recu = direction.get("/api/export/mouvements", params={"matiere": matiere_id})
    assert recu.status_code == 200 and "attachment" in recu.headers["content-disposition"]
    assert _lignes(recu.content) == _lignes(attendu)


def test_l_alimentation_n_a_aucun_prix_dans_le_flux(
    base: Session, alimentation: TestClient
) -> None:
    """Barrière 4 : aucune colonne de prix ni de valeur dans le flux de l'alimentation."""
    matiere_id, _ = _journal(base)
    entete = _lignes(
        alimentation.get("/api/export/mouvements", params={"matiere": matiere_id}).content
    )[0]
    assert not {"Prix / kg", "Valeur (MAD)", "Cours", "Devise"} & set(entete)


def test_les_lignes_sont_un_generateur_consomme_une_a_une(base: Session) -> None:
    """`tous_en_flux` est un générateur (aucune liste de tout l'historique) ; il rend les mêmes
    lignes valorisées que `tous`."""
    matiere_id, _ = _journal(base)
    f = mouvements.FiltresMouvements(matiere_id=matiere_id)
    flux = mouvements.tous_en_flux(base, f, avec_valeur=True)
    assert isinstance(flux, types.GeneratorType)
    en_flux = list(flux)
    en_une_fois = mouvements.tous(base, f, avec_valeur=True)
    assert [(x.id, x.quantite, x.valeur_mad) for x in en_flux] == [
        (x.id, x.quantite, x.valeur_mad) for x in en_une_fois
    ]
