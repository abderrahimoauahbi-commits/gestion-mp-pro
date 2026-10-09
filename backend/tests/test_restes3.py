"""Restes serveur des lots D2, D3, D6, D8 (RESTES3) : chacun son essai, qui échoue sans lui.

- D2 (F3) : un montage d'un seul étage fait au bureau (route 99) est rangé « bureau » ;
- D3 (F1, F4) : `en_retard`, `arrivee_attendue`, `arrivee_sous_7_jours` comptés au serveur ;
  la fiche dit pourquoi la date promise est figée ; `/export/plan-achat` suit les filtres liés
  de l'écran ;
- D6 (F3) : `OrdreResume.metier` = « CODE · Nom » ;
- D8/D1 (F8) : `GET /periode-close/brouillons` annonce ce qui empêcherait la clôture.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from fabrique import *  # noqa: F403
from fabriques.pilotage import Usine, fixture_usine, jour  # noqa: F401
from fabriques.qualites import creer_metier, lire_metier, version_etage
from fabriques.referentiels import devise, matiere, qualite
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from gestion_mp_pro.modeles import Ordre, Reception
from gestion_mp_pro.modeles.enums import StatutDocument
from gestion_mp_pro.services import donnees_moteur


def ok(r: Any, code: int = 200) -> Any:
    assert r.status_code == code, r.text
    return r.json()


def test_un_montage_d_un_etage_fait_au_bureau_est_dit_bureau(
    direction: TestClient, alimentation: TestClient
) -> None:
    poil = matiere(poids_bobine="2")
    x = creer_metier(direction)
    fil = {"matiere_id": poil, "metres": 2, "pourcentage": 100, "commentaire": ""}
    corps = {"etages": [{"etage": 1, "version": version_etage(x, 1), "fils": [fil]}],
             "confirme_effet_kg": 40}  # fmt: skip
    ok(direction.put(f"/api/metiers/{x['id']}/montage", json=corps))
    x = lire_metier(direction, x["id"])
    telephone = {"version": version_etage(x, 2), "fils": [fil], "confirme_effet_kg": 40}
    ok(alimentation.put(f"/api/metiers/{x['id']}/etages/2", json=telephone))
    historique = ok(direction.get(f"/api/metiers/{x['id']}/montages"))
    assert [(h["origine"], h["etage"]) for h in historique] == [("téléphone", 2), ("bureau", None)]
    assert historique[1]["apres"][0]["etage"] == 1


def test_retard_et_arrivees_a_7_jours_comptes_au_serveur(usine: Usine) -> None:
    u = usine
    f, m = u.fournisseur("Filature des arrivées"), u.matiere("FIL ARRIVEE")
    retard = u.commande(f, m, 100, le=jour(-40), promise=jour(-10))
    proche = u.commande(f, m, 100, le=jour(-5), promise=jour(3))
    loin = u.commande(f, m, 100, le=jour(-5), promise=jour(30))
    liste = {c["id"]: c for c in u.get(u.achats, "/api/commandes")}
    vus = {i: (liste[i]["en_retard"], liste[i]["arrivee_attendue"],
               liste[i]["arrivee_sous_7_jours"]) for i in (retard, proche, loin)}  # fmt: skip
    assert vus == {
        retard: (True, jour(-10).isoformat(), False),
        proche: (False, jour(3).isoformat(), True),
        loin: (False, jour(30).isoformat(), False),
    }
    en_retard = {c["id"] for c in u.get(u.achats, "/api/commandes", filtre="retard")}
    assert {i for i, c in liste.items() if c["en_retard"]} == en_retard
    fiche = u.get(u.achats, f"/api/commandes/{retard}")
    assert fiche["promesse_figee"] == (
        f"La livraison était promise pour le {jour(-10):%d/%m/%Y}, date passée : elle ne se "
        "repousse plus. Noter la nouvelle date annoncée par le fournisseur dans l'arrivée "
        "prévue (ETA) du suivi d'expédition."
    )
    assert u.get(u.achats, f"/api/commandes/{loin}")["promesse_figee"] is None
    assert fiche["kg_en_quarantaine"] == 0


def _references(r: Any) -> list[str]:
    assert r.status_code == 200, r.text
    lignes = list(load_workbook(BytesIO(r.content)).worksheets[0].iter_rows(values_only=True))
    entete = next(x for x in lignes if x and x[0] == "Référence")
    return sorted(str(x[0]) for x in lignes[lignes.index(entete) + 1 :] if x and x[0])


def test_le_classeur_du_plan_d_achat_suit_les_filtres_de_l_ecran(usine: Usine) -> None:
    u = usine
    a, b = u.fournisseur("Filature A"), u.fournisseur("Filature B")
    ra, rb = u.matiere("PLAN A", fournisseur_id=a), u.matiere("PLAN B", fournisseur_id=b)
    qualite(recette={ra: 50, rb: 50}, poids={"Poil": "1"}, base_mensuelle="1000", s=u.s)
    u.s.flush()
    donnees_moteur.invalider()

    def exporter(**filtres: str) -> list[str]:
        params = {"onglet": "references", **filtres}
        return _references(u.achats.get("/api/export/plan-achat", params=params))

    assert exporter() == ["PLAN A", "PLAN B"]
    assert exporter(fournisseur="Filature A") == ["PLAN A"]
    assert exporter(statut="RUPTURE") == ["PLAN A", "PLAN B"]
    assert exporter(statut="CRITIQUE") == []
    assert exporter(priorite="Tier 1") == ["PLAN A", "PLAN B"]
    assert exporter(priorite="Tier 4") == []


def test_le_metier_d_un_ordre_se_lit_code_et_nom(usine: Usine) -> None:
    u = usine
    metier = u.metier()
    metier.nom = "Jacquard du fond"
    o = u.s.get(Ordre, u.ordre("OF-METIER"))
    assert o is not None
    o.metier_id = metier.id
    (lu,) = [x for x in u.get(u.achats, "/api/ordres") if x["id"] == o.id]
    assert lu["metier"] == f"{metier.code} · Jacquard du fond"
    metier.nom = ""
    (lu,) = [x for x in u.get(u.achats, "/api/ordres") if x["id"] == o.id]
    assert lu["metier"] == metier.code


def test_les_brouillons_de_la_periode_s_annoncent_avant_la_cloture(usine: Usine) -> None:
    u = usine
    f = u.fournisseur("Filature de septembre")
    u.s.add(Reception(date=jour(-23), fournisseur_id=f, devise="MAD",
                      statut=StatutDocument.brouillon))  # fmt: skip
    u.s.flush()
    lu = u.get(u.direction, "/api/periode-close/brouillons", date="2026-09-30")
    assert (lu["jusqu_au"], lu["total"]) == ("2026-09-30", 1)
    assert {g["genre"]: g["nombre"] for g in lu["genres"]} == {
        "reception": 1, "ordre": 0, "operation": 0, "inventaire": 0,
    }  # fmt: skip
    assert lu["message"].startswith("Des brouillons sont datés de la période à clore : 1 récep")
    avant = u.get(u.direction, "/api/periode-close/brouillons", date="2026-09-14")
    assert (avant["total"], avant["message"]) == (0, None)
    assert (
        u.achats.get("/api/periode-close/brouillons", params={"date": "2026-09-30"}).status_code
        == 403
    )


def test_les_cours_sont_dates(direction: TestClient, achats: TestClient) -> None:
    """Hérité : test_direction::test_les_cours_sont_dates — le cours dit son jour ; son
    historique, le plus récent d'abord, dit qui l'a changé (RESTES3 : preuve laissée par D8)."""
    code = devise(cours="10")
    lue = next(d for d in ok(achats.get("/api/devises")) if d["code"] == code)
    corps = {"nom": lue["nom"], "cours": 12, "actif": True, "version": lue["version"]}
    ok(achats.put(f"/api/devises/{code}", json=corps))
    lue = next(d for d in ok(achats.get("/api/devises")) if d["code"] == code)
    assert lue["cours"] == 12 and lue["cours_le"].startswith("2026-10-08")
    historique = ok(achats.get(f"/api/devises/{code}/cours"))["elements"]
    assert [h["cours"] for h in historique[:2]] == [12, 10]
    assert historique[0]["par"] == "Achats"
