"""Le plan d'achat par l'API (lot D3, routes 149, 150, 151, 242) : les lignes portent ce que
l'écran suit (matière, fournisseur, stock, commandé, unités possibles) ; le budget est la somme
des lignes ; un brouillon est déduit sans être attendu ; une matière désactivée ne se rachète
pas ; ce qui attend en quarantaine reste attendu sur le bon (D-02) ; la préparation propose
l'unité de commande ; les classeurs disent la liste affichée et le pilotage complet. Les preuves
du moteur (domaine 11) sont portées par le lot S3 (`tests/moteur/`).
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

import pytest
from fabrique import *  # noqa: F403
from fabriques.achats import Achat, bon, geste, ok, recevoir
from fabriques.referentiels import qualite
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from gestion_mp_pro.db import fabrique as fabrique_de_sessions
from gestion_mp_pro.modeles import Matiere
from gestion_mp_pro.services import donnees_moteur


@pytest.fixture
def usine() -> tuple[Achat, int, str]:
    """Une matière en recette (1 kg/m², 1 000 m² par mois), sans stock : à commander."""
    a = Achat.nouveau(devise="MAD", delai=30, pays="MAROC")
    m = a.matiere("PLANFIL", prix="10", poids_bobine="2.5", bobines_palette="100")
    qualite(recette={m: 100}, poids={"Poil": "1.0"}, base_mensuelle="1000")
    with fabrique_de_sessions()() as s:
        ref = s.scalar(select(Matiere.reference).where(Matiere.id == m))
    assert ref is not None
    donnees_moteur.invalider()  # écrit sans journal : la signature du cache ne bouge pas
    return a, m, ref


def ligne_du_plan(client: TestClient, ref: str) -> dict[str, Any] | None:
    """La ligne du plan d'achat d'une référence (ou `None`)."""
    p = ok(client.get("/api/plan-achat"))
    return next((x for x in p["lignes"] if x["reference"] == ref), None)


def test_le_plan_d_achat_porte_ce_que_l_ecran_suit(
    achats: TestClient, usine: tuple[Achat, int, str]
) -> None:
    """Hérité : test_audit2_achat::test_le_budget_est_la_somme_des_lignes_affichees (par l'API)."""
    a, m, ref = usine
    p = ok(achats.get("/api/plan-achat"))
    x = next(x for x in p["lignes"] if x["reference"] == ref)
    assert x["matiere_id"] == m and x["fournisseur_id"] == a.fournisseur_id
    assert x["quantite"] > 0 and x["stock_disponible"] == 0 and x["commande_kg"] == 0
    assert {u["unite"] for u in x["unites_possibles"]} == {"kg", "Bobine", "Palette"}
    g = next(g for g in p["fournisseurs"] if g["fournisseur_id"] == a.fournisseur_id)
    assert g["montant_mad"] == pytest.approx(x["montant_mad"])
    assert p["indicateurs"]["budget_mad"] == pytest.approx(
        sum(y["montant_mad"] for y in p["lignes"])
    )
    assert p["indicateurs"]["budget_mad"] == pytest.approx(
        sum(y["montant_mad"] for y in p["fournisseurs"])
    )
    assert achats.get("/api/plan-achat", params={"date": "2999-01-01"}).status_code == 422


def test_un_brouillon_est_deduit_sans_etre_attendu(
    achats: TestClient, usine: tuple[Achat, int, str]
) -> None:
    """Hérité : test_audit2_achat::test_un_brouillon_n_est_pas_repropose_mais_n_est_pas_attendu
    (par l'API : le bon en préparation a son identifiant, pour l'ouvrir)."""
    a, m, ref = usine
    avant = ligne_du_plan(achats, ref)
    assert avant is not None
    b = bon(achats, a, {m: 300})
    apres = ligne_du_plan(achats, ref)
    assert apres is not None
    assert apres["en_preparation"] == 300 and apres["commande_kg"] == 0
    assert apres["bons_en_preparation"] == [{"id": b["id"], "numero": "Brouillon"}]
    # Déduit, puis arrondi à la palette (250 kg) : à une palette près.
    assert abs(apres["quantite"] - (avant["quantite"] - 300)) < 250
    b = geste(achats, b, "envoyer")
    parti = ligne_du_plan(achats, ref)
    assert parti is not None and parti["en_preparation"] == 0 and parti["commande_kg"] == 300


def test_une_matiere_desactivee_ne_se_rachete_pas(
    achats: TestClient, usine: tuple[Achat, int, str]
) -> None:
    """Hérité : test_audit2_achat::test_une_matiere_ou_un_fournisseur_desactive_ne_se_rachete_pas
    (par l'API : la ligne quitte le plan)."""
    _, m, ref = usine
    assert ligne_du_plan(achats, ref) is not None
    with fabrique_de_sessions()() as s, s.begin():
        x = s.get(Matiere, m)
        assert x is not None
        x.actif = False
    donnees_moteur.invalider()
    assert ligne_du_plan(achats, ref) is None


def test_la_quarantaine_ne_declenche_pas_de_rachat(
    achats: TestClient, usine: tuple[Achat, int, str]
) -> None:
    """Hérité : test_defauts_d::test_d02_la_quarantaine_ne_declenche_pas_de_rachat (D-02) :
    reçu sous réserve, en quarantaine, le bon l'attend encore ; rien n'est racheté."""
    a, m, ref = usine
    b = geste(achats, bon(achats, a, {m: 1000}), "envoyer")
    avant = ligne_du_plan(achats, ref)
    assert avant is not None and avant["commande_kg"] == 1000
    recevoir(b["lignes"][0]["id"], 1000, qualite="reserve")
    apres = ligne_du_plan(achats, ref)
    assert apres is not None and apres["commande_kg"] == 1000 and apres["stock_disponible"] == 0
    assert apres["quantite"] == pytest.approx(avant["quantite"])
    f = ok(achats.get(f"/api/commandes/{b['id']}"))
    assert f["statut"] == "envoyee" and f["livraison"] == "sous_reserve"
    assert f["lignes"][0]["en_quarantaine"] == 1000


def test_la_preparation_propose_l_unite_de_commande(
    achats: TestClient, alimentation: TestClient, usine: tuple[Achat, int, str]
) -> None:
    """Route 150 : les suggestions chez un fournisseur, l'unité de départ (la palette déclarée
    par la fiche), au kilo supérieur ; 404 pour un fournisseur inconnu ; rien pour l'atelier."""
    a, m, _ = usine
    p = ok(achats.get(f"/api/plan-achat/fournisseurs/{a.fournisseur_id}"))
    assert p["fournisseur_id"] == a.fournisseur_id and p["delai_jours"] == 30
    (x,) = p["lignes"]
    assert x["matiere_id"] == m and x["unite_commande"] == "Palette"
    assert x["depart"]["unite"] == "Palette" and x["depart"]["contenance"] == 250
    assert x["depart"]["quantite"] == pytest.approx(x["quantite_suggeree"], abs=1)
    r = achats.get("/api/plan-achat/fournisseurs/999999")
    assert r.status_code == 404 and r.json()["detail"] == "Fournisseur inconnu."
    for url in ("/api/plan-achat", f"/api/plan-achat/fournisseurs/{a.fournisseur_id}",
                "/api/export/plan-achat", "/api/export/pilotage"):  # fmt: skip
        assert alimentation.get(url).status_code == 403, url


def _feuille(r: Any) -> list[tuple[Any, ...]]:
    """Les lignes de la première feuille d'un classeur."""
    assert r.status_code == 200, r.text
    return list(load_workbook(BytesIO(r.content)).worksheets[0].iter_rows(values_only=True))


def test_les_classeurs_du_plan_d_achat(achats: TestClient, usine: tuple[Achat, int, str]) -> None:
    """Route 242 : l'onglet affiché, cherché comme à l'écran ; 151 : le classeur complet."""
    a, _, ref = usine
    lignes = _feuille(
        achats.get("/api/export/plan-achat", params={"onglet": "references", "q": ref})
    )
    entete = next(x for x in lignes if x and x[0] == "Référence")
    assert [x[0] for x in lignes[lignes.index(entete) + 1 :]] == [ref]
    nom = ok(achats.get(f"/api/plan-achat/fournisseurs/{a.fournisseur_id}"))["fournisseur"]
    lignes = _feuille(
        achats.get("/api/export/plan-achat", params={"onglet": "fournisseurs", "q": nom})
    )
    entete = next(x for x in lignes if x and x[0] == "Fournisseur")
    assert [x[0] for x in lignes[lignes.index(entete) + 1 :]] == [nom]
    r = achats.get("/api/export/plan-achat", params={"onglet": "autre"})
    assert r.status_code == 422
    r = achats.get("/api/export/pilotage")
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Stock", "Plan d'achat", "Par fournisseur", "Besoins", "Contrôles"]
    assert len(list(wb["Contrôles"].iter_rows(values_only=True))) >= 38
