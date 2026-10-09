"""Le bon fournisseur (lot D3, routes 164 et 165) : A4, en anglais, sans aucun prix, montant ni
devise, sur le papier à en-tête de l'usine (« HEAD OFFICE », « BANKS », ICE, R.C, N.I.F,
C.N.S.S, patente) ; le même contenu en classeur ; le classeur des bons (166). Preuves héritées
de `test_achats`, `test_audit2_pilotage`, `test_exports`, `test_defauts_a` (A-18, A-19).
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

import pytest
from fabrique import *  # noqa: F403
from fabriques.achats import Achat, bon, geste, ok
from fabriques.documents import reglages
from fabriques.referentiels import unique
from fastapi.testclient import TestClient
from openpyxl import load_workbook

SOCIETE = {
    "nom_entreprise": "POLYFASHIONS CARPET S.A.R.L",
    "forme_juridique": "SOCIETE ANONYME AU CAPITAL DE 20 000 000 DH",
    "adresse": "Zone industrielle Allée 3 Azit 79",
    "ville": "TANGER",
    "banques": "ATTWAFA BANK : 90D012433 · B.M.C.I : 01063 000619 001 82",
    "ice": "001568075000050",
    "rc": "7253",
    "identifiant_fiscal": "04901102",
    "cnss": "2662296",
    "patente": "57221265",
}
FRANCAIS = ("SIÈGE", "SIEGE", "BANQUES", "Livraison", "Fournisseur", "Remarques", "Quantité")


def imprime(html: str) -> str:
    """Le texte imprimé du document : sans la barre et les avertissements de l'écran, ni balises."""
    sans_ecran = re.sub(
        r'<(div|p) class="[^"]*ecran-seulement[^"]*">.*?</\1>', "", html, flags=re.S
    )
    corps = sans_ecran.split("<body", 1)[1]
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", corps))


def valeurs(r: Any) -> list[Any]:
    """Toutes les valeurs du classeur du bon."""
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    return [v for ligne in wb.active.iter_rows(values_only=True) for v in ligne if v is not None]


@pytest.fixture
def eur() -> Achat:
    """Une filature turque en euros : le bleu en bobines de 2,5 kg à 3 EUR."""
    a = Achat.nouveau(
        devise="EUR",
        delai=45,
        pays="TURKEY",
        contact="Ayşe",
        incoterm="CIF",
        nom=unique("FILATURE"),
    )
    a.matiere("BLEU", prix="3", poids_bobine="2.5", bobines_palette="96")
    return a


def test_le_bon_imprime_et_l_excel_du_fournisseur(achats: TestClient, eur: Achat) -> None:
    """Hérité : test_achats::test_le_bon_imprime_et_l_excel_du_fournisseur."""
    lignes = [{"matiere_id": eur.matieres["BLEU"], "quantite": 1120, "unite_achat": "Bobine",
               "contenance": 2.5}]  # fmt: skip
    b = ok(
        achats.post(
            "/api/commandes", json={"fournisseur_id": eur.fournisseur_id, "lignes": lignes}
        ),
        201,
    )[0]
    b = geste(achats, b, "envoyer")
    b = geste(achats, b, "prix", lignes=[{"ligne_id": b["lignes"][0]["id"], "prix": 3.0}],
              mettre_a_jour_catalogue=False)  # fmt: skip
    with reglages(SOCIETE):
        r = achats.get(f"/api/documents/commandes/{b['id']}/bon")
        assert r.status_code == 200 and r.headers["content-type"].startswith("text/html")
        assert r.headers["content-disposition"] == "inline"
        page = r.text
        texte = imprime(page)
        assert "Bobbin of 2.5 kg" in texte and "448" in texte and "1,120" in texte
        assert "size: A4" in page and "Purchase order" in texte and "Expected delivery" in texte
        assert "08 Oct 2026" in texte and b["numero"] in texte and "CIF" in texte
        for interdit in ("Price", "Prix", "3,360", "3.36", "EUR", "MAD", "Amount", "Total amount"):
            assert interdit not in texte, interdit
        x = valeurs(achats.get(f"/api/export/commandes/{b['id']}/bon"))
    assert "Bobbin of 2.5 kg" in x and "Expected delivery" in x and 1120 in x and 448 in x
    assert 3360 not in x and 3 not in x[x.index("No.") :]
    assert not any(re.search(r"Price|Prix|EUR|MAD|Montant|Amount", str(v)) for v in x)


def test_le_bon_suit_le_papier_a_en_tete_de_l_usine(achats: TestClient, eur: Achat) -> None:
    """Hérité : test_audit2_pilotage::test_le_bon_suit_le_papier_a_en_tete_de_l_usine ; et le
    bon anglais ne contient pas de français (A-18 : « HEAD OFFICE », « BANKS »)."""
    b = geste(achats, bon(achats, eur, {eur.matieres["BLEU"]: 100}), "envoyer")
    with reglages(SOCIETE):
        page = achats.get(f"/api/documents/commandes/{b['id']}/bon").text
        x = " ".join(str(v) for v in valeurs(achats.get(f"/api/export/commandes/{b['id']}/bon")))
    texte = imprime(page)
    for attendu in ("POLYFASHIONS CARPET S.A.R.L", "AU CAPITAL DE 20 000 000 DH", "HEAD OFFICE",
                    "TANGER", "BANKS", "ATTWAFA BANK : 90D012433", "B.M.C.I : 01063 000619 001 82",
                    "ICE : 001568075000050", "R.C : 7253", "N.I.F : 04901102",
                    "C.N.S.S : 2662296", "PATENTE : 57221265"):  # fmt: skip
        assert attendu in texte, attendu
        assert attendu in x, attendu
    for mot in FRANCAIS:
        assert mot not in texte and mot not in x, mot
    assert 'class="nom"' in page and "#d40000" in page.lower() or "rouge-societe" in page


def test_le_bon_de_commande_part_chez_le_fournisseur(achats: TestClient) -> None:
    """Hérité : test_exports::test_le_bon_de_commande_part_chez_le_fournisseur (anglais, sans
    prix ; l'ancienne route française avec prix n'existe plus) ; remarques au fournisseur à
    part, gardées en texte (jamais une formule)."""
    a = Achat.nouveau(devise="MAD", delai=20)
    m = a.matiere("EXPORT POIL BLEU", prix="2.55", poids_bobine="2", bobines_palette="100")
    ligne = {"matiere_id": m, "quantite": 1800, "unite_achat": "Palette", "contenance": 200}
    corps = {"fournisseur_id": a.fournisseur_id, "observations": '=HYPERLINK("http://x")',
             "lignes": [ligne]}  # fmt: skip
    b = geste(achats, ok(achats.post("/api/commandes", json=corps), 201)[0], "envoyer")
    assert achats.get(f"/api/export/bon-de-commande/{b['id']}").status_code == 404
    with reglages({"nom_entreprise": "TAPIS ESSAI SARL", "ice": "001122334455667"}):
        r = achats.get(f"/api/export/commandes/{b['id']}/bon")
    assert r.status_code == 200 and b["numero"] in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Purchase order"]
    feuille = wb["Purchase order"]
    lignes = [x for x in feuille.iter_rows(values_only=True)]
    assert lignes[0][0] == "TAPIS ESSAI SARL"
    textes = " ".join(str(v) for x in lignes for v in x if v is not None)
    assert f"Purchase order {b['numero']}" in textes and "Pallet of 200 kg" in textes
    entete = next(x for x in lignes if x and x[0] == "No.")
    ligne = lignes[lignes.index(entete) + 1]
    assert ligne[1].startswith("EXPORT POIL BLEU")
    assert ligne[entete.index("Qty")] == pytest.approx(9)
    assert ligne[entete.index("Total weight (kg)")] == pytest.approx(1800)
    nombres = [v for x in lignes for v in x if isinstance(v, int | float)]
    assert not any(v == pytest.approx(2.55) or v == pytest.approx(1800 * 2.55) for v in nombres)
    observation = next(
        c for x in feuille.iter_rows() for c in x if c.value == '=HYPERLINK("http://x")'
    )
    assert observation.data_type == "s"
    assert any(x and x[0] == "Remarks" for x in lignes)


def test_le_bon_ne_part_pas_sans_raison_sociale_et_un_brouillon_le_dit(
    achats: TestClient, eur: Achat
) -> None:
    """Hérité : test_defauts_a::test_l_excel_du_bon_ne_part_pas_sans_raison_sociale (A-18) ;
    un brouillon s'imprime « Draft — not sent », filigrane DRAFT, daté « on sending »."""
    b = bon(achats, eur, {eur.matieres["BLEU"]: 100})
    with reglages({"nom_entreprise": ""}):
        r = achats.get(f"/api/export/commandes/{b['id']}/bon")
        assert r.status_code == 409 and "raison sociale" in r.json()["detail"]
        page = achats.get(f"/api/documents/commandes/{b['id']}/bon").text
    assert "elle manque en tête du bon" in page and "ecran-seulement" in page
    texte = imprime(page)
    assert "Draft — not sent" in texte and "on sending" in texte and "to be agreed" in texte
    assert 'class="brouillon"' in page and 'content: "DRAFT"' in page
    assert ok(achats.get(f"/api/commandes/{b['id']}"))["raison_sociale_reglee"] is False


def test_les_metres_d_un_cuir_sont_indicatifs(achats: TestClient) -> None:
    """Un cuir, des franges, un ruban : kilos, et mètres « à titre indicatif » (R-CMD-52)."""
    a = Achat.nouveau(devise="MAD")
    cuir = a.matiere("CUIR", role="Cuir", densite="0.25", unite_prix="ml", prix="2.5")
    b = geste(achats, bon(achats, a, {cuir: 300}), "envoyer")
    with reglages(SOCIETE):
        texte = imprime(achats.get(f"/api/documents/commandes/{b['id']}/bon").text)
        x = valeurs(achats.get(f"/api/export/commandes/{b['id']}/bon"))
    assert "Length (m, indicative)" in texte and "≈ 1,200 m" in texte and "By weight" in texte
    assert "Length (m, indicative)" in x and 1200 in x


def test_qui_imprime_et_ce_qui_est_introuvable(
    achats: TestClient, alimentation: TestClient, direction: TestClient, eur: Achat
) -> None:
    """Hérité : test_exports::test_l_alimentation_n_exporte_pas_un_bon ;
    test_achats::test_l_atelier_ne_lit_ni_commandes_ni_receptions (partie commandes)."""
    b = bon(achats, eur, {eur.matieres["BLEU"]: 10})
    for url in (f"/api/documents/commandes/{b['id']}/bon",
                f"/api/export/commandes/{b['id']}/bon", "/api/commandes",
                f"/api/commandes/{b['id']}", "/api/export/commandes"):  # fmt: skip
        assert alimentation.get(url).status_code == 403, url
    with reglages(SOCIETE):
        assert direction.get(f"/api/export/commandes/{b['id']}/bon").status_code == 200
    r = achats.get("/api/documents/commandes/999999/bon")
    assert r.status_code == 404 and r.json()["detail"] == "Commande introuvable."
    r = achats.get("/api/export/commandes/999999/bon")
    assert r.status_code == 404 and r.json()["detail"] == "Bon de commande introuvable."


def test_l_export_des_commandes_dit_le_prix_dans_la_devise_du_bon(
    achats: TestClient, eur: Achat
) -> None:
    """Hérité : test_defauts_a::test_l_export_des_commandes_dit_le_prix_dans_la_devise_du_bon
    (A-19) ; test_achats::test_les_listes_s_exportent_comme_a_l_ecran (partie commandes)."""
    from gestion_mp_pro.db import fabrique as fabrique_de_sessions
    from gestion_mp_pro.modeles import Matiere

    m = eur.matiere("EXPORT DEVISE", prix="3")
    b = bon(achats, eur, {m: 100}, observations="=SOMME(A1)")
    with fabrique_de_sessions()() as s, s.begin():
        x = s.get(Matiere, m)
        assert x is not None
        x.prix, x.devise = 30, "MAD"  # type: ignore[assignment]
    estime = ok(achats.get(f"/api/commandes/{b['id']}"))["montant_estime"]
    r = achats.get("/api/export/commandes", params={"statut": "brouillon", "q": "export devise"})
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Commandes", "Lignes"]
    lignes = list(wb["Lignes"].iter_rows(values_only=True))
    entete = next(x for x in lignes if x and x[0] == "N°")
    (ligne,) = [x for x in lignes if x and str(x[3]).startswith("EXPORT DEVISE")]
    assert ligne[entete.index("Devise")] == "EUR"
    assert ligne[entete.index("Prix / kg")] * 100 == pytest.approx(estime)
    assert ligne[entete.index("Prix / kg")] != pytest.approx(30)
    bons = list(wb["Commandes"].iter_rows(values_only=True))
    assert any("onglet : Brouillons" in str(x[0]) for x in bons if x and x[0])
    entete = next(x for x in bons if x and x[0] == "N°")
    rangs = [x for x in bons[bons.index(entete) + 1 :]]
    assert [x[entete.index("Statut")] for x in rangs] == ["Brouillon"]
    cellule = next(c for x in wb["Commandes"].iter_rows() for c in x if c.value == "=SOMME(A1)")
    assert cellule.data_type == "s"
    assert achats.get("/api/export/commandes", params={"statut": "inconnu"}).status_code == 422
