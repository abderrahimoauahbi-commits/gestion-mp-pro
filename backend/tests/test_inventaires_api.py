"""Routes des inventaires (lot D5 ; API-PRO 124–133, 239, 240, 245) : un parcours complet sur un
magasin à soi (données validées en base), la feuille de comptage sur le papier à en-tête, les
classeurs, les droits et les montants masqués de l'alimentation.
"""

from __future__ import annotations

from collections.abc import Iterator
from io import BytesIO
from typing import Any

import pytest
from fabrique import *  # noqa: F403
from fabriques.operations import JOUR_DES_INVENTAIRES, Monde, monde
from fabriques.stock import JOUR, autre_magasin
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_prix_masques import fuites

from gestion_mp_pro.db import fabrique


@pytest.fixture
def m() -> Iterator[Monde]:
    """Le décor des essais hérités, dans un magasin à soi, validé en base."""
    with fabrique()() as s:
        decor = monde(s, le=JOUR_DES_INVENTAIRES, courant=autre_magasin(s))
        s.commit()
        yield decor


def _ok(r: Any, code: int = 200) -> Any:
    assert r.status_code == code, r.text
    return r.json()


def _ouvrir(client: TestClient, m: Monde, **criteres: Any) -> dict[str, Any]:
    corps = {
        "date": JOUR_DES_INVENTAIRES.isoformat(),
        "magasin_id": m.courant.id,
        "criteres": criteres,
    }
    return _ok(client.post("/api/inventaires", json=corps), 201)


def _ligne(i: dict[str, Any], lot: str) -> dict[str, Any]:
    return next(x for x in i["lignes"] if x["lot"] == lot)


def test_parcours_complet_par_l_api(
    m: Monde, achats: TestClient, alimentation: TestClient, direction: TestClient
) -> None:
    """Ouvrir, compter, lot trouvé, valider, annuler ; l'alimentation lit sans montant et
    n'écrit pas ; version périmée → 409 ; écarts en kilos et en dirhams."""
    i = _ouvrir(achats, m, matieres=[m.a1.id])
    assert i["numero"].startswith("INV-") and i["perimetre"] == "1 matière(s)"
    assert i["magasin"] == m.courant.code and i["statut"] == "ouvert"
    assert alimentation.post("/api/inventaires", json={"date": JOUR.isoformat()}).status_code == 403
    l1 = _ligne(i, m.lots["E-L1"].numero)
    corps = {
        "version": i["version"],
        "lignes": [{"ligne_id": l1["id"], "compte": 450, "justification": "palette abîmée"}],
    }
    i = _ok(achats.put(f"/api/inventaires/{i['id']}/comptages", json=corps))
    r = achats.put(f"/api/inventaires/{i['id']}/comptages", json=corps)
    assert r.status_code == 409 and r.json()["code"] == "conflit_version"
    vu = _ok(alimentation.get(f"/api/inventaires/{i['id']}"))
    assert fuites(vu) == [] and vu["lignes"][0]["theorique"] is not None
    assert _ok(direction.get(f"/api/inventaires/{i['id']}"))["ecart_valeur_mad"] == pytest.approx(
        -50 * (500 * 20 + 300 * 22 + 100 * 20) / 900  # au CMUP du moment (quarantaine comprise)
    )
    trouve = {
        "version": i["version"],
        "matiere_id": m.a1.id,
        "lot_interne": True,
        "compte": 4,
        "justification": "sans étiquette",
    }
    i = _ok(achats.post(f"/api/inventaires/{i['id']}/lots-trouves", json=trouve))
    assert any(x["lot_trouve"] and x["lot"] == "LOT INTERNE" for x in i["lignes"])
    v = _ok(achats.post(f"/api/inventaires/{i['id']}/valider", json={"version": i["version"]}))
    assert v["statut"] == "valide" and v["ecart_total"] == pytest.approx(-46)
    assert any(x["lot"].startswith("LOT-") for x in v["lignes"] if x["lot_trouve"])
    sans = achats.post(f"/api/inventaires/{i['id']}/annuler", json={"version": v["version"]})
    assert (
        sans.status_code == 422
        and sans.json()["detail"] == "Dire pourquoi l'inventaire est annulé."
    )
    a = _ok(
        achats.post(
            f"/api/inventaires/{i['id']}/annuler",
            json={"version": v["version"], "motif": "mauvais magasin"},
        )
    )
    assert (a["statut"], a["annule_par"], a["annulation_motif"]) == (
        "annule",
        "Achats",
        "mauvais magasin",
    )
    liste = _ok(alimentation.get("/api/inventaires", params={"statut": "annule"}))
    assert any(x["id"] == i["id"] for x in liste) and fuites(liste) == []
    assert achats.get("/api/inventaires/999999").status_code == 404


def test_feuille_de_comptage(m: Monde, achats: TestClient, alimentation: TestClient) -> None:
    """Hérité : test_stock_atelier::test_inventaire_tolerance_perimetre_et_feuille (feuille) et
    ::test_inventaire_second_comptage (feuille de recomptage) : A4 paysage, papier à en-tête, à
    l'aveugle par défaut ; `theorique=true` l'ajoute ; le recomptage ne garde que les lignes à
    recompter, avec le théorique ; l'abandonné le dit."""
    i = _ouvrir(achats, m, roles=[m.a3.role_id])
    url = f"/api/documents/inventaires/{i['id']}/feuille"
    r = alimentation.get(url)
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/html")
    assert r.headers["content-disposition"] == "inline"
    assert m.lots["E-F1"].numero in r.text and "Théorique" not in r.text
    assert "size: A4 landscape" in r.text and "Compté par" in r.text
    assert "Théorique" in achats.get(url, params={"theorique": True}).text
    ligne = i["lignes"][0]
    i = _ok(
        achats.put(
            f"/api/inventaires/{i['id']}/comptages",
            json={"version": i["version"], "lignes": [{"ligne_id": ligne["id"], "compte": 900}]},
        )
    )
    recompte = achats.get(url, params={"recomptage": True}).text
    assert "Feuille de recomptage" in recompte and "Théorique" in recompte
    assert "1 000,0" in recompte.replace(" ", " ").replace("\xa0", " ")
    a = _ok(
        achats.post(
            f"/api/inventaires/{i['id']}/abandonner",
            json={"version": i["version"], "motif": "reporté"},
        )
    )
    assert a["statut"] == "abandonne" and "Abandonné" in achats.get(url).text
    assert achats.get("/api/documents/inventaires/999999/feuille").status_code == 404


def _rangs(r: Any) -> list[tuple[Any, ...]]:
    assert r.status_code == 200, r.text
    ws = load_workbook(BytesIO(r.content)).active
    assert ws is not None
    return list(ws.iter_rows(values_only=True))


def test_classeurs_des_inventaires(m: Monde, achats: TestClient, alimentation: TestClient) -> None:
    """Routes 239, 240 : la liste filtrée par état ; les lignes avec écarts en kilos et en
    dirhams ; réservés à la direction et aux achats."""
    i = _ouvrir(achats, m, matieres=[m.a2.id])
    t1 = i["lignes"][0]
    i = _ok(
        achats.put(
            f"/api/inventaires/{i['id']}/comptages",
            json={
                "version": i["version"],
                "lignes": [{"ligne_id": t1["id"], "compte": 390, "justification": "humidité"}],
            },
        )
    )
    i = _ok(achats.post(f"/api/inventaires/{i['id']}/valider", json={"version": i["version"]}))
    rangs = _rangs(achats.get(f"/api/export/inventaires/{i['id']}"))
    entete = next(x for x in rangs if x[0] == "Référence")
    donnees = rangs[rangs.index(entete) + 1 :]
    assert donnees[0][entete.index("Écart (kg)")] == -10
    assert donnees[0][entete.index("Valeur de l'écart (MAD)")] == pytest.approx(-80)
    liste = _rangs(achats.get("/api/export/inventaires", params={"statut": "valide"}))
    assert any(i["numero"] in x for x in liste)
    assert alimentation.get("/api/export/inventaires").status_code == 403
    assert alimentation.get(f"/api/export/inventaires/{i['id']}").status_code == 403


def test_defauts_et_tournant(m: Monde, achats: TestClient, alimentation: TestClient) -> None:
    """Routes 125, 126 : direction et achats ; la classe ABC n'est jamais lue par l'atelier."""
    d = _ok(achats.get("/api/inventaires/defauts"))
    assert d["tolerance"] == pytest.approx(0.02) and d["cadence_jours"] == {
        "A": 30,
        "B": 90,
        "C": 365,
    }
    dues = _ok(achats.get("/api/inventaires/a-compter", params={"magasin": m.courant.id}))
    assert {m.a1.reference, m.a2.reference, m.a3.reference} <= {x["reference"] for x in dues}
    assert alimentation.get("/api/inventaires/a-compter").status_code == 403
