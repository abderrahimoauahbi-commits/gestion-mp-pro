"""Étapes 1 à 6 du scénario : le plan révisé, le plan d'achat, le bon, son approbation, le bon
imprimé, les prix confirmés en deux fois, les jalons d'expédition.

Calculs à la main (CALCULS § 3.6, 3.7, 3.10, 3.12), au 10 novembre 2026 : 20 jours restants dans
le mois (30 − 10) ; délai du fournisseur 21 j, couverture 30 j ; perte 2 % ; recettes à 100 %.

- Besoin mensuel, plan révisé (Q1 3 000 m², Q2 1 500 m²) : M1 = (1,2 × 3 000 + 1,0 × 1 500) × 1,02
  = 5 202 kg ; M2 = (0,3 × 3 000 + 0,3 × 1 500) × 1,02 = 1 377 ; M3 = (0,5 × 3 000 + 0,4 × 1 500)
  × 1,02 = 2 142.
- Demande du délai [0 ; 21 j] = b × 20/30 + b × 1/30 = 0,7 b ; stock projeté = 0 − 0,7 b.
- Sorties prévues [21 ; 51 j] = b × 29/30 + b × 1/30 = b ; stock minimum recalculé = b × 30/30
  = b (sécurité 30 j) ; besoin net = b + max(b, b) + 0,7 b = 2,7 b.
- M1 : 2,7 × 5 202 = 14 045,4 → 29 palettes de 500 → 14 500 kg ; 36 250 USD ; 362 500 MAD.
- M2 : 2,7 × 1 377 = 3 717,9 kg (au kilo) ; 11 897,28 USD ; 118 972,80 MAD ; bon : 3 718 kg.
- M3 : 2,7 × 2 142 = 5 783,4 → 15 palettes de 400 → 6 000 kg ; 10 800 EUR ; 118 800 MAD.
"""

from __future__ import annotations

from typing import Any

from scenario_monde import JOURS, Monde
from scenario_outils import KG, MAD

BESOIN = {"M1": 5202.0, "M2": 1377.0, "M3": 2142.0}
BESOIN_AVANT = {"M1": 3978.0, "M2": 1071.0, "M3": 1632.0}
"""Au plan de départ (Q1 2 000 m²) : (1,2 × 2 000 + 1 500) × 1,02 = 3 978, etc."""
PROJETE = {"M1": -3641.4, "M2": -963.9, "M3": -1499.4}
A_COMMANDER = {"M1": 14500.0, "M2": 3717.9, "M3": 6000.0}
UNITES = {"M1": 29.0, "M2": 3717.9, "M3": 15.0}
MONTANT_DEVISE = {"M1": 36250.0, "M2": 11897.28, "M3": 10800.0}
MONTANT_MAD = {"M1": 362500.0, "M2": 118972.8, "M3": 118800.0}
BON = {
    "M1": (14500.0, "Palette", 500.0),
    "M2": (3718.0, "kg", 1.0),
    "M3": (6000.0, "Palette", 400.0),
}
"""Le brouillon tiré de la proposition (`depart`) : kilos, unité, contenance."""
MONTANT_BON_USD = 14500 * 2.5 + 3718 * 3.2  # 36 250 + 11 897,60 = 48 147,60 USD
MONTANT_BON_EUR = 6000 * 1.8  # 10 800 EUR


def _ma_ligne(lignes: list[dict[str, Any]], m: Monde, cle: str) -> dict[str, Any] | None:
    return next((x for x in lignes if x.get("matiere_id") == m.ids[cle]), None)


def etape_plan(m: Monde) -> None:
    """Étape 1 (10 novembre, direction) : Q1 passe de 2 000 à 3 000 m² par mois."""
    v, d = m.verif, m.D
    m.jour("plan_revise")
    q1 = m.ids["Q1"]
    realise_avant = d.get("/production/realise")
    for cle in ("M1", "M2", "M3"):
        lu = d.get(f"/stock/{m.ids[cle]}")["ligne"]
        v.egal(f"{cle} : besoin du mois avant révision", lu["besoins"][0], BESOIN_AVANT[cle], KG)
    fiche = d.get(f"/qualites/{q1}")
    corps = {
        "base_mensuelle": 3000,
        "coefficients": [1] * 12,
        "motif": "commande export : 1 000 m² de plus par mois",
        "version": fiche["version"],
    }
    v.refus(
        "les achats ne révisent pas le plan",
        m.A.appeler("PUT", f"/qualites/{q1}/plan", corps),
        403,
        "droit_refuse",
        "la direction",
    )
    d.put(f"/qualites/{q1}/plan", corps)
    versions = d.get(f"/qualites/{q1}/plan/versions")
    v.egal("Q1 : trois versions du plan (d'avant, départ, révision)", len(versions), 3)
    v.egal("Q1 : version en vigueur", versions[0]["base_mensuelle"], 3000.0, 1e-9)
    v.egal("Q1 : révisée le 10", versions[0]["valide_depuis"][:10], "2026-11-10")
    v.egal("Q1 : motif gardé", versions[0]["motif"], corps["motif"])
    m.avant["revision_q1"] = versions[0]["valide_depuis"]
    m.avant["plan_q1_depart"] = versions[1]["valide_depuis"]
    for cle in ("M1", "M2", "M3"):
        lu = d.get(f"/stock/{m.ids[cle]}")["ligne"]
        v.egal(f"{cle} : besoins des 12 mois", lu["besoins"], [BESOIN[cle]] * 12)
        v.egal(f"{cle} : besoin annuel", lu["besoin_annuel"], 12 * BESOIN[cle], KG)
    realise = d.get("/production/realise")
    _mois_clos_inchanges(m, realise_avant, realise)


def _mois_clos_inchanges(m: Monde, avant: dict[str, Any], apres: dict[str, Any]) -> None:
    """Les mois clos ne bougent pas ; le mois en cours prend le plan d'aujourd'hui."""
    v = m.verif
    clos = avant["mois_clos"]
    v.egal("mois clos (réalisé)", apres["mois_clos"], clos)
    plan_avant = {x["qualite_id"]: x["plan"] for x in avant["lignes"]}
    plan_apres = {x["qualite_id"]: x["plan"] for x in apres["lignes"]}
    v.egal(
        "mois clos de toutes les qualités inchangés",
        {q: p[:clos] for q, p in plan_apres.items()},
        {q: p[:clos] for q, p in plan_avant.items()},
    )
    q1 = m.ids["Q1"]
    v.egal("Q1 : mois en cours au plan d'avant", plan_avant[q1][clos], 2000.0, 1e-9)
    v.egal("Q1 : mois en cours au plan révisé", plan_apres[q1][clos], 3000.0, 1e-9)
    autres = {q: p for q, p in plan_apres.items() if q != q1}
    v.egal(
        "les autres qualités ne bougent pas",
        autres,
        {q: p for q, p in plan_avant.items() if q != q1},
    )


def etape_plan_achat(m: Monde) -> None:
    """Étape 2 (10 novembre, achats) : le plan d'achat, puis le bon tiré de la proposition."""
    v, a = m.verif, m.A
    plan = a.get("/plan-achat")
    for cle in ("M1", "M2", "M3"):
        x = _ma_ligne(plan["lignes"], m, cle)
        if not v.vrai(f"{cle} : proposé au plan d'achat", x is not None):
            continue
        assert x is not None
        m.avant["faits"].setdefault("plan_achat", {})[cle] = x["quantite"]
        v.egal(f"{cle} : quantité proposée", x["quantite"], A_COMMANDER[cle], KG)
        v.egal(f"{cle} : unités (palettes ou kg)", x["unites_commande"], UNITES[cle], KG)
        v.egal(f"{cle} : montant en devise", x["montant_devise"], MONTANT_DEVISE[cle], MAD)
        v.egal(f"{cle} : montant en dirhams", x["montant_mad"], MONTANT_MAD[cle], MAD)
        v.egal(f"{cle} : stock projeté", x["stock_projete"], PROJETE[cle], KG)
        v.egal(f"{cle} : consommation mensuelle", x["conso_mensuelle"], BESOIN[cle], KG)
        v.egal(f"{cle} : statut", x["statut"], "RUPTURE")  # code du moteur (RESTES3)
        v.egal(f"{cle} : tier (rupture)", x["tier"], "🔥 TIER 1")
    g = next(x for x in plan["fournisseurs"] if x["fournisseur_id"] == m.ids["F1"])
    v.egal("fournisseur : USD", g["montants_par_devise"]["USD"], 36250 + 11897.28, MAD)
    v.egal("fournisseur : EUR", g["montants_par_devise"]["EUR"], 10800.0, MAD)
    v.egal("fournisseur : dirhams", g["montant_mad"], 362500 + 118972.8 + 118800, MAD)
    v.egal("fournisseur : urgence (ruptures)", g["urgence"], "🚨 BC URGENT")
    v.refus("l'alimentation ne lit pas le plan d'achat", m.L.appeler("GET", "/plan-achat"), 403)
    prep = a.get(f"/plan-achat/fournisseurs/{m.ids['F1']}")
    lignes = []
    for cle, (kg, unite, contenance) in BON.items():
        x = _ma_ligne(prep["lignes"], m, cle)
        assert x is not None
        v.egal(f"{cle} : départ du bon (kg)", x["depart"]["quantite"], kg, KG)
        v.egal(f"{cle} : départ du bon (unité)", x["depart"]["unite"], unite)
        lignes.append(
            {
                "matiere_id": m.ids[cle],
                "quantite": kg,
                "unite_achat": unite,
                "contenance": contenance,
            }
        )
    corps = {
        "fournisseur_id": m.ids["F1"],
        "date_livraison_prevue": "2026-12-01",
        "observations": "Livraison en un camion si possible.",
        "lignes": lignes,
    }
    bons = a.ok("POST", "/commandes", corps, attendu=201)
    _verifier_brouillons(m, bons)


def _verifier_brouillons(m: Monde, bons: list[dict[str, Any]]) -> None:
    """Un bon par devise, brouillons sans numéro ; la proposition les déduit, pas le projeté."""
    v, a = m.verif, m.A
    v.egal("un bon par devise", sorted(b["devise"] for b in bons), ["EUR", "USD"])
    par = {b["devise"]: b for b in bons}
    m.ids["BC_USD"], m.ids["BC_EUR"] = par["USD"]["id"], par["EUR"]["id"]
    for devise, montant, cours in (("USD", MONTANT_BON_USD, 10.0), ("EUR", MONTANT_BON_EUR, 11.0)):
        b = par[devise]
        v.egal(f"bon {devise} : brouillon", b["statut"], "brouillon")
        v.egal(f"bon {devise} : sans numéro", b["numero"], None)
        v.egal(f"bon {devise} : montant estimé", b["montant_estime"], montant, MAD)
        v.egal(f"bon {devise} : en dirhams", b["montant_estime_mad"], montant * cours, MAD)
    unites = {x["matiere_id"]: x["unites"] for b in bons for x in b["lignes"]}
    v.egal("M1 : 29 palettes", unites[m.ids["M1"]], 29.0, 1e-9)
    v.egal("M3 : 15 palettes", unites[m.ids["M3"]], 15.0, 1e-9)
    plan = a.get("/plan-achat")
    for cle in ("M1", "M2", "M3"):
        v.egal(
            f"{cle} : plus reproposé (brouillon déduit)", _ma_ligne(plan["lignes"], m, cle), None
        )
        lu = a.get(f"/stock/{m.ids[cle]}")["ligne"]
        v.egal(f"{cle} : brouillon en préparation", lu["bc_en_preparation"], BON[cle][0], KG)
        v.egal(
            f"{cle} : projeté inchangé (un brouillon n'engage rien)",
            lu["stock_projete"],
            PROJETE[cle],
            KG,
        )


def numero_suivant(m: Monde, serie: str, liste: list[dict[str, Any]], cle: str = "numero") -> str:
    """Le numéro qui suit le plus grand de la série pour 2026 (`BC-2026-0041` → `…-0042`)."""
    prefixe = f"{serie}-2026-"
    rangs = [int(x[cle][len(prefixe) :]) for x in liste if (x.get(cle) or "").startswith(prefixe)]
    return f"{prefixe}{max(rangs, default=0) + 1:04d}"


def etape_approbation(m: Monde) -> None:
    """Étape 3 (10 novembre) : le bon EUR part ; le bon USD attend la direction, revient, part."""
    v, a, d = m.verif, m.A, m.D
    tous = a.get("/commandes?statut=clos")["elements"] + a.get("/commandes")["elements"]
    premier = numero_suivant(m, "BC", tous)
    second = premier[:-4] + f"{int(premier[-4:]) + 1:04d}"
    eur, usd = m.ids["BC_EUR"], m.ids["BC_USD"]
    jour = JOURS["plan_revise"].isoformat()
    b = a.post(f"/commandes/{eur}/envoyer", {"version": _version(m, eur), "date": jour})
    v.egal("bon EUR (118 800 MAD < 200 000) : parti", b["statut"], "envoyee")
    v.egal("bon EUR : numéro suivant", b["numero"], premier)
    b = a.post(f"/commandes/{usd}/envoyer", {"version": _version(m, usd), "date": jour})
    m.avant["faits"]["envoi"] = ("envoyee", premier, b["statut"], b["numero"])
    v.egal("bon USD (481 476 MAD > 200 000) : à approuver", b["statut"], "a_approuver")
    v.egal("bon USD : pas de numéro avant l'approbation", b["numero"], None)
    corps = {"version": b["version"], "date": jour}
    v.refus(
        "les achats n'approuvent pas",
        a.appeler("POST", f"/commandes/{usd}/approuver", corps),
        403,
        "droit_refuse",
        "la direction",
    )
    notes = [n["texte"] for n in d.get("/notifications?non_lues=true")["elements"]]
    v.vrai("la direction est prévenue", any("à approuver" in t for t in notes))
    motif = "Chaîne : confirmer d'abord la teinte avec l'atelier"
    b = d.post(f"/commandes/{usd}/renvoyer", {"version": b["version"], "motif": motif})
    v.egal("renvoyé : brouillon", b["statut"], "brouillon")
    v.egal("renvoyé : motif lu sur la fiche", (b["renvoi"] or {}).get("motif"), motif)
    b = a.post(f"/commandes/{usd}/envoyer", {"version": b["version"], "date": jour})
    v.egal("renvoyé de nouveau : à approuver", b["statut"], "a_approuver")
    b = d.post(f"/commandes/{usd}/approuver", {"version": b["version"], "date": jour})
    v.egal("approuvé : parti", b["statut"], "envoyee")
    v.egal("approuvé : numéro sans trou", b["numero"], second)
    v.egal("montant approuvé (MAD)", b["montant_approuve"], MONTANT_BON_USD * 10.0, MAD)
    decisions = [x["decision"] for x in b["decisions"]]
    m.avant["faits"]["approbation"] = (b["statut"], b["numero"])
    v.egal(
        "décisions du bon",
        decisions[-4:] if decisions[0] != "approbation" else decisions[:4][::-1],
        ["demande_approbation", "renvoi", "demande_approbation", "approbation"],
    )
    v.refus("l'alimentation ne lit pas les bons", m.L.appeler("GET", f"/commandes/{usd}"), 403)
    m.avant["numeros_bc"] = (premier, second)


def _version(m: Monde, cid: int) -> int:
    return int(m.A.get(f"/commandes/{cid}")["version"])


SANS_PRIX = ("USD", "EUR", "MAD", "Price", "price", "Amount", "amount", "2.5", "3.2", "1.8")


def etape_bon_imprime(m: Monde) -> None:
    """Étape 4 (10 novembre, achats) : bon fournisseur A4, anglais, papier à en-tête, sans prix."""
    v, a = m.verif, m.A
    usd = m.ids["BC_USD"]
    numero = m.avant["numeros_bc"][1]
    r = a.appeler("GET", f"/documents/commandes/{usd}/bon")
    v.egal("bon imprimé : HTML", r.statut, 200)
    html = str(r.corps)
    v.vrai("bon imprimé : « size: A4 »", "size: A4" in html)
    texte = _texte(html)
    for morceau in (
        "Purchase order",
        numero,
        "HEAD OFFICE",
        "BANKS",
        "Tapis de l'Atlas (scénario)",
        "Pallet of 500 kg",
        "14,500 kg",
        "3,718 kg",
        "Total ordered 18,218 kg",
        "Expected delivery 01 Dec 2026",
    ):
        v.vrai(f"bon imprimé : « {morceau} »", morceau in texte)
    for interdit in SANS_PRIX:
        v.vrai(f"bon imprimé : aucun « {interdit} »", interdit not in texte)
    r = a.appeler("GET", f"/export/commandes/{usd}/bon")
    v.egal("bon en Excel", r.statut, 200)
    cellules = _cellules(r.corps)
    v.vrai("Excel : le numéro", numero in cellules)
    v.vrai("Excel : 29 palettes", any(c in {"29", "29.0"} for c in cellules))
    for interdit in SANS_PRIX[:6]:
        v.vrai(f"Excel : aucun « {interdit} »", not any(interdit in c for c in cellules))
    v.refus(
        "l'alimentation n'imprime pas de bon",
        m.L.appeler("GET", f"/documents/commandes/{usd}/bon"),
        403,
    )


def _texte(html: str) -> str:
    import re

    html = re.sub(r"<(style|script).*?</\1>", " ", html, flags=re.S)
    from html import unescape

    return re.sub(r"\s+", " ", unescape(re.sub(r"<[^>]+>", " ", html)))


def _cellules(octets: bytes) -> list[str]:
    import io

    from openpyxl import load_workbook

    classeur = load_workbook(io.BytesIO(octets), read_only=True)
    return [
        str(c)
        for f in classeur.worksheets
        for ligne in f.iter_rows(values_only=True)
        for c in ligne
        if c is not None
    ]


def _ligne(bon: dict[str, Any], m: Monde, cle: str) -> dict[str, Any]:
    return next(x for x in bon["lignes"] if x["matiere_id"] == m.ids[cle])


def _confirmer(m: Monde, cid: int, prix: dict[str, float], promesse: str | None) -> Any:
    bon = m.A.get(f"/commandes/{cid}")
    corps: dict[str, Any] = {
        "version": bon["version"],
        "mettre_a_jour_catalogue": False,
        "lignes": [
            {"ligne_id": _ligne(bon, m, k)["id"], "prix": p, "unite": "kg"} for k, p in prix.items()
        ],
    }
    if promesse:
        corps["date_livraison_prevue"] = promesse
    return m.A.appeler("POST", f"/commandes/{cid}/prix", corps)


MONTANT_USD_HAUSSE = 14500 * 3.0 + 3718 * 3.2  # 43 500 + 11 897,60 = 55 397,60 USD


def etape_prix(m: Monde) -> None:
    """Étape 5 (11 et 12 novembre) : prix confirmés en deux fois, cours changé entre les deux."""
    v, a, d = m.verif, m.A, m.D
    usd, eur = m.ids["BC_USD"], m.ids["BC_EUR"]
    m.jour("confirmation_1")
    b = _confirmer(m, eur, {"M3": 1.8}, "2026-11-25").corps
    v.egal("bon EUR : tout confirmé", b["statut"], "confirmee")
    v.egal("bon EUR : cours figé (11)", b["cours"], 11.0, 1e-9)
    b = _confirmer(m, usd, {"M2": 3.2}, "2026-11-25").corps
    v.egal("bon USD : M1 encore à confirmer", b["statut"], "envoyee")
    v.egal("bon USD : cours figé à la 1re confirmation", b["cours"], 10.0, 1e-9)
    v.egal("bon USD : date promise", b["date_livraison_prevue"], "2026-11-25")
    v.egal("bon USD : une ligne à confirmer", b["prix_a_confirmer"], 1)
    m.jour("cours_change")
    fiche = next(x for x in a.get("/devises") if x["code"] == "USD")
    a.put(
        "/devises/USD",
        {"nom": fiche["nom"], "cours": 10.5, "actif": True, "version": fiche["version"]},
    )
    r = _confirmer(m, usd, {"M1": 3.0}, "2026-12-05")
    v.refus("date promise figée", r, 422, "date_figee", "est figée depuis la confirmation des prix")
    b = _confirmer(m, usd, {"M1": 3.0}, None).corps
    v.egal("hausse de 20 % : repasse par la direction", b["statut"], "a_approuver")
    m.avant["faits"]["hausse"] = (b["statut"], b["cours"])
    v.egal("hausse : le numéro reste", b["numero"], m.avant["numeros_bc"][1])
    v.egal("hausse : cours toujours figé (10, pas 10,5)", b["cours"], 10.0, 1e-9)
    v.egal("cours du jour", b["taux_du_jour"], 10.5, 1e-9)
    v.egal("hausse : montant estimé", b["montant_estime"], MONTANT_USD_HAUSSE, MAD)
    v.egal(
        "hausse : en dirhams au cours figé", b["montant_estime_mad"], MONTANT_USD_HAUSSE * 10, MAD
    )
    v.egal("hausse : date promise inchangée", b["date_livraison_prevue"], "2026-11-25")
    b = d.post(f"/commandes/{usd}/approuver", {"version": b["version"], "date": None})
    v.egal("approuvée : confirmée (tous les prix)", b["statut"], "confirmee")
    v.egal("approuvée : montant", b["montant_approuve"], MONTANT_USD_HAUSSE * 10, MAD)
    v.egal("M1 : prix confirmé", _ligne(b, m, "M1")["prix"], 3.0, 1e-9)
    v.egal("M1 : montant de ligne", _ligne(b, m, "M1")["montant"], 43500.0, MAD)


def etape_expedition(m: Monde) -> None:
    """Étape 6 (13 novembre, achats) : jalons du camion ; la date promise ne bouge pas."""
    v, a = m.verif, m.A
    m.jour("expedition")
    jalons = {
        "proforma_le": "2026-11-11",
        "embarquee_le": "2026-11-13",
        "eta": "2026-11-19",
        "conteneur": "TIR 34-SCN-2611",
        "navire": "",
    }
    for cle in ("BC_USD", "BC_EUR"):
        cid = m.ids[cle]
        corps = {"version": a.get(f"/commandes/{cid}")["version"]} | jalons
        r = m.L.appeler("PUT", f"/commandes/{cid}/expedition", corps)
        v.refus(f"{cle} : l'alimentation ne suit pas l'expédition", r, 403)
        b = a.put(f"/commandes/{cid}/expedition", corps)
        v.egal(f"{cle} : ETA", b["expedition"]["eta"], "2026-11-19")
        v.egal(f"{cle} : embarqué", b["expedition"]["embarquee_le"], "2026-11-13")
        v.egal(f"{cle} : date promise inchangée", b["date_livraison_prevue"], "2026-11-25")
        v.egal(f"{cle} : attendu", b["livraison"], "attendu")
    b = a.put(
        f"/commandes/{m.ids['BC_USD']}/expedition",
        {"version": a.get(f"/commandes/{m.ids['BC_USD']}")["version"], "navire": None},
    )
    v.egal("champ absent : inchangé", b["expedition"]["eta"], "2026-11-19")
